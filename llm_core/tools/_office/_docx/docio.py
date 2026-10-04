# SPDX-License-Identifier: MIT
"""Opening, reading and editing existing .docx files."""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

import docx

from .errors import InvalidSpec
from .ooxml.util import qn

_HEADING = re.compile(r"^Heading (\d)$", re.IGNORECASE)


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


def load_bytes(source: str, what: str = "document") -> bytes:
    """Resolve a path, data URI, base64 payload or allowed URL to bytes.

    The rules live in ``assets``: paths are confined to the workspace and the
    configured asset directories, and outbound fetches are refused unless an
    operator turned them on.
    """
    from .assets import load_bytes as _load

    return _load(source, what)


def open_document(source: str, *, notes: list[str] | None = None):
    """Open a document, a catalogue template, or anything in between.

    ``source`` may be a catalogue name, a path inside an allowed directory, or a
    base64 payload. Template *formats* (.dotx, macro-enabled) are normalised on
    the way in, because refusing the exact file an IT department distributes is
    not a defensible position for a server that claims to support templates.
    """
    from . import templates

    blob, template_notes, catalogue_name = templates.load(source)
    if notes is not None:
        if catalogue_name:
            notes.append(f"Opened the {catalogue_name!r} template from the catalogue.")
        notes.extend(template_notes)
    try:
        return docx.Document(io.BytesIO(blob))
    except Exception as exc:
        raise InvalidSpec(f"Could not open the document: {exc}") from exc


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _level_for_name(name: str) -> int | None:
    """The heading level a style *name* implies, or None. The single rule."""
    if name.lower() == "title":
        return 0
    match = _HEADING.match(name)
    return int(match.group(1)) if match else None


def style_levels(document) -> tuple[dict[str, int | None], str | None]:
    """Map every paragraph style id to the heading level its name implies.

    ``paragraph.style`` is the obvious way to ask this, and it is why reading a
    document used to cost 150 ms: for a paragraph with no explicit ``w:pStyle``
    python-docx resolves the default style by scanning the whole styles part,
    once per paragraph. Building the answer for every style up front turns that
    quadratic walk into one pass, and the map is cheap enough (~0.4 ms) to
    rebuild per call — so there is no cache to invalidate when a tool adds a
    style mid-document.

    Names go through ``BabelFish`` because ``w:name`` holds Word's internal name
    ("heading 1") while ``style.name`` — and therefore the rule above — sees the
    UI name ("Heading 1").
    """
    from docx.styles import BabelFish

    levels: dict[str, int | None] = {}
    default_id: str | None = None
    for style in document.styles.element.style_lst:
        # type 1 is WD_STYLE_TYPE.PARAGRAPH; a character style can never be a heading.
        if style.type is not None and int(style.type) != 1:
            continue
        style_id = style.styleId
        if style_id is None:
            continue
        levels[style_id] = _level_for_name(BabelFish.internal2ui(style.name_val or ""))
        if style.default and default_id is None:
            default_id = style_id
    return levels, default_id


def heading_level_for(paragraph, levels: dict[str, int | None], default_id: str | None) -> int | None:
    """``heading_level`` against a prebuilt map. Identical results, ~140x faster."""
    # A direct w:outlineLvl of 9 means "not in the outline" — it is how the
    # TOC's own "Contents" title stays out of the contents. Honour it here so
    # the outline tool and the TOC agree on what counts as a heading.
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is not None:
        outline = properties.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                value = int(outline.get(qn("w:val"), "9"))
            except (TypeError, ValueError):
                value = 9
            return value + 1 if 0 <= value <= 8 else None
        style = properties.find(qn("w:pStyle"))
        if style is not None:
            return levels.get(style.get(qn("w:val")) or "")
    # No w:pStyle at all means the document default applies.
    return levels.get(default_id or "")


def heading_level(paragraph) -> int | None:
    """Kept for callers holding a single paragraph; prefer the map for a whole
    document, which is where the cost actually lives."""
    document = paragraph.part.document
    levels, default_id = style_levels(document)
    return heading_level_for(paragraph, levels, default_id)


def read_text(document, include_tables: bool = True) -> str:
    """Flatten the document to plain text, in reading order."""
    lines: list[str] = []
    body = document.element.body
    tables = {id(t._tbl): t for t in document.tables}
    paragraphs = {id(p._p): p for p in document.paragraphs}
    levels, default_id = style_levels(document)

    for child in body.iterchildren():
        if child.tag == qn("w:p") and id(child) in paragraphs:
            paragraph = paragraphs[id(child)]
            level = heading_level_for(paragraph, levels, default_id)
            text = paragraph.text.strip()
            if not text:
                continue
            lines.append(f"{'#' * max(1, level or 1)} {text}" if level is not None else text)
        elif child.tag == qn("w:tbl") and include_tables and id(child) in tables:
            for row in tables[id(child)].rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def get_outline(document) -> list[dict[str, Any]]:
    outline = []
    levels, default_id = style_levels(document)
    for index, paragraph in enumerate(document.paragraphs):
        level = heading_level_for(paragraph, levels, default_id)
        text = paragraph.text.strip()
        if level is not None and text:
            outline.append({"index": index, "level": level, "text": text})
    return outline


def read_tables(document) -> list[dict[str, Any]]:
    result = []
    for index, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        result.append(
            {
                "index": index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "style": table.style.name if table.style is not None else None,
                "data": rows,
            }
        )
    return result


def describe(document) -> dict[str, Any]:
    properties = document.core_properties
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "sections": len(document.sections),
        "inline_shapes": len(document.inline_shapes),
        "headings": len(get_outline(document)),
        "core_properties": {
            "title": properties.title or None,
            "author": properties.author or None,
            "subject": properties.subject or None,
            "keywords": properties.keywords or None,
            "created": properties.created.isoformat() if properties.created else None,
            "modified": properties.modified.isoformat() if properties.modified else None,
        },
    }


# ---------------------------------------------------------------------------
# Editing
# ---------------------------------------------------------------------------


#: What an unfinished document leaves behind. Checked by default, because the
#: failure they cause — a contract that ships saying "[À COMPLÉTER]" — is the one
#: nobody notices until the recipient does.
LEFTOVER_PATTERNS: tuple[tuple[str, str], ...] = (
    (r"\{\{.*?\}\}", "an unfilled Jinja placeholder"),
    (r"\{%.*?%\}", "an unrendered Jinja tag"),
    (r"<<[^<>\n]{1,60}>>", "an unfilled <<placeholder>>"),
    (r"\[[A-ZÀ-Þ][A-ZÀ-Þ0-9 _'’-]{2,40}\]", "a bracketed placeholder such as [XXX]"),
    (r"\bTODO\b|\bTBD\b|\bFIXME\b", "a TODO left in the text"),
    (r"(?i)\bà compléter\b|\bto be completed\b", "a 'to be completed' marker"),
    (r"(?i)\blorem ipsum\b", "placeholder Lorem ipsum text"),
)


def iter_located_paragraphs(document) -> Iterable[tuple[dict[str, Any], Any]]:
    """Every paragraph in the document, each with where it actually sits.

    ``read_text`` walks the body only, which makes it the wrong tool for checking
    a document: the header is exactly where a stale quarter or an unreplaced
    client name survives a find-and-replace. Everything that verifies uses this
    instead, so what is written and what is checked are the same surface.

    ``paragraph_index`` matches ``document.paragraphs``, so a match can be handed
    straight to ``docx_set_paragraph_text`` or ``docx_delete_paragraph``.
    """
    for index, paragraph in enumerate(document.paragraphs):
        yield {"where": "body", "paragraph_index": index}, paragraph

    for table_index, table in enumerate(document.tables):
        yield from _iter_table(table, {"where": "table", "table_index": table_index})

    for section_index, section in enumerate(document.sections):
        for name, story in _stories(section):
            # Skipping linked stories does two jobs. It stops one stale string in a
            # shared header being reported once per section — and, more importantly,
            # it avoids *reading* a header that does not exist: python-docx creates
            # the part on access, so merely scanning a document would add six empty
            # header and footer parts to it. An empty but defined first-page header
            # is not nothing in Word — it means "no header on the first page", which
            # is how a corporate template silently loses its own.
            if getattr(story, "is_linked_to_previous", False):
                continue
            base = {"where": name, "section": section_index}
            for index, paragraph in enumerate(story.paragraphs):
                yield {**base, "paragraph_index": index}, paragraph
            for table_index, table in enumerate(getattr(story, "tables", []) or []):
                yield from _iter_table(table, {**base, "table_index": table_index})


def _iter_table(table, base: dict[str, Any]) -> Iterable[tuple[dict[str, Any], Any]]:
    # Maps id -> the element itself, and the value is load-bearing: lxml builds a
    # proxy object per element on access and frees it once the last reference
    # goes, so a bare set of ids would let CPython recycle an address and make a
    # distinct cell look like one already seen. That silently dropped table cells
    # from every find and check, depending on when the collector happened to run.
    seen: dict[int, Any] = {}
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            # A merged cell is reported once, at its first position.
            element = cell._tc
            if id(element) in seen:
                continue
            seen[id(element)] = element
            for paragraph in cell.paragraphs:
                yield {**base, "row": row_index, "column": column_index}, paragraph


def _stories(section) -> Iterable[tuple[str, Any]]:
    candidates = (
        ("header", "header"), ("footer", "footer"),
        ("header", "first_page_header"), ("footer", "first_page_footer"),
        ("header", "even_page_header"), ("footer", "even_page_footer"),
    )
    for name, attribute in candidates:
        story = getattr(section, attribute, None)
        if story is None:
            continue
        yield (name if attribute in {"header", "footer"} else attribute), story


def compile_pattern(pattern: str, *, regex: bool, ignore_case: bool) -> re.Pattern:
    flags = re.IGNORECASE if ignore_case else 0
    try:
        return re.compile(pattern if regex else re.escape(pattern), flags)
    except re.error as exc:
        raise InvalidSpec(
            f"{pattern!r} is not a valid regular expression: {exc}. Pass regex=false to "
            "search for it literally."
        ) from exc


def find_pattern(
    document,
    pattern: str,
    *,
    regex: bool = False,
    ignore_case: bool = False,
    limit: int = 200,
) -> list[dict[str, Any]]:
    """Every occurrence of ``pattern``, with where it is and what surrounds it."""
    compiled = compile_pattern(pattern, regex=regex, ignore_case=ignore_case)
    found: list[dict[str, Any]] = []
    for location, paragraph in iter_located_paragraphs(document):
        text = paragraph.text
        if not text:
            continue
        for match in compiled.finditer(text):
            found.append({**location, "match": match.group(0), "excerpt": _excerpt(text, match)})
            if len(found) >= limit:
                return found
    return found


def _excerpt(text: str, match: re.Match, window: int = 40) -> str:
    start = max(0, match.start() - window)
    end = min(len(text), match.end() + window)
    return ("…" if start else "") + text[start:end].strip() + ("…" if end < len(text) else "")


def check_document(
    document,
    *,
    required: list[str] | None = None,
    forbidden: list[str] | None = None,
    regex: bool = False,
    check_leftovers: bool = True,
    ignore_case: bool = False,
) -> dict[str, Any]:
    """Report what a document is missing and what it should not still contain.

    Reports, never blocks: a document mid-draft is *meant* to be incomplete, and a
    server that refuses to hand it back would be wrong about that. The caller
    decides what to do with the findings.
    """
    # One traversal, every pattern. This used to walk the whole document once per
    # required pattern, once per forbidden one, once per leftover pattern, and
    # once more just to count what it had scanned — a dozen passes over the same
    # paragraphs, each of them re-reading header and footer parts.
    #
    #: (compiled, bucket, cap) — the cap is per pattern, matching what the old
    #: per-pattern find_pattern(limit=...) calls enforced individually.
    probes: list[tuple[re.Pattern, list[dict[str, Any]], int]] = []
    required_hits: list[list[dict[str, Any]]] = []
    forbidden_hits: list[list[dict[str, Any]]] = []
    leftover_hits: list[list[dict[str, Any]]] = []

    for pattern in required or []:
        hits: list[dict[str, Any]] = []
        required_hits.append(hits)
        probes.append((compile_pattern(pattern, regex=regex, ignore_case=ignore_case), hits, 1))

    for pattern in forbidden or []:
        hits = []
        forbidden_hits.append(hits)
        probes.append((compile_pattern(pattern, regex=regex, ignore_case=ignore_case), hits, 200))

    if check_leftovers:
        for pattern, _explanation in LEFTOVER_PATTERNS:
            hits = []
            leftover_hits.append(hits)
            # ignore_case stays False regardless of the caller's flag, as it did
            # when these went through find_pattern's default: the patterns that
            # need it carry their own inline (?i).
            probes.append((compile_pattern(pattern, regex=True, ignore_case=False), hits, 20))

    scanned = {"paragraphs": 0, "headers_and_footers": 0, "table_cells": 0}
    for location, paragraph in iter_located_paragraphs(document):
        where = location["where"]
        if where == "body":
            scanned["paragraphs"] += 1
        elif where == "table":
            scanned["table_cells"] += 1
        else:
            scanned["headers_and_footers"] += 1

        text = paragraph.text
        if not text:
            continue
        for compiled, hits, cap in probes:
            if len(hits) >= cap:
                continue
            for match in compiled.finditer(text):
                hits.append(
                    {**location, "match": match.group(0), "excerpt": _excerpt(text, match)}
                )
                if len(hits) >= cap:
                    break

    violations: list[dict[str, Any]] = []
    missing = [
        pattern for pattern, hits in zip(required or [], required_hits) if not hits
    ]

    for pattern, hits in zip(forbidden or [], forbidden_hits):
        if hits:
            violations.append(
                {"pattern": pattern, "kind": "forbidden", "count": len(hits), "occurrences": hits}
            )

    for (pattern, explanation), hits in zip(LEFTOVER_PATTERNS, leftover_hits):
        if hits:
            violations.append(
                {
                    "pattern": pattern,
                    "kind": "leftover",
                    "means": explanation,
                    "count": len(hits),
                    "occurrences": hits,
                }
            )

    return {
        "ok": not missing and not violations,
        "missing": missing,
        "violations": violations,
        "scanned": scanned,
    }


def set_paragraph_text(document, index: int, text: str) -> str:
    """Replace a paragraph's text, keeping its style and its first run's formatting."""
    from .ooxml.util import sanitize_text

    paragraphs = document.paragraphs
    if not 0 <= index < len(paragraphs):
        raise InvalidSpec(
            f"Paragraph index {index} is out of range (the document has {len(paragraphs)}). "
            "Read the document again (docx_read) for the current paragraph numbers."
        )
    paragraph = paragraphs[index]
    previous = paragraph.text
    value = sanitize_text(text)

    runs = paragraph.runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        paragraph.add_run(value)
    return previous


def read_stories(document) -> dict[str, list[dict[str, Any]]]:
    """Header and footer text, which ``read_text`` deliberately does not carry."""
    stories: dict[str, list[dict[str, Any]]] = {"headers": [], "footers": []}
    for location, paragraph in iter_located_paragraphs(document):
        where = location["where"]
        if where in {"body", "table"}:
            continue
        text = paragraph.text.strip()
        if not text:
            continue
        bucket = "headers" if "header" in where else "footers"
        stories[bucket].append({"section": location["section"], "kind": where, "text": text})
    return stories


def iter_all_paragraphs(document) -> Iterable[Any]:
    """Body paragraphs plus everything inside tables, headers and footers.

    One traversal, shared with the verification tools. Two traversals that were
    meant to agree is how a find-and-replace ends up writing somewhere the check
    never looks.
    """
    for _, paragraph in iter_located_paragraphs(document):
        yield paragraph


def replace_text(
    document,
    search: str,
    replacement: str,
    *,
    regex: bool = False,
    ignore_case: bool = False,
    include_headers: bool = True,
) -> int:
    """Search and replace across the document, preserving run formatting.

    Word splits a sentence across runs at arbitrary points (a spell-check pass is
    enough to do it), so a match routinely spans several runs. Matches inside a
    single run are edited in place; a match that spans runs collapses onto the
    first run it touches, which is the only way to keep *some* formatting rather
    than none.

    ``include_headers=False`` excludes headers and footers — and nothing else.
    Table cells are part of the document body as far as a reader is concerned, so
    they are always covered; a replace that quietly skipped every table would be a
    trap, not an option.
    """
    from .ooxml.util import sanitize_text

    replacement = sanitize_text(replacement)
    flags = re.IGNORECASE if ignore_case else 0
    pattern = re.compile(search if regex else re.escape(search), flags)
    if include_headers:
        paragraphs: Iterable[Any] = iter_all_paragraphs(document)
    else:
        paragraphs = (
            paragraph
            for location, paragraph in iter_located_paragraphs(document)
            if location["where"] in {"body", "table"}
        )

    total = 0
    for paragraph in paragraphs:
        total += _replace_in_paragraph(paragraph, pattern, replacement)
    return total


def _replace_in_paragraph(paragraph, pattern: re.Pattern, replacement: str) -> int:
    runs = paragraph.runs
    if not runs:
        return 0
    full = "".join(run.text or "" for run in runs)
    if not pattern.search(full):
        return 0

    # Character span covered by each run.
    spans, cursor = [], 0
    for run in runs:
        length = len(run.text or "")
        spans.append((cursor, cursor + length))
        cursor += length

    replaced = 0
    for match in reversed(list(pattern.finditer(full))):
        start, end = match.span()
        touched = [i for i, (low, high) in enumerate(spans) if low < end and high > start]
        if not touched:
            continue
        first = touched[0]
        low, _ = spans[first]
        text = runs[first].text or ""
        head = text[: start - low]
        tail_run = touched[-1]
        tail_low, _ = spans[tail_run]
        tail = (runs[tail_run].text or "")[end - tail_low :]

        runs[first].text = head + match.expand(replacement) + (tail if tail_run == first else "")
        for index in touched[1:]:
            runs[index].text = "" if index != tail_run else tail
        replaced += 1

        # Offsets past this match are stale, but we iterate backwards so earlier
        # matches keep valid spans; recompute the ones we just touched.
        cursor = 0
        for i, run in enumerate(runs):
            length = len(run.text or "")
            spans[i] = (cursor, cursor + length)
            cursor += length
    return replaced


def delete_paragraph(document, index: int) -> str:
    paragraphs = document.paragraphs
    if not 0 <= index < len(paragraphs):
        raise InvalidSpec(
            f"Paragraph index {index} is out of range (the document has {len(paragraphs)})."
        )
    paragraph = paragraphs[index]
    text = paragraph.text
    element = paragraph._p
    element.getparent().remove(element)
    return text


def set_metadata(document, **values: Any) -> dict[str, Any]:
    from .ooxml.util import sanitize_text

    properties = document.core_properties
    applied = {}
    for key, value in values.items():
        if value is None:
            continue
        value = sanitize_text(value)
        if not hasattr(properties, key):
            raise InvalidSpec(
                f"Unknown metadata field {key!r}. Valid fields: title, subject, author, "
                "keywords, category, comments, language."
            )
        setattr(properties, key, value)
        applied[key] = value
    return applied
