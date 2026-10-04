# SPDX-License-Identifier: MIT
"""Headers, footers, page numbering and watermarks."""

from __future__ import annotations

from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Cm

from .fields import add_text_with_fields
from .layout import content_width_cm
from .util import align_paragraph, frag, make, parse_color, qn, sanitize_text

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"

#: The VML text-path shape Word uses for every watermark. Readers other than Word
#: need the shapetype spelled out, so we ship it with the shape.
_WORDART_SHAPETYPE = (
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e">'
    "<v:formulas>"
    '<v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/>'
    '<v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/>'
    '<v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/>'
    '<v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/>'
    '<v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/>'
    "</v:formulas>"
    '<v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/>'
    '<v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>'
    "</v:shapetype>"
)


def _story(section, which: str, variant: str):
    """Return the header/footer story object for a section."""
    attribute = {
        ("header", "default"): "header",
        ("header", "first"): "first_page_header",
        ("header", "even"): "even_page_header",
        ("footer", "default"): "footer",
        ("footer", "first"): "first_page_footer",
        ("footer", "even"): "even_page_footer",
    }.get((which, variant))
    if attribute is None:
        raise ValueError(f"variant must be 'default', 'first' or 'even', got {variant!r}.")
    story = getattr(section, attribute)
    if variant == "first":
        section.different_first_page_header_footer = True
    if variant == "even":
        _enable_even_odd(section)
    story.is_linked_to_previous = False
    return story


def _enable_even_odd(section) -> None:
    """``w:evenAndOddHeaders`` is a document-level setting, not a section one."""
    from .fields import SETTINGS_ORDER
    from .util import insert_in_order

    settings = section.part.package.main_document_part.document.settings.element
    for existing in settings.findall(qn("w:evenAndOddHeaders")):
        settings.remove(existing)
    insert_in_order(settings, make("w:evenAndOddHeaders", **{"w:val": "true"}), SETTINGS_ORDER)


def _clear(story) -> None:
    for paragraph in list(story.paragraphs):
        element = paragraph._p
        element.getparent().remove(element)


def set_header_footer(
    section,
    which: str = "header",
    *,
    variant: str = "default",
    text: str | None = None,
    left: str | None = None,
    center: str | None = None,
    right: str | None = None,
    align: str | None = None,
    style: str | None = None,
    font_size_pt: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    rule: bool = False,
    clear: bool = True,
):
    """Write a header or footer.

    Passing ``left`` / ``center`` / ``right`` lays the three out on one line using
    tab stops sized to the section's text width, which is how Word itself does it.
    ``{PAGE}`` and ``{NUMPAGES}`` tokens anywhere in the text become live fields.
    """
    if which not in {"header", "footer"}:
        raise ValueError(f"which must be 'header' or 'footer', got {which!r}.")
    text, left, center, right = (
        sanitize_text(v) if isinstance(v, str) else v for v in (text, left, center, right)
    )
    story = _story(section, which, variant)
    if clear:
        _clear(story)

    paragraph = story.add_paragraph()
    if style:
        try:
            paragraph.style = style
        except KeyError:
            pass

    run_format = {
        "size_pt": font_size_pt,
        "color": color,
        "bold": bold,
    }
    run_format = {k: v for k, v in run_format.items() if v is not None}

    if left is not None or center is not None or right is not None:
        width = content_width_cm(section)
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Cm(width / 2), WD_TAB_ALIGNMENT.CENTER)
        tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
        segments = [left or "", center or "", right or ""]
        for index, segment in enumerate(segments):
            if index:
                paragraph.add_run("\t")
            if segment:
                add_text_with_fields(paragraph, segment, **run_format)
    elif text is not None:
        add_text_with_fields(paragraph, text, **run_format)
        align_paragraph(paragraph, align)

    if rule:
        from .fields import horizontal_rule

        horizontal_rule(paragraph, color=parse_color(color) or "BFBFBF")
    return paragraph


def add_page_numbers(
    section,
    template: str = "{PAGE} / {NUMPAGES}",
    *,
    align: str = "center",
    variant: str = "default",
    font_size_pt: float | None = 9,
    color: str | None = "808080",
) -> None:
    """Convenience wrapper: a footer that is just a page-number expression."""
    paragraph = set_header_footer(
        section,
        "footer",
        variant=variant,
        text=template,
        align=align,
        font_size_pt=font_size_pt,
        color=color,
    )
    align_paragraph(paragraph, align)


def add_watermark(
    section,
    text: str,
    *,
    color: str = "D9D9D9",
    font: str = "Calibri",
    rotation: int = -45,
    width_pt: float = 468,
    height_pt: float = 117,
    variant: str = "default",
    opacity: float | None = None,
) -> None:
    """Stamp diagonal WordArt across every page of the section.

    A Word watermark lives in the header story, not the body — that is what makes
    it repeat on every page and sit behind the text.
    """
    story = _story(section, "header", variant)
    # Word's watermark gallery replaces the current watermark; stacking a second
    # WordArt shape renders both superimposed. Any paragraph holding a VML
    # text-path is a watermark — ours or one authored in Word.
    for existing in list(story.paragraphs):
        if existing._p.find(f".//{{{V_NS}}}textpath") is not None:
            existing._p.getparent().remove(existing._p)
    paragraph = story.add_paragraph()
    fill = parse_color(color) or "D9D9D9"
    escaped = (
        sanitize_text(str(text))
        .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    )
    fill_opacity = "" if opacity is None else f' opacity="{max(0.0, min(1.0, float(opacity)))}"'
    style = (
        "position:absolute;margin-left:0;margin-top:0;"
        f"width:{width_pt}pt;height:{height_pt}pt;rotation:{int(rotation)};"
        "z-index:-251654144;mso-position-horizontal:center;"
        "mso-position-horizontal-relative:margin;mso-position-vertical:center;"
        "mso-position-vertical-relative:margin"
    )
    pict = frag(
        f'<w:pict xmlns:w="{W_NS}" xmlns:v="{V_NS}" xmlns:o="{O_NS}">'
        f"{_WORDART_SHAPETYPE}"
        f'<v:shape id="docxMcpWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
        f'style="{style}" fillcolor="#{fill}"{fill_opacity} stroked="f">'
        f'<v:textpath style="font-family:&quot;{font}&quot;;font-size:1pt" string="{escaped}"/>'
        "</v:shape></w:pict>"
    )
    paragraph.add_run()._r.append(pict)


def add_header_image(section, image_path, *, width_cm: float = 3.0, align: str = "right",
                     variant: str = "default"):
    """Put a logo in the header."""
    story = _story(section, "header", variant)
    paragraph = story.add_paragraph()
    paragraph.add_run().add_picture(image_path, width=Cm(float(width_cm)))
    align_paragraph(paragraph, align)
    return paragraph
