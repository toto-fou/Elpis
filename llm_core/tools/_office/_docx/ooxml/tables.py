# SPDX-License-Identifier: MIT
"""Table plumbing python-docx leaves out: shading, borders, widths, header repeat."""

from __future__ import annotations

from typing import Any, Iterable, Sequence

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Cm

from .util import frag, insert_in_order, make, parse_color, qn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

#: CT_TcPr sequence.
TCPR_ORDER = (
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers",
)
#: CT_TrPr sequence.
TRPR_ORDER = (
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
    "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
)
#: CT_TblPr sequence.
TBLPR_ORDER = (
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
)

_VALIGN = {
    "top": WD_ALIGN_VERTICAL.TOP,
    "center": WD_ALIGN_VERTICAL.CENTER,
    "centre": WD_ALIGN_VERTICAL.CENTER,
    "middle": WD_ALIGN_VERTICAL.CENTER,
    "bottom": WD_ALIGN_VERTICAL.BOTTOM,
}

BORDER_EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")


def shade_cell(cell, color: str | None) -> None:
    """Fill a cell with a solid colour."""
    fill = parse_color(color)
    if not fill:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shading = make("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})
    insert_in_order(tc_pr, shading, TCPR_ORDER)


def set_cell_vertical_alignment(cell, value: str | None) -> None:
    if not value:
        return
    try:
        cell.vertical_alignment = _VALIGN[str(value).lower()]
    except KeyError:
        raise ValueError(f"Unknown vertical alignment {value!r}. Use top, center or bottom.") from None


def set_cell_borders(cell, *, color: str = "BFBFBF", size: int = 4, style: str = "single",
                     edges: Iterable[str] = ("top", "left", "bottom", "right")) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(existing)
    hex_color = parse_color(color) or "BFBFBF"
    parts = "".join(
        f'<w:{edge} w:val="{style}" w:sz="{int(size)}" w:space="0" w:color="{hex_color}"/>'
        for edge in edges
        if edge in BORDER_EDGES
    )
    insert_in_order(tc_pr, frag(f'<w:tcBorders xmlns:w="{W_NS}">{parts}</w:tcBorders>'), TCPR_ORDER)


def set_table_borders(
    table,
    *,
    color: str = "BFBFBF",
    size: int = 4,
    style: str = "single",
    edges: Iterable[str] = BORDER_EDGES,
) -> None:
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(existing)
    hex_color = parse_color(color) or "BFBFBF"
    parts = "".join(
        f'<w:{edge} w:val="{style}" w:sz="{int(size)}" w:space="0" w:color="{hex_color}"/>'
        for edge in edges
        if edge in BORDER_EDGES
    )
    insert_in_order(tbl_pr, frag(f'<w:tblBorders xmlns:w="{W_NS}">{parts}</w:tblBorders>'), TBLPR_ORDER)


def clear_table_borders(table) -> None:
    set_table_borders(table, style="none", size=0, edges=BORDER_EDGES)


def set_column_widths(table, widths_cm: Sequence[float | None]) -> None:
    """Fix column widths.

    Word only honours per-cell widths when autofit is off, and it needs the width
    written on *every* cell of the column — setting it on the column object alone
    is silently ignored.
    """
    table.autofit = False
    try:
        table.allow_autofit = False
    except AttributeError:  # pragma: no cover — older python-docx
        pass
    for index, width in enumerate(widths_cm):
        if width is None or index >= len(table.columns):
            continue
        emu = Cm(float(width))
        table.columns[index].width = emu
        for cell in table.columns[index].cells:
            cell.width = emu


def repeat_header_row(row) -> None:
    """Mark a row as a header so it repeats when the table spans pages."""
    tr_pr = row._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:tblHeader")):
        tr_pr.remove(existing)
    insert_in_order(tr_pr, make("w:tblHeader", **{"w:val": "true"}), TRPR_ORDER)


def set_row_cant_split(row, value: bool = True) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:cantSplit")):
        tr_pr.remove(existing)
    if value:
        insert_in_order(tr_pr, make("w:cantSplit"), TRPR_ORDER)


def set_cell_margins(table, *, top: float = 0.1, bottom: float = 0.1,
                     left: float = 0.19, right: float = 0.19) -> None:
    """Uniform inner padding, in centimetres."""
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(existing)
    parts = "".join(
        f'<w:{edge} w:w="{int(Cm(float(value)).twips)}" w:type="dxa"/>'
        for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right))
    )
    insert_in_order(tbl_pr, frag(f'<w:tblCellMar xmlns:w="{W_NS}">{parts}</w:tblCellMar>'), TBLPR_ORDER)


def set_table_look(table, *, first_row: bool = True, first_column: bool = False,
                   banded_rows: bool = True, banded_columns: bool = False) -> None:
    """Tell the table style which conditional formats to apply.

    Without ``w:tblLook`` a built-in style like "Light Grid Accent 1" renders flat:
    the style defines header and banding formats, and this element enables them.
    """
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblLook")):
        tbl_pr.remove(existing)
    element = make(
        "w:tblLook",
        **{
            "w:val": "04A0",
            "w:firstRow": "1" if first_row else "0",
            "w:lastRow": "0",
            "w:firstColumn": "1" if first_column else "0",
            "w:lastColumn": "0",
            "w:noHBand": "0" if banded_rows else "1",
            "w:noVBand": "0" if banded_columns else "1",
        },
    )
    insert_in_order(tbl_pr, element, TBLPR_ORDER)


def merge_region(table, row_start: int, col_start: int, row_end: int, col_end: int):
    """Merge a rectangular region, given inclusive 0-based grid coordinates."""
    rows, cols = len(table.rows), len(table.columns)
    for name, value, limit in (
        ("row_start", row_start, rows), ("row_end", row_end, rows),
        ("col_start", col_start, cols), ("col_end", col_end, cols),
    ):
        if not 0 <= value < limit:
            raise ValueError(
                f"merge {name}={value} is outside the table "
                f"({rows} rows x {cols} columns)."
            )
    if row_end < row_start or col_end < col_start:
        raise ValueError("merge end coordinates must not be before the start coordinates.")
    # python-docx concatenates the content of every merged cell. A grid spec
    # repeats the anchor value in the covered cells, so without this the result
    # reads "EMEA\nEMEA"; only genuinely different content is worth stacking.
    anchor = table.cell(row_start, col_start)
    anchor_text = anchor.text.strip()
    # Maps id -> the element itself, and the value is load-bearing: lxml frees an
    # element proxy once the last reference goes, so a bare set of ids would let
    # CPython recycle an address and skip a cell that was never seen — leaving
    # exactly the duplicated "EMEA\nEMEA" content this loop exists to clear.
    seen: dict[int, Any] = {id(anchor._tc): anchor._tc}
    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            cell = table.cell(row, col)
            element = cell._tc
            if id(element) in seen:
                continue
            seen[id(element)] = element
            if cell.text.strip() in ("", anchor_text):
                cell.text = ""
    merged = anchor.merge(table.cell(row_end, col_end))
    _drop_trailing_empty_paragraphs(merged)
    return merged


def _drop_trailing_empty_paragraphs(cell) -> None:
    """Swallowed cells each leave an empty w:p behind, which pads the row height."""
    tc = cell._tc
    paragraphs = tc.findall(qn("w:p"))
    while len(paragraphs) > 1:
        last = paragraphs[-1]
        has_content = "".join(last.itertext()).strip() or any(
            last.find(f".//{qn(tag)}") is not None
            for tag in ("w:drawing", "w:pict", "w:object")
        )
        if has_content:
            break
        tc.remove(last)
        paragraphs.pop()
