# SPDX-License-Identifier: MIT
"""The engine: turn a :class:`DocumentSpec` into a python-docx ``Document``.

One block type, one method. Everything a block can need — a chart part, a
numbering definition, a caption counter — is reached through :class:`BuildContext`
so the granular tools in ``docx_mcp.tools`` can build a single block against an
existing document using exactly the same code path as ``docx_build``.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Iterable, Sequence

import docx
from docx.shared import Cm, Pt

from ..errors import InvalidSpec
from ..ooxml import (
    charts as chart_ooxml,
    fields,
    headers as header_ooxml,
    layout,
    lists,
    notes,
    refs,
    tables as table_ooxml,
)
from ..ooxml.util import align_paragraph, apply_run_format, make, parse_color, qn, rgb, table_alignment
from . import images as image_io, spec as S

MONO_SHADING = "F4F4F2"

CALLOUT_STYLES = {
    "info": ("2A78D6", "EAF2FC", "ℹ"),
    "note": ("4A3AA7", "EFEDF9", "✎"),
    "success": ("008300", "E9F5EA", "✔"),
    "warning": ("EDA100", "FDF4E3", "▲"),
    "danger": ("E34948", "FCECEC", "■"),
}


@dataclass
class BuildContext:
    """Shared state for a build. Also what the granular add_* tools carry."""

    document: Any
    theme: S.ThemeSpec = field(default_factory=S.ThemeSpec)
    notes: list[str] = field(default_factory=list)

    def note(self, message: str) -> None:
        if message not in self.notes:
            self.notes.append(message)

    @property
    def section(self):
        return self.document.sections[-1]

    def content_width_cm(self) -> float:
        return layout.content_width_cm(self.section)

    def palette(self) -> Sequence[str]:
        return self.theme.chart_palette or chart_ooxml.DEFAULT_PALETTE


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------


def build_document(spec: S.DocumentSpec, base: bytes | None = None,
                   keep_base_content: bool = False) -> tuple[Any, BuildContext]:
    """Construit le document ; ``base`` = octets d'un modèle (.docx/.dotx normalisé)
    dont on garde styles, numérotation, en-têtes et pieds, et mise en page —
    sauf ce que ``spec`` redéfinit (``page``, ``header``, ``footer`` non nuls).
    Le texte du modèle est retiré sauf ``keep_base_content``."""
    document = docx.Document(io.BytesIO(base)) if base else docx.Document()
    if base and not keep_base_content:
        body = document.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    context = BuildContext(document=document, theme=spec.theme)

    apply_metadata(document, spec.meta)
    apply_theme(document, spec.theme)
    if spec.page is not None:
        apply_page(document.sections[0], spec.page)
    if spec.header:
        apply_header(document.sections[0], spec.header)
    if spec.footer:
        apply_footer(document.sections[0], spec.footer)
    if spec.watermark:
        header_ooxml.add_watermark(
            document.sections[0],
            spec.watermark.text,
            color=spec.watermark.color,
            font=spec.watermark.font,
            rotation=spec.watermark.rotation,
        )

    add_blocks(context, spec.blocks)
    finalize(document)
    return document, context


def finalize(document) -> None:
    """Everything that must happen after the last block: fields, then field refresh."""
    fields.populate_toc(document)
    fields.set_update_fields_on_open(document)


def apply_metadata(document, meta: S.MetaSpec) -> None:
    properties = document.core_properties
    for attribute in ("title", "subject", "author", "keywords", "category", "comments", "language"):
        value = getattr(meta, attribute, None)
        if value:
            setattr(properties, attribute, value)


def apply_theme(document, theme: S.ThemeSpec) -> None:
    """Push the theme onto the style definitions, so it applies to everything."""
    styles = document.styles
    if theme.body_font or theme.base_font_size_pt:
        normal = styles["Normal"]
        if theme.body_font:
            normal.font.name = theme.body_font
            _set_style_east_asian(normal, theme.body_font)
        if theme.base_font_size_pt:
            normal.font.size = Pt(float(theme.base_font_size_pt))

    accent = parse_color(theme.accent) if theme.accent else None
    for name in ("Title", "Subtitle", *(f"Heading {i}" for i in range(1, 10))):
        try:
            style = styles[name]
        except KeyError:
            continue
        if theme.heading_font:
            style.font.name = theme.heading_font
            _set_style_east_asian(style, theme.heading_font)
        if accent and name not in {"Subtitle"}:
            style.font.color.rgb = rgb(accent)


def _set_style_east_asian(style, name: str) -> None:
    from ..ooxml.util import qn, sub

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = sub(rpr, "w:rFonts")
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def apply_page(section, page: S.PageSpec) -> None:
    layout.apply_page_setup(
        section,
        size=page.size,
        width_cm=page.width_cm,
        height_cm=page.height_cm,
        orientation=page.orientation,
        margins_cm=page.margins_cm,
        columns=page.columns,
        column_spacing_cm=page.column_spacing_cm,
    )


def apply_header(section, header: S.HeaderSpec) -> None:
    if header.image:
        stream, _ = image_io.resolve_image(header.image)
        header_ooxml.add_header_image(section, stream, width_cm=header.image_width_cm)
    if any(v is not None for v in (header.text, header.left, header.center, header.right)):
        header_ooxml.set_header_footer(
            section, "header", text=header.text, left=header.left, center=header.center,
            right=header.right, align=header.align, font_size_pt=header.font_size_pt,
            color=header.color, rule=header.rule, clear=not header.image,
        )
    if header.different_first_page:
        section.different_first_page_header_footer = True
        header_ooxml.set_header_footer(
            section, "header", variant="first", text=header.first_page_text or "",
            font_size_pt=header.font_size_pt, color=header.color,
        )


def apply_footer(section, footer: S.FooterSpec) -> None:
    text = footer.text
    if footer.page_numbers and not any((footer.text, footer.left, footer.center, footer.right)):
        text = footer.page_number_format
    header_ooxml.set_header_footer(
        section, "footer", text=text, left=footer.left, center=footer.center,
        right=footer.right, align=footer.align, font_size_pt=footer.font_size_pt,
        color=footer.color, rule=footer.rule,
    )
    if footer.different_first_page:
        section.different_first_page_header_footer = True


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------


def add_blocks(context: BuildContext, blocks: Iterable[S.Block]) -> int:
    count = 0
    for block in blocks:
        add_block(context, block)
        count += 1
    return count


def add_block(context: BuildContext, block: S.Block):
    handler = _HANDLERS.get(type(block))
    if handler is None:  # pragma: no cover — the discriminated union prevents this
        raise InvalidSpec(f"No handler for block type {type(block).__name__}.")
    return handler(context, block)


# -- text ------------------------------------------------------------------------


def _write_runs(context: BuildContext, paragraph, text: str | None, runs: Sequence[S.RunSpec]):
    """Fill a paragraph from either a plain string or a list of formatted runs."""
    if runs:
        for run_spec in runs:
            _write_run(context, paragraph, run_spec)
    elif text:
        paragraph.add_run(text)
    return paragraph


def _write_run(context: BuildContext, paragraph, run_spec: S.RunSpec):
    # Drop unset fields: passing color=None through to add_hyperlink would override
    # its blue-and-underlined default with "no formatting".
    formatting = {
        key: value
        for key, value in run_spec.model_dump(exclude={"text", "link", "code", "footnote"}).items()
        if value is not None
    }
    if run_spec.code:
        formatting["font"] = formatting.get("font") or context.theme.mono_font
        formatting["highlight"] = formatting.get("highlight") or MONO_SHADING

    if run_spec.link:
        link = str(run_spec.link)
        if link.startswith("#"):
            run = refs.add_hyperlink(
                context.document, paragraph, run_spec.text, anchor=link[1:], **formatting
            )
        else:
            run = refs.add_hyperlink(
                context.document, paragraph, run_spec.text, url=link, **formatting
            )
    else:
        run = paragraph.add_run(run_spec.text)
        apply_run_format(run, formatting)

    if run_spec.footnote:
        notes.add_footnote(context.document, paragraph, run_spec.footnote)
    return run


def _heading(context: BuildContext, block: S.HeadingBlock):
    document = context.document
    paragraph = document.add_heading("", level=block.level)
    if block.style:
        try:
            paragraph.style = block.style
        except KeyError:
            context.note(f"Style {block.style!r} is not in this document; used the heading style.")
    _write_runs(context, paragraph, block.text, block.runs)
    align_paragraph(paragraph, block.align)
    if block.color:
        for run in paragraph.runs:
            run.font.color.rgb = rgb(block.color)
    if block.page_break_before:
        paragraph.paragraph_format.page_break_before = True
    if block.bookmark:
        refs.add_bookmark(document, paragraph, block.bookmark)
    return paragraph


def _paragraph(context: BuildContext, block: S.ParagraphBlock):
    paragraph = context.document.add_paragraph()
    if block.style:
        try:
            paragraph.style = block.style
        except KeyError:
            context.note(f"Style {block.style!r} is not in this document; used Normal.")
    _write_runs(context, paragraph, block.text, block.runs)
    _apply_paragraph_format(paragraph, block)
    if block.bookmark:
        refs.add_bookmark(context.document, paragraph, block.bookmark)
    return paragraph


def _apply_paragraph_format(paragraph, block) -> None:
    align_paragraph(paragraph, getattr(block, "align", None))
    fmt = paragraph.paragraph_format
    if getattr(block, "space_before_pt", None) is not None:
        fmt.space_before = Pt(float(block.space_before_pt))
    if getattr(block, "space_after_pt", None) is not None:
        fmt.space_after = Pt(float(block.space_after_pt))
    if getattr(block, "line_spacing", None) is not None:
        fmt.line_spacing = float(block.line_spacing)
    if getattr(block, "indent_left_cm", None) is not None:
        fmt.left_indent = Cm(float(block.indent_left_cm))
    if getattr(block, "indent_right_cm", None) is not None:
        fmt.right_indent = Cm(float(block.indent_right_cm))
    if getattr(block, "first_line_indent_cm", None) is not None:
        fmt.first_line_indent = Cm(float(block.first_line_indent_cm))
    if getattr(block, "keep_with_next", None) is not None:
        fmt.keep_with_next = bool(block.keep_with_next)
    if getattr(block, "page_break_before", False):
        fmt.page_break_before = True


def _list(context: BuildContext, block: S.ListBlock):
    if not block.items:
        return None
    num_id = lists.create_numbering(context.document, ordered=block.ordered)
    written = []
    for item in block.items:
        item_spec = S.ListItemSpec(text=item) if isinstance(item, str) else item
        paragraph = context.document.add_paragraph()
        if block.style:
            try:
                paragraph.style = block.style
            except KeyError:
                context.note(f"Style {block.style!r} is not in this document; used Normal.")
        _write_runs(context, paragraph, item_spec.text, item_spec.runs)
        lists.apply_list_item(paragraph, num_id, item_spec.level)
        if block.space_after_pt is not None:
            paragraph.paragraph_format.space_after = Pt(float(block.space_after_pt))
        written.append(paragraph)
    return written


def _code(context: BuildContext, block: S.CodeBlock):
    paragraph = context.document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.left_indent = Cm(0.4)
    lines = (block.text or "").split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        apply_run_format(
            run, {"font": block.font or context.theme.mono_font, "size_pt": block.size_pt}
        )
    _shade_paragraph(paragraph, MONO_SHADING)
    if block.caption:
        refs.add_caption(context.document, block.caption, label="Listing", align="left")
    return paragraph


def _shade_paragraph(paragraph, color: str) -> None:
    from ..ooxml.util import frag, insert_in_order, qn

    ppr = paragraph._p.get_or_add_pPr()
    for existing in ppr.findall(qn("w:shd")):
        ppr.remove(existing)
    shading = frag(
        '<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'w:val="clear" w:color="auto" w:fill="{parse_color(color)}"/>'
    )
    insert_in_order(ppr, shading, fields.PPR_ORDER)


def _quote(context: BuildContext, block: S.QuoteBlock):
    paragraph = context.document.add_paragraph()
    if block.style:
        try:
            paragraph.style = block.style
        except KeyError:
            paragraph.paragraph_format.left_indent = Cm(1.0)
            context.note(f"Style {block.style!r} is not in this document; indented instead.")
    paragraph.add_run(block.text)
    if block.author:
        attribution = context.document.add_paragraph()
        run = attribution.add_run(f"— {block.author}")
        apply_run_format(run, {"italic": True, "color": "808080", "size_pt": 9})
        attribution.paragraph_format.left_indent = Cm(1.0)
    return paragraph


def _separate_from_previous_table(document) -> None:
    """Two adjacent w:tbl elements read as ONE table in Word — keep them apart."""
    from ..ooxml.util import qn

    children = [
        child for child in document.element.body if child.tag != qn("w:sectPr")
    ]
    if children and children[-1].tag == qn("w:tbl"):
        document.add_paragraph()


def _callout(context: BuildContext, block: S.CalloutBlock):
    """A one-cell shaded table — the reliable way to draw a boxed note in Word."""
    accent, fill, glyph = CALLOUT_STYLES[block.variant]
    _separate_from_previous_table(context.document)
    table = context.document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    cell.text = ""
    table_ooxml.shade_cell(cell, fill)
    table_ooxml.set_cell_borders(cell, color=fill, size=2)
    # A thick left edge is what makes it read as a callout rather than a table.
    table_ooxml.set_cell_borders(cell, color=accent, size=18, edges=("left",))
    table_ooxml.set_cell_margins(table, top=0.18, bottom=0.18, left=0.28, right=0.28)
    # Un encadré ne se coupe pas entre deux pages (titre seul en bas de page).
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(make("w:cantSplit"))

    paragraph = cell.paragraphs[0]
    if block.title:
        title_run = paragraph.add_run(f"{glyph}  {block.title}")
        apply_run_format(title_run, {"bold": True, "color": accent})
        paragraph = cell.add_paragraph()
    elif glyph:
        apply_run_format(paragraph.add_run(f"{glyph}  "), {"bold": True, "color": accent})
    paragraph.add_run(block.text)
    context.document.add_paragraph()  # breathing room after the box
    return table


# -- tables ----------------------------------------------------------------------


def _check_merges_do_not_overlap(merges) -> None:
    """Merging into an already-merged cell chains the spans: two regions that
    share even one cell silently collapse into a single giant cell."""
    for index, merge in enumerate(merges):
        for other in merges[:index]:
            rows_meet = merge.row_start <= other.row_end and other.row_start <= merge.row_end
            cols_meet = merge.col_start <= other.col_end and other.col_start <= merge.col_end
            if rows_meet and cols_meet:
                raise InvalidSpec(
                    f"Merge regions overlap: rows {other.row_start}-{other.row_end} x "
                    f"cols {other.col_start}-{other.col_end} and rows "
                    f"{merge.row_start}-{merge.row_end} x cols "
                    f"{merge.col_start}-{merge.col_end} share at least one cell."
                )


def _table(context: BuildContext, block: S.TableBlock):
    grid: list[list[Any]] = []
    if block.header:
        grid.append(list(block.header))
    grid.extend([list(row) for row in block.rows])
    if not grid:
        raise InvalidSpec("A table block needs 'header' and/or 'rows'.")

    columns = max(len(row) for row in grid)
    # Validated before anything is inserted: this used to raise from inside the
    # merge loop, by which point a half-built table was already in the document
    # and the next successful call would save it.
    _check_merges_do_not_overlap(block.merges)
    _separate_from_previous_table(context.document)
    table = context.document.add_table(rows=len(grid), cols=columns)

    if block.style:
        try:
            table.style = block.style
        except KeyError:
            context.note(
                f"Table style {block.style!r} is not in this document; used Table Grid."
            )
            try:
                table.style = "Table Grid"
            except KeyError:
                pass

    for row_index, row in enumerate(grid):
        for col_index in range(columns):
            value = row[col_index] if col_index < len(row) else None
            is_header = bool(block.header) and row_index == 0
            _fill_cell(
                context,
                table.cell(row_index, col_index),
                value,
                header=is_header,
                block=block,
                first_column=col_index == 0,
            )

    table_ooxml.set_table_look(
        table, first_row=bool(block.header), banded_rows=block.banded_rows,
        first_column=block.first_column_bold,
    )
    table_ooxml.set_cell_margins(table)
    if block.widths_cm:
        table_ooxml.set_column_widths(table, block.widths_cm)
    elif not block.autofit:
        width = context.content_width_cm()
        table_ooxml.set_column_widths(table, [width / columns] * columns)
    alignment = table_alignment(block.align)
    if alignment is not None:
        table.alignment = alignment
    if block.header and block.repeat_header:
        table_ooxml.repeat_header_row(table.rows[0])

    for merge in block.merges:
        table_ooxml.merge_region(
            table, merge.row_start, merge.col_start, merge.row_end, merge.col_end
        )

    if block.caption:
        refs.add_caption(context.document, block.caption, label="Table", bookmark=block.bookmark)
    elif block.bookmark:
        refs.add_bookmark(context.document, context.document.add_paragraph(), block.bookmark)
    return table


def _fill_cell(context, cell, value, *, header: bool, block: S.TableBlock, first_column: bool):
    cell_spec = value if isinstance(value, S.CellSpec) else S.CellSpec(text=_as_text(value))
    paragraph = cell.paragraphs[0]
    paragraph.text = ""

    if cell_spec.runs:
        for run_spec in cell_spec.runs:
            _write_run(context, paragraph, run_spec)
    else:
        run = paragraph.add_run(cell_spec.text or "")
        apply_run_format(
            run,
            {
                "bold": cell_spec.bold if cell_spec.bold is not None
                else (True if header or (block.first_column_bold and first_column) else None),
                "italic": cell_spec.italic,
                "color": cell_spec.color or (block.header_color if header else None),
                "size_pt": cell_spec.size_pt or block.font_size_pt,
            },
        )
    align_paragraph(paragraph, cell_spec.align)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)

    fill = cell_spec.fill or (block.header_fill if header else None)
    if fill:
        table_ooxml.shade_cell(cell, fill)
    table_ooxml.set_cell_vertical_alignment(cell, cell_spec.valign or "center")


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# -- media -----------------------------------------------------------------------


def _image(context: BuildContext, block: S.ImageBlock):
    stream, pixel_size = image_io.resolve_image(block.source)
    width_cm, height_cm = image_io.fit_size(
        pixel_size, block.width_cm, block.height_cm, context.content_width_cm()
    )
    paragraph = context.document.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(
        stream,
        width=Cm(width_cm) if width_cm else None,
        height=Cm(height_cm) if height_cm else None,
    )
    align_paragraph(paragraph, block.align)
    if block.alt_text:
        _set_alt_text(run, block.alt_text)
    if block.caption:
        refs.add_caption(context.document, block.caption, label="Figure", bookmark=block.bookmark)
    elif block.bookmark:
        refs.add_bookmark(context.document, paragraph, block.bookmark)
    return paragraph


def _set_alt_text(run, text: str) -> None:
    from ..ooxml.util import qn

    for doc_pr in run._r.iter(qn("wp:docPr")):
        doc_pr.set("descr", text)
        doc_pr.set("title", text[:100])


def _chart(context: BuildContext, block: S.ChartBlock):
    # Graphiques natifs seulement : les types sans équivalent OOXML arrivent
    # déjà en image (rendu ECharts côté serveur, voir ``_office/graphiques.py``).
    key = str(block.chart_type or "").strip().lower().replace("-", "_").replace(" ", "_")
    if key not in chart_ooxml.NATIVE_TYPES:
        raise InvalidSpec(
            f"Unknown chart type {block.chart_type!r}. Native types: "
            f"{', '.join(chart_ooxml.native_chart_types())}."
        )
    request = chart_ooxml.ChartRequest(
        chart_type=key,
        categories=list(block.categories),
        series=[
            chart_ooxml.SeriesData(
                name=s.name, values=s.values, points=s.points, color=s.color
            )
            for s in block.series
        ],
        title=block.title,
        x_title=block.x_title,
        y_title=block.y_title,
        legend=block.legend,
        data_labels=block.data_labels,
        data_label_position=block.data_label_position,
        data_label_format=block.data_label_format,
        number_format=block.number_format,
        gridlines=block.gridlines,
        y_min=block.y_min,
        y_max=block.y_max,
        width_cm=min(block.width_cm, context.content_width_cm()),
        height_cm=block.height_cm,
        palette=block.palette or context.theme.chart_palette,
        gap_width=block.gap_width,
        overlap=block.overlap,
        hole_size=block.hole_size,
        smooth=block.smooth,
    )
    paragraph = chart_ooxml.add_native_chart(context.document, request)
    align_paragraph(paragraph, "center")

    if len(block.series) > len(context.palette()):
        context.note(
            f"{len(block.series)} series exceeds the {len(context.palette())}-colour palette, so "
            "colours repeat. Consider grouping the tail into an 'Other' series or splitting the "
            "chart into small multiples."
        )

    if block.caption:
        refs.add_caption(context.document, block.caption, label="Figure", bookmark=block.bookmark)
    return paragraph


# -- structure -------------------------------------------------------------------


def _toc(context: BuildContext, block: S.TocBlock):
    if block.title:
        heading = context.document.add_heading(block.title, level=1)
        # A TOC heading that lists itself is noise.
        heading.paragraph_format.element.get_or_add_pPr()
        _suppress_outline(heading)
    start = fields.add_toc(context.document, tuple(block.levels))
    context.note(fields.TOC_NOTE)
    if block.page_break_after:
        fields.add_page_break(context.document.add_paragraph())
    return start


def _suppress_outline(paragraph) -> None:
    """Keep a heading out of the table of contents."""
    from ..ooxml.util import insert_in_order, make, qn

    ppr = paragraph._p.get_or_add_pPr()
    for existing in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(existing)
    insert_in_order(ppr, make("w:outlineLvl", **{"w:val": "9"}), fields.PPR_ORDER)
    for existing in ppr.findall(qn("w:pStyle")):
        ppr.remove(existing)
    insert_in_order(ppr, make("w:pStyle", **{"w:val": "Heading1"}), fields.PPR_ORDER)


def _cover(context: BuildContext, block: S.CoverBlock):
    document = context.document
    accent = block.accent or context.theme.accent or "1F4E79"

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    if block.logo:
        stream, size = image_io.resolve_image(block.logo)
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(stream, width=Cm(4.5))
        align_paragraph(paragraph, "center")

    title = document.add_paragraph()
    apply_run_format(title.add_run(block.title), {"size_pt": 30, "bold": True, "color": accent,
                                                  "font": context.theme.heading_font})
    align_paragraph(title, "center")
    title.paragraph_format.space_after = Pt(6)

    if block.subtitle:
        subtitle = document.add_paragraph()
        apply_run_format(subtitle.add_run(block.subtitle), {"size_pt": 14, "color": "595959"})
        align_paragraph(subtitle, "center")

    rule = document.add_paragraph()
    fields.horizontal_rule(rule, color=accent, size=12)
    rule.paragraph_format.space_before = Pt(14)
    rule.paragraph_format.space_after = Pt(14)

    for value, size in ((block.author, 11), (block.date, 10)):
        if value:
            line = document.add_paragraph()
            apply_run_format(line.add_run(value), {"size_pt": size, "color": "595959"})
            align_paragraph(line, "center")

    if block.page_break_after:
        fields.add_page_break(document.add_paragraph())
    return title


def _caption(context: BuildContext, block: S.CaptionBlock):
    return refs.add_caption(
        context.document, block.text, label=block.label, bookmark=block.bookmark
    )


def _page_break(context: BuildContext, block: S.PageBreakBlock):
    return fields.add_page_break(context.document.add_paragraph())


def _rule(context: BuildContext, block: S.RuleBlock):
    # Parsed before the paragraph exists: parse_color raises on a bad value, and
    # adding first left an empty paragraph behind that the next successful call
    # would persist.
    color = parse_color(block.color) or "BFBFBF"
    paragraph = context.document.add_paragraph()
    fields.horizontal_rule(paragraph, color=color)
    return paragraph


def _spacer(context: BuildContext, block: S.SpacerBlock):
    paragraph = context.document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(float(block.height_pt))
    return paragraph


def _section(context: BuildContext, block: S.SectionBlock):
    section = layout.add_section(context.document, block.start)
    layout.apply_page_setup(
        section,
        size=block.size,
        orientation=block.orientation,
        margins_cm=block.margins_cm,
        columns=block.columns,
        column_spacing_cm=block.column_spacing_cm,
        column_line=block.column_line,
    )
    if block.restart_page_numbering_at is not None:
        layout.restart_page_numbering(section, block.restart_page_numbering_at)
    if block.header:
        apply_header(section, block.header)
    if block.footer:
        apply_footer(section, block.footer)
    return section


def _markdown(context: BuildContext, block: S.MarkdownBlock):
    from .mdconv import markdown_to_blocks

    return add_blocks(
        context, markdown_to_blocks(block.text, heading_offset=block.heading_offset)
    )


def _html(context: BuildContext, block: S.HtmlBlock):
    from .mdconv import html_to_blocks

    return add_blocks(context, html_to_blocks(block.html, heading_offset=block.heading_offset))


_HANDLERS = {
    S.HeadingBlock: _heading,
    S.ParagraphBlock: _paragraph,
    S.ListBlock: _list,
    S.TableBlock: _table,
    S.ImageBlock: _image,
    S.ChartBlock: _chart,
    S.MarkdownBlock: _markdown,
    S.HtmlBlock: _html,
    S.CodeBlock: _code,
    S.QuoteBlock: _quote,
    S.CalloutBlock: _callout,
    S.TocBlock: _toc,
    S.CoverBlock: _cover,
    S.CaptionBlock: _caption,
    S.PageBreakBlock: _page_break,
    S.RuleBlock: _rule,
    S.SpacerBlock: _spacer,
    S.SectionBlock: _section,
}
