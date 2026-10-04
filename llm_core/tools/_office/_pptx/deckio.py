# SPDX-License-Identifier: MIT
"""Opening, reading and editing existing .pptx files.

Reading matters more for decks than for documents: the common enterprise task is
not "make me a deck" but "take last quarter's deck and update it". So the reader
reports positions and shape identities, not just text — enough for a caller to
name the shape it wants to change.
"""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

from pptx import Presentation

from .assets import load_bytes
from .errors import InvalidSpec, SlideNotFound
from .ooxml.util import emu_to_cm, sanitize_text

# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


def open_presentation(source: str):
    blob = load_bytes(source, "presentation")
    if blob[:2] != b"PK":
        raise InvalidSpec(
            "That is not a .pptx file (no ZIP signature). The legacy .ppt format is not "
            "supported — save it as .pptx first."
        )
    try:
        return Presentation(io.BytesIO(blob))
    except Exception as exc:
        raise InvalidSpec(f"Could not open the presentation: {exc}") from exc


def slide_at(presentation, index: int):
    slides = presentation.slides
    if not 0 <= index < len(slides):
        raise SlideNotFound(index, len(slides))
    return slides[index]


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def iter_text_frames(slide, include_tables: bool = True) -> Iterable[Any]:
    """Every text frame on a slide, groups and table cells included."""
    for shape in _iter_shapes(slide.shapes):
        if shape.has_text_frame:
            yield shape.text_frame
        if include_tables and getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _iter_shapes(shapes) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # MSO_SHAPE_TYPE.GROUP
            yield from _iter_shapes(shape.shapes)


def slide_title(slide) -> str | None:
    """The slide's title: its title placeholder, or the topmost text if it has none."""
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type in (1, 3):  # TITLE, CENTER_TITLE
            text = (placeholder.text_frame.text or "").strip()
            if text:
                return text.split("\n")[0]
    best, best_top = None, None
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = (shape.text_frame.text or "").strip()
        if not text:
            continue
        top = int(shape.top or 0)
        if best_top is None or top < best_top:
            best, best_top = text.split("\n")[0], top
    return best


def shape_kind(shape) -> str:
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.shape_type is not None and str(shape.shape_type).startswith("PICTURE"):
        return "picture"
    if shape.is_placeholder:
        return "placeholder"
    if shape.has_text_frame and (shape.text_frame.text or "").strip():
        return "text"
    return "shape"


def describe_slide(presentation, index: int, detail: bool = False) -> dict[str, Any]:
    slide = slide_at(presentation, index)
    payload: dict[str, Any] = {
        "index": index,
        "slide_id": slide.slide_id,
        "layout": slide.slide_layout.name,
        "title": slide_title(slide),
        "shapes": len(slide.shapes),
        "has_notes": slide.has_notes_slide and bool(
            (slide.notes_slide.notes_text_frame.text or "").strip()
        ),
    }
    if not detail:
        return payload

    payload["notes"] = (
        slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    )
    inventory = []
    for position, shape in enumerate(slide.shapes):
        entry: dict[str, Any] = {
            "shape_index": position,
            "shape_id": shape.shape_id,
            "name": shape.name,
            "kind": shape_kind(shape),
            "x_cm": emu_to_cm(shape.left),
            "y_cm": emu_to_cm(shape.top),
            "w_cm": emu_to_cm(shape.width),
            "h_cm": emu_to_cm(shape.height),
        }
        if shape.is_placeholder:
            entry["placeholder_idx"] = shape.placeholder_format.idx
        if shape.has_text_frame:
            entry["text"] = shape.text_frame.text
        if getattr(shape, "has_table", False):
            entry["table"] = {
                "rows": len(shape.table.rows),
                "columns": len(shape.table.columns),
            }
        if getattr(shape, "has_chart", False):
            chart = shape.chart
            entry["chart"] = {
                "chart_type": str(chart.chart_type),
                "series": [series.name for series in chart.series],
            }
        inventory.append(entry)
    payload["inventory"] = inventory
    return payload


def outline(presentation) -> list[dict[str, Any]]:
    return [describe_slide(presentation, index) for index in range(len(presentation.slides))]


def read_text(presentation, include_notes: bool = True, include_tables: bool = True) -> str:
    """Flatten the whole deck to text, one block per slide."""
    blocks: list[str] = []
    for index, slide in enumerate(presentation.slides):
        lines = [f"--- Slide {index + 1}: {slide_title(slide) or '(untitled)'} ---"]
        title = slide_title(slide)
        for frame in iter_text_frames(slide, include_tables=include_tables):
            text = (frame.text or "").strip()
            if not text or text == title:
                continue
            lines.append(text)
        if include_notes and slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                lines.append(f"[notes] {notes}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def read_tables(presentation) -> list[dict[str, Any]]:
    result = []
    for index, slide in enumerate(presentation.slides):
        for position, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_table", False):
                continue
            table = shape.table
            result.append({
                "slide_index": index,
                "shape_index": position,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "data": [[cell.text for cell in row.cells] for row in table.rows],
            })
    return result


def describe(presentation) -> dict[str, Any]:
    properties = presentation.core_properties
    charts = sum(
        1 for slide in presentation.slides for shape in slide.shapes
        if getattr(shape, "has_chart", False)
    )
    tables = sum(
        1 for slide in presentation.slides for shape in slide.shapes
        if getattr(shape, "has_table", False)
    )
    return {
        "slides": len(presentation.slides),
        "layouts": sum(len(master.slide_layouts) for master in presentation.slide_masters),
        "masters": len(presentation.slide_masters),
        "charts": charts,
        "tables": tables,
        "slide_size_cm": {
            "width": emu_to_cm(presentation.slide_width),
            "height": emu_to_cm(presentation.slide_height),
        },
        "core_properties": {
            "title": properties.title or None,
            "author": properties.author or None,
            "subject": properties.subject or None,
            "keywords": properties.keywords or None,
            "created": properties.created.isoformat() if properties.created else None,
            "modified": properties.modified.isoformat() if properties.modified else None,
        },
    }


# ---------------------------------------------------------------------------
# Editing
# ---------------------------------------------------------------------------


def set_metadata(presentation, **values: Any) -> dict[str, Any]:
    properties = presentation.core_properties
    applied = {}
    for key, value in values.items():
        if value is None:
            continue
        value = sanitize_text(value)
        if not hasattr(properties, key):
            raise InvalidSpec(
                f"Unknown metadata field {key!r}. Valid fields: title, subject, author, "
                "keywords, category, comments."
            )
        setattr(properties, key, value)
        applied[key] = value
    return applied


def replace_text(
    presentation,
    search: str,
    replacement: str,
    *,
    regex: bool = False,
    ignore_case: bool = False,
    include_notes: bool = True,
    slide_indices: list[int] | None = None,
) -> int:
    """Search and replace across the deck, preserving run formatting.

    PowerPoint splits a sentence across runs at arbitrary points, so a match
    routinely spans several. Matches inside a single run are edited in place; a
    match that spans runs collapses onto the first run it touches, which is the
    only way to keep *some* formatting rather than none.
    """
    replacement = sanitize_text(replacement)
    flags = re.IGNORECASE if ignore_case else 0
    try:
        pattern = re.compile(search if regex else re.escape(search), flags)
    except re.error as exc:
        raise InvalidSpec(f"Invalid regular expression {search!r}: {exc}") from exc

    total = 0
    for index, slide in enumerate(presentation.slides):
        if slide_indices is not None and index not in slide_indices:
            continue
        for frame in iter_text_frames(slide):
            for paragraph in frame.paragraphs:
                total += _replace_in_paragraph(paragraph, pattern, replacement)
        if include_notes and slide.has_notes_slide:
            for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
                total += _replace_in_paragraph(paragraph, pattern, replacement)
    return total


def _replace_in_paragraph(paragraph, pattern: re.Pattern, replacement: str) -> int:
    runs = paragraph.runs
    if not runs:
        return 0
    full = "".join(run.text or "" for run in runs)
    if not pattern.search(full):
        return 0

    spans, cursor = [], 0
    for run in runs:
        length = len(run.text or "")
        spans.append((cursor, cursor + length))
        cursor += length

    replaced = 0
    for match in reversed(list(pattern.finditer(full))):
        start, end = match.span()
        touched = [i for i, (low, high) in enumerate(spans) if low < end and high > start]
        if not touched:
            continue
        first = touched[0]
        low, _ = spans[first]
        head = (runs[first].text or "")[: start - low]
        tail_run = touched[-1]
        tail_low, _ = spans[tail_run]
        tail = (runs[tail_run].text or "")[end - tail_low:]

        runs[first].text = head + match.expand(replacement) + (tail if tail_run == first else "")
        for position in touched[1:]:
            runs[position].text = "" if position != tail_run else tail
        replaced += 1

        cursor = 0
        for i, run in enumerate(runs):
            length = len(run.text or "")
            spans[i] = (cursor, cursor + length)
            cursor += length
    return replaced


def find_shape(slide, shape_index: int | None = None, name: str | None = None):
    shapes = list(slide.shapes)
    if shape_index is not None:
        if not 0 <= shape_index < len(shapes):
            raise InvalidSpec(
                f"shape_index {shape_index} is out of range — this slide has {len(shapes)} "
                "shapes. Read the slide with pptx_read(slide=N) for its shapes."
            )
        return shapes[shape_index]
    if name:
        for shape in shapes:
            if (shape.name or "").strip().lower() == str(name).strip().lower():
                return shape
        available = ", ".join(repr(shape.name) for shape in shapes)
        raise InvalidSpec(f"No shape named {name!r} on this slide. Shapes: {available or '(none)'}.")
    raise InvalidSpec("Give either shape_index or name.")
