# SPDX-License-Identifier: MIT
"""Native, editable PowerPoint charts — real chart objects, not pictures.

python-pptx writes the ``c:chartSpace`` part and the embedded worksheet for us,
which is what makes *Edit Data* work in PowerPoint. What it does not do is make
the result look considered: the default is a heavy grid, a 10-colour cycle that
fails colour-vision tests, and a legend on every chart including the ones with a
single series. Everything below the ``add_chart`` call is that second half.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Sequence

from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_DATA_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
    XL_TICK_MARK,
)

from ..errors import InvalidSpec
from .util import (
    Emu,
    Pt,
    drop,
    frag,
    insert_in_order,
    parse_color,
    qn,
    rgb,
)

C_NS = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------

#: Validated categorical palette, shared with the sibling docx server so a report
#: and its deck read as one system. Lightness band PASS, chroma floor PASS, worst
#: adjacent colour-vision-deficient ΔE 9.1, worst adjacent normal-vision ΔE 19.6.
#: Assigned in fixed series order — never cycled by rank.
DEFAULT_PALETTE: tuple[str, ...] = (
    "2A78D6",  # blue
    "EB6834",  # orange
    "1BAF7A",  # aqua
    "EDA100",  # yellow
    "E87BA4",  # magenta
    "008300",  # green
    "4A3AA7",  # violet
    "E34948",  # red
)

#: Recessive furniture — grid and axis lines must never compete with the marks.
GRIDLINE_COLOR = "E4E4E1"
AXIS_LINE_COLOR = "BFBFBF"
TEXT_COLOR = "404040"
DEFAULT_TEXT_PT = 11.0
DEFAULT_TITLE_PT = 13.0


# ---------------------------------------------------------------------------
# Chart types
# ---------------------------------------------------------------------------

NATIVE_TYPES: dict[str, XL_CHART_TYPE] = {
    # column / bar
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_clustered": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "column_stacked_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_clustered": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "bar_stacked_100": XL_CHART_TYPE.BAR_STACKED_100,
    # line
    "line": XL_CHART_TYPE.LINE,
    "line_markers": XL_CHART_TYPE.LINE_MARKERS,
    "line_stacked": XL_CHART_TYPE.LINE_STACKED,
    "line_stacked_100": XL_CHART_TYPE.LINE_STACKED_100,
    "line_markers_stacked": XL_CHART_TYPE.LINE_MARKERS_STACKED,
    "line_markers_stacked_100": XL_CHART_TYPE.LINE_MARKERS_STACKED_100,
    # circular
    "pie": XL_CHART_TYPE.PIE,
    "pie_exploded": XL_CHART_TYPE.PIE_EXPLODED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "donut": XL_CHART_TYPE.DOUGHNUT,
    "doughnut_exploded": XL_CHART_TYPE.DOUGHNUT_EXPLODED,
    # area
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "area_stacked_100": XL_CHART_TYPE.AREA_STACKED_100,
    # radar
    "radar": XL_CHART_TYPE.RADAR,
    "radar_markers": XL_CHART_TYPE.RADAR_MARKERS,
    "radar_filled": XL_CHART_TYPE.RADAR_FILLED,
    # xy
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "xy_scatter": XL_CHART_TYPE.XY_SCATTER,
    "scatter_lines": XL_CHART_TYPE.XY_SCATTER_LINES,
    "scatter_lines_no_markers": XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
    "scatter_smooth": XL_CHART_TYPE.XY_SCATTER_SMOOTH,
    "scatter_smooth_no_markers": XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
    "bubble": XL_CHART_TYPE.BUBBLE,
    "bubble_3d": XL_CHART_TYPE.BUBBLE_THREE_D_EFFECT,
}

#: CT_ChartSpace — c:spPr sits between the chart and the text properties.
CHART_SPACE_ORDER = (
    "date1904", "lang", "roundedCorners", "style", "clrMapOvr", "pivotSource", "protection",
    "chart", "spPr", "txPr", "externalData", "printSettings", "userShapes", "extLst",
)
#: CT_DoughnutChart
DOUGHNUT_ORDER = ("varyColors", "ser", "dLbls", "firstSliceAng", "holeSize", "extLst")

XY_TYPES = {k for k in NATIVE_TYPES if k.startswith(("scatter", "xy_"))}
BUBBLE_TYPES = {"bubble", "bubble_3d"}
CIRCULAR_TYPES = {"pie", "pie_exploded", "doughnut", "donut", "doughnut_exploded"}
#: Series colour lands on the stroke, not the fill.
LINE_TYPES = {k for k in NATIVE_TYPES if k.startswith("line")} | XY_TYPES | {"radar", "radar_markers"}
STACKED_TYPES = {k for k in NATIVE_TYPES if "stacked" in k}

_LEGEND_POSITIONS = {
    "right": XL_LEGEND_POSITION.RIGHT,
    "left": XL_LEGEND_POSITION.LEFT,
    "top": XL_LEGEND_POSITION.TOP,
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "top_right": XL_LEGEND_POSITION.CORNER,
    "corner": XL_LEGEND_POSITION.CORNER,
}

_LABEL_POSITIONS = {
    "center": XL_DATA_LABEL_POSITION.CENTER,
    "inside_end": XL_DATA_LABEL_POSITION.INSIDE_END,
    "inside_base": XL_DATA_LABEL_POSITION.INSIDE_BASE,
    "outside_end": XL_DATA_LABEL_POSITION.OUTSIDE_END,
    "above": XL_DATA_LABEL_POSITION.ABOVE,
    "below": XL_DATA_LABEL_POSITION.BELOW,
    "best_fit": XL_DATA_LABEL_POSITION.BEST_FIT,
}


def native_chart_types() -> list[str]:
    return sorted(NATIVE_TYPES)


def normalise_type(value: str) -> str:
    return str(value or "").strip().lower().replace("-", "_").replace(" ", "_")


# ---------------------------------------------------------------------------
# Request
# ---------------------------------------------------------------------------


@dataclass
class SeriesData:
    name: str
    values: Sequence[float | None] | None = None
    points: Sequence[Sequence[float]] | None = None
    color: str | None = None


@dataclass
class ChartRequest:
    chart_type: str = "column_clustered"
    categories: Sequence[Any] = field(default_factory=list)
    series: Sequence[SeriesData] = field(default_factory=list)
    title: str | None = None
    x_title: str | None = None
    y_title: str | None = None
    legend: str = "auto"
    data_labels: bool | None = None
    data_label_position: str | None = None
    number_format: str | None = None
    data_label_format: str | None = None
    gridlines: bool = True
    y_min: float | None = None
    y_max: float | None = None
    palette: Sequence[str] | None = None
    gap_width: int | None = None
    overlap: int | None = None
    hole_size: int | None = None
    smooth: bool | None = None
    font: str | None = None
    text_color: str = TEXT_COLOR
    font_size_pt: float = DEFAULT_TEXT_PT


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------


def build_chart_data(request: ChartRequest):
    key = normalise_type(request.chart_type)
    if not request.series:
        raise InvalidSpec("A chart needs at least one series.")

    if key in BUBBLE_TYPES:
        data = BubbleChartData(number_format=request.number_format or "General")
        for series in request.series:
            if not series.points:
                raise InvalidSpec(
                    f"Bubble series {series.name!r} needs 'points' as [x, y, size] triples."
                )
            handle = data.add_series(series.name)
            for point in series.points:
                if len(point) < 3:
                    raise InvalidSpec(
                        f"Bubble series {series.name!r} has a point with {len(point)} value(s); "
                        "each point must be [x, y, size]."
                    )
                handle.add_data_point(float(point[0]), float(point[1]), float(point[2]))
        return data

    if key in XY_TYPES:
        data = XyChartData(number_format=request.number_format or "General")
        for series in request.series:
            if not series.points:
                raise InvalidSpec(
                    f"Scatter series {series.name!r} needs 'points' as [x, y] pairs."
                )
            handle = data.add_series(series.name)
            for point in series.points:
                if len(point) < 2:
                    raise InvalidSpec(
                        f"Scatter series {series.name!r} has a point with one value; "
                        "each point must be [x, y]."
                    )
                handle.add_data_point(float(point[0]), float(point[1]))
        return data

    data = CategoryChartData(number_format=request.number_format or "General")
    categories = list(request.categories)
    lengths = {len(series.values or []) for series in request.series}
    if not categories:
        width = max(lengths) if lengths else 0
        categories = [f"Item {i + 1}" for i in range(width)]
    data.categories = categories
    for series in request.series:
        values = list(series.values or [])
        if not values:
            raise InvalidSpec(
                f"Series {series.name!r} has no 'values'. Category charts need one value "
                "per category; scatter and bubble charts use 'points' instead."
            )
        if len(values) != len(categories):
            raise InvalidSpec(
                f"Series {series.name!r} has {len(values)} value(s) but there are "
                f"{len(categories)} categories. They must line up."
            )
        data.add_series(series.name, values, request.number_format or None)
    return data


# ---------------------------------------------------------------------------
# Placement
# ---------------------------------------------------------------------------


def add_chart(slide, request: ChartRequest, left: int, top: int, width: int, height: int):
    """Insert a native chart and style it. Returns the graphic frame."""
    key = normalise_type(request.chart_type)
    chart_type = NATIVE_TYPES.get(key)
    if chart_type is None:
        raise InvalidSpec(
            f"{request.chart_type!r} has no native PowerPoint equivalent. "
            f"Native types: {', '.join(native_chart_types())}."
        )
    data = build_chart_data(request)
    frame = slide.shapes.add_chart(chart_type, Emu(left), Emu(top), Emu(width), Emu(height), data)
    style_chart(frame.chart, request)
    return frame


# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------


def style_chart(chart, request: ChartRequest) -> None:
    key = normalise_type(request.chart_type)
    palette = [parse_color(c) for c in (request.palette or DEFAULT_PALETTE)]
    text_color = parse_color(request.text_color) or TEXT_COLOR

    _base_text(chart, request, text_color)
    _title(chart, request, text_color)
    _series_colors(chart, key, palette)
    _series_overrides(chart, key, request)
    _axes(chart, key, request, text_color)
    _legend(chart, key, request, text_color)
    _plot_options(chart, key, request)
    _data_labels(chart, key, request, text_color)
    _no_border(chart)


def _base_text(chart, request: ChartRequest, text_color: str) -> None:
    font = chart.font
    font.size = Pt(request.font_size_pt)
    font.color.rgb = rgb(text_color)
    if request.font:
        font.name = request.font


def _title(chart, request: ChartRequest, text_color: str) -> None:
    if not request.title:
        chart.has_title = False
        return
    chart.has_title = True
    frame = chart.chart_title.text_frame
    frame.text = request.title
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(DEFAULT_TITLE_PT)
        paragraph.font.bold = True
        paragraph.font.color.rgb = rgb(text_color)
        if request.font:
            paragraph.font.name = request.font


def _series_colors(chart, key: str, palette: Sequence[str]) -> None:
    """Colour in series order — never by rank, never cycled by value."""
    for index, series in enumerate(chart.series):
        color = palette[index % len(palette)]
        if key in CIRCULAR_TYPES:
            # One series, many slices: colour the points instead.
            for point_index, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(palette[point_index % len(palette)])
                point.format.line.color.rgb = rgb("FFFFFF")
                point.format.line.width = Pt(1.5)
            continue
        if key in LINE_TYPES:
            series.format.line.color.rgb = rgb(color)
            series.format.line.width = Pt(2.25)
            if key in XY_TYPES and "lines" not in key and "smooth" not in key:
                # A pure scatter has no line; python-pptx would draw one anyway.
                series.format.line.fill.background()
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(color)
            series.format.line.fill.background()


def _series_overrides(chart, key: str, request: ChartRequest) -> None:
    """Honour per-series ``color`` overrides after the palette pass."""
    for index, (series, spec) in enumerate(zip(chart.series, request.series)):
        if not spec.color:
            continue
        color = parse_color(spec.color)
        if key in CIRCULAR_TYPES:
            points = list(series.points)
            if index < len(points):
                points[index].format.fill.solid()
                points[index].format.fill.fore_color.rgb = rgb(color)
            continue
        if key in LINE_TYPES:
            series.format.line.color.rgb = rgb(color)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(color)


def _axes(chart, key: str, request: ChartRequest, text_color: str) -> None:
    if key in CIRCULAR_TYPES:
        return
    try:
        category_axis = chart.category_axis
        value_axis = chart.value_axis
    except (ValueError, NotImplementedError):  # pragma: no cover — radar/odd types
        return

    for axis in (category_axis, value_axis):
        axis.has_minor_gridlines = False
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.tick_labels.font.size = Pt(request.font_size_pt)
        axis.tick_labels.font.color.rgb = rgb(text_color)
        if request.font:
            axis.tick_labels.font.name = request.font

    # Only the value axis carries gridlines: horizontal rules a reader can follow.
    category_axis.has_major_gridlines = False
    value_axis.has_major_gridlines = bool(request.gridlines)
    if request.gridlines:
        line = value_axis.major_gridlines.format.line
        line.color.rgb = rgb(GRIDLINE_COLOR)
        line.width = Pt(0.75)

    category_axis.format.line.color.rgb = rgb(AXIS_LINE_COLOR)
    category_axis.format.line.width = Pt(0.75)
    # The value-axis spine duplicates the gridlines; drop it.
    value_axis.format.line.fill.background()
    value_axis.tick_label_position = XL_TICK_LABEL_POSITION.NEXT_TO_AXIS

    if request.number_format:
        value_axis.tick_labels.number_format = request.number_format
        value_axis.tick_labels.number_format_is_linked = False
    if request.y_min is not None:
        value_axis.minimum_scale = float(request.y_min)
    if request.y_max is not None:
        value_axis.maximum_scale = float(request.y_max)

    _axis_title(category_axis, request.x_title, request, text_color)
    _axis_title(value_axis, request.y_title, request, text_color)


def _axis_title(axis, text: str | None, request: ChartRequest, text_color: str) -> None:
    if not text:
        axis.has_title = False
        return
    axis.has_title = True
    frame = axis.axis_title.text_frame
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(max(9.0, request.font_size_pt - 0.5))
        paragraph.font.bold = False
        paragraph.font.color.rgb = rgb(text_color)
        if request.font:
            paragraph.font.name = request.font


def _legend(chart, key: str, request: ChartRequest, text_color: str) -> None:
    """One rule: a legend only exists when the reader would otherwise be lost.

    A single series is named by the title. A pie is labelled on its slices. Both
    get no legend box, and the space goes to the marks instead.
    """
    mode = str(request.legend or "auto").lower()
    if mode == "none":
        chart.has_legend = False
        return
    if mode == "auto":
        if key in CIRCULAR_TYPES or len(request.series) <= 1:
            chart.has_legend = False
            return
        position = XL_LEGEND_POSITION.BOTTOM
    else:
        position = _LEGEND_POSITIONS.get(mode)
        if position is None:
            raise InvalidSpec(
                f"Unknown legend position {request.legend!r}. Use auto, none, right, left, "
                "top, bottom or top_right."
            )
    chart.has_legend = True
    chart.legend.position = position
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(request.font_size_pt)
    chart.legend.font.color.rgb = rgb(text_color)
    if request.font:
        chart.legend.font.name = request.font


def _plot_options(chart, key: str, request: ChartRequest) -> None:
    for plot in chart.plots:
        if hasattr(plot, "gap_width"):
            plot.gap_width = int(
                request.gap_width if request.gap_width is not None
                else (60 if key in STACKED_TYPES else 80)
            )
        if hasattr(plot, "overlap"):
            plot.overlap = int(
                request.overlap if request.overlap is not None
                else (100 if key in STACKED_TYPES else -10)
            )
        if key in CIRCULAR_TYPES:
            plot.vary_by_categories = True

    if key in {"doughnut", "donut", "doughnut_exploded"}:
        _hole_size(chart, request.hole_size if request.hole_size is not None else 62)
    if request.smooth is not None:
        for series in chart.series:
            try:
                series.smooth = bool(request.smooth)
            except (AttributeError, NotImplementedError):  # pragma: no cover
                pass


def _hole_size(chart, percent: int) -> None:
    for element in chart._chartSpace.iter(qn("c:doughnutChart")):
        drop(element, "c:holeSize")
        insert_in_order(
            element,
            frag(f'<c:holeSize {C_NS} val="{max(10, min(90, int(percent)))}"/>'),
            DOUGHNUT_ORDER,
        )


DLBLS_ORDER = (
    "dLbl", "delete", "numFmt", "spPr", "txPr", "dLblPos", "showLegendKey", "showVal",
    "showCatName", "showSerName", "showPercent", "showBubbleSize", "separator",
    "showLeaderLines", "leaderLines",
)


def set_label_flags(plot, *, value=None, category=None, percent=None, series=None) -> None:
    """Say *what* a data label shows.

    ``has_data_labels = True`` only creates the ``c:dLbls`` element — python-pptx
    writes every ``show*`` flag as 0, so a chart with labels turned on shows
    nothing at all until these are set.
    """
    dlbls = plot._element.get_or_add_dLbls()
    for tag, wanted in (
        ("c:showVal", value), ("c:showCatName", category),
        ("c:showPercent", percent), ("c:showSerName", series),
    ):
        if wanted is None:
            continue
        drop(dlbls, tag)
        insert_in_order(dlbls, frag(f'<{tag} {C_NS} val="{1 if wanted else 0}"/>'), DLBLS_ORDER)


def _data_labels(chart, key: str, request: ChartRequest, text_color: str) -> None:
    wanted = request.data_labels
    if wanted is None:
        # Circular charts are unreadable without them; everything else is cleaner without.
        wanted = key in CIRCULAR_TYPES
    circular = key in CIRCULAR_TYPES
    for plot in chart.plots:
        # python-pptx models no c:dLbls on a scatter or bubble plot, so neither
        # adding nor removing one is possible there.
        if not hasattr(plot._element, "get_or_add_dLbls"):
            if request.data_labels:
                raise InvalidSpec(
                    f"Data labels are not available on a {key!r} chart. Label the points in "
                    "the series names, or use a category chart type."
                )
            continue
        plot.has_data_labels = bool(wanted)
        if not wanted:
            continue
        labels = plot.data_labels
        labels.font.size = Pt(max(9.0, request.font_size_pt - 1))
        # A doughnut has nowhere to put a label but inside its own ring, so the
        # label has to read on the series colour; a pie can push them outside.
        inside_ring = key in {"doughnut", "donut", "doughnut_exploded"}
        labels.font.color.rgb = rgb("FFFFFF" if inside_ring else text_color)
        if request.font:
            labels.font.name = request.font
        fmt = request.data_label_format or request.number_format
        if fmt:
            labels.number_format = fmt
            labels.number_format_is_linked = False

        if circular:
            # A pie with no legend needs its slices named, and a share reads
            # better as a percentage than as the raw number.
            set_label_flags(plot, value=False, category=True, percent=True)
        else:
            set_label_flags(plot, value=True, category=False, percent=False)

        position = request.data_label_position
        if position:
            resolved = _LABEL_POSITIONS.get(str(position).lower())
            if resolved is None:
                raise InvalidSpec(
                    f"Unknown data label position {position!r}. Use one of: "
                    f"{', '.join(sorted(_LABEL_POSITIONS))}."
                )
            labels.position = resolved
        elif key in {"pie", "pie_exploded"}:
            labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
        elif key in STACKED_TYPES:
            labels.position = XL_DATA_LABEL_POSITION.CENTER
        # A doughnut accepts no dLblPos at all — PowerPoint refuses the file.


def _no_border(chart) -> None:
    """Charts sit on the slide, not in a box."""
    chart_space = chart._chartSpace
    drop(chart_space, "c:spPr")
    insert_in_order(
        chart_space,
        frag(f"<c:spPr {C_NS} {A_NS}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>"),
        CHART_SPACE_ORDER,
    )
