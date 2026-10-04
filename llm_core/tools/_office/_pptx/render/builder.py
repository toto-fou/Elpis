# SPDX-License-Identifier: MIT
"""The engine: turn a :class:`DeckSpec` into a python-pptx ``Presentation``.

One slide layout, one function. Each is handed a :class:`RenderContext` and asks
the geometry engine for its bands, so ``pptx_build`` and the granular
``pptx_add_slide`` tools go through exactly the same code path — a slide added
later is indistinguishable from one built in the first call.
"""

from __future__ import annotations

import io
from typing import Any, Iterable, Sequence

from pptx import Presentation
from pptx.util import Emu, Pt

from .. import templates
from ..errors import InvalidSpec
from ..ooxml import shapes as shape_ooxml, slides as slide_ooxml, text as text_ooxml, theme as theme_ooxml
from ..ooxml.util import anchor, mix, parse_color, rgb
from . import elements as E, spec as S
from .layout import Frame, Geometry, cm, fit_font_size, text_height_cm
from .theme import Theme, resolve_theme

#: Slide sizes, in centimetres.
SIZE_PRESETS = {
    "16:9": (33.867, 19.05),
    "16:10": (33.867, 21.167),
    "4:3": (25.4, 19.05),
    "a4": (21.0, 29.7),
    "a4_landscape": (29.7, 21.0),
    "letter": (27.94, 21.59),
}

#: Layouts that carry no footer furniture unless the caller insists.
CHROMELESS = {"title", "cover", "section", "divider", "closing", "end", "thanks"}


# ---------------------------------------------------------------------------
# Deck assembly
# ---------------------------------------------------------------------------


def build_deck(spec: S.DeckSpec) -> tuple[Any, E.RenderContext]:
    presentation, template_notes = _open_base(spec.template)
    _apply_size(presentation, spec.size, has_template=bool(spec.template))
    if spec.template:
        template_notes += _strip_template_slides(presentation, spec.template_slides)

    geometry = Geometry.from_presentation(presentation)
    theme_spec = _theme_for(spec, presentation)
    theme = resolve_theme(theme_spec, slide_height_cm=geometry.height_cm)
    ctx = E.RenderContext(
        presentation=presentation, theme=theme, geometry=geometry, footer=spec.footer
    )
    ctx.notes.extend(template_notes)

    _apply_metadata(presentation, spec.meta)
    if _should_write_master(spec):
        _apply_theme_to_master(presentation, theme, bool(spec.template))
    if spec.footer:
        slide_ooxml.enable_header_footer(
            presentation, slide_number=spec.footer.slide_numbers, footer=bool(spec.footer.text)
        )

    add_slides(ctx, spec.slides)
    return presentation, ctx


def _open_base(template: str | None) -> tuple[Any, list[str]]:
    """Open the deck a build starts from: blank, or a corporate template.

    ``template`` may be a catalogue name — the normal case — or a path or base64
    payload. Template *formats* (.potx, .ppsx, macro-enabled) are normalised on
    the way in, because refusing the exact file an IT department distributes is
    not a defensible position for a server that claims to support templates.
    """
    if not template:
        return Presentation(), []

    blob, notes, catalogue_name = templates.load(template)
    try:
        presentation = Presentation(io.BytesIO(blob))
    except Exception as exc:
        raise InvalidSpec(f"Could not open the template: {exc}") from exc
    if catalogue_name:
        notes = [f"Built on the {catalogue_name!r} template from the catalogue.", *notes]
    return presentation, notes


def _strip_template_slides(presentation, mode: str) -> list[str]:
    """Drop the slides a corporate template ships with.

    A house template almost always carries example or instruction slides. They
    are there to be read by whoever builds a deck by hand, and they have no place
    in the deck the server produces — so they go, unless asked to stay.
    """
    count = len(presentation.slides)
    if not count or mode == "keep":
        return []
    for index in reversed(range(count)):
        slide_ooxml.delete_slide(presentation, index)
    return [
        f"Removed {count} slide(s) the template shipped with. Pass "
        'template_slides="keep" to build on top of them instead.'
    ]


def _theme_for(spec: S.DeckSpec, presentation) -> S.ThemeSpec:
    """Take the design tokens from the template's own theme, unless told otherwise.

    A corporate template *is* the brand definition. Building on it and then
    painting the slides in this server's default blue is the single most likely
    way to produce something that looks conformant and is not — so anything the
    caller did not set explicitly comes from the template.
    """
    theme_spec = spec.theme
    wanted = theme_spec.source or ("template" if spec.template else "preset")
    if not spec.template or wanted != "template":
        return theme_spec

    brand = theme_ooxml.read_theme(presentation) or {}
    colors = brand.get("colors") or {}
    explicit = theme_spec.model_fields_set
    updates: dict[str, Any] = {}

    def take(field: str, value: str | None, *, hex_color: bool = True) -> None:
        if field in explicit or not value:
            return
        updates[field] = f"#{value}" if hex_color else value

    take("accent", colors.get("accent1"))
    take("accent2", colors.get("accent2"))
    take("text_color", colors.get("dk1") or colors.get("tx1"))
    take("background", colors.get("lt1") or colors.get("bg1"))
    take("heading_font", brand.get("major_font"), hex_color=False)
    take("body_font", brand.get("minor_font"), hex_color=False)

    return theme_spec.model_copy(update=updates) if updates else theme_spec


def _should_write_master(spec: S.DeckSpec) -> bool:
    """Whether to push our palette into the theme part.

    Explicit wins. Otherwise: never for a template — overwriting the accents of
    the brand you were handed is exactly backwards — and always for a deck built
    from scratch, so shapes a human adds later inherit the design.
    """
    if "write_to_master" in spec.theme.model_fields_set:
        return spec.theme.write_to_master
    return not spec.template


def _apply_size(presentation, size, has_template: bool = False) -> None:
    if isinstance(size, str):
        size = S.SizeSpec(preset=size)
    if size is None:
        size = S.SizeSpec()
    if size.width_cm and size.height_cm:
        slide_ooxml.set_slide_size(presentation, cm(size.width_cm), cm(size.height_cm))
        return
    if size.preset is None:
        return
    # A template's own page size is a deliberate corporate choice; only override it
    # when the caller asked for something other than the default.
    if has_template and size.preset == "16:9" and presentation.slide_width:
        return
    width, height = SIZE_PRESETS[size.preset]
    slide_ooxml.set_slide_size(presentation, cm(width), cm(height))


def _apply_metadata(presentation, meta: S.MetaSpec) -> None:
    properties = presentation.core_properties
    for attribute in ("title", "subject", "author", "keywords", "category", "comments"):
        value = getattr(meta, attribute, None)
        if value:
            setattr(properties, attribute, value)


def _apply_theme_to_master(presentation, theme: Theme, has_template: bool) -> None:
    """Write the brand into the theme part, so later manual edits inherit it."""
    palette = list(theme.palette())
    theme_ooxml.apply_colors(
        presentation,
        accents=palette[:6],
        dark=theme.text,
        light=theme.surface,
        hyperlink=theme.accent2,
        followed_hyperlink=mix(theme.accent2, theme.text, 0.4),
        scheme_name=f"pptx-mcp {theme.preset}",
    )
    if not has_template:
        # A corporate template's typography is its own; only a from-scratch deck
        # gets its fonts rewritten.
        theme_ooxml.apply_fonts(presentation, major=theme.heading_font, minor=theme.body_font)


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------


def add_slides(ctx: E.RenderContext, slides: Iterable[S.Slide]) -> int:
    count = 0
    for slide_spec in slides:
        add_slide(ctx, slide_spec)
        count += 1
    return count


def add_slide(ctx: E.RenderContext, spec: S.Slide):
    layout = _resolve_layout(ctx, spec)
    slide = ctx.presentation.slides.add_slide(layout)
    if spec.layout_name is None:
        # A blank layout may still ship placeholders; the builders draw their own.
        for shape in list(slide.shapes):
            if shape.is_placeholder:
                shape._element.getparent().remove(shape._element)

    _apply_background(ctx, slide, spec)

    handler = _HANDLERS.get(type(spec))
    if handler is None:  # pragma: no cover — the discriminated union prevents this
        raise InvalidSpec(f"No builder for slide layout {type(spec).__name__}.")
    # A builder reports the area it left free, so extra elements — and the grid
    # they may be placed on — land under the title rather than across it.
    free = handler(ctx, slide, spec) or ctx.geometry.content()

    for element in spec.elements:
        E.render_element(ctx, slide, element, free)

    _apply_footer(ctx, slide, spec)

    if spec.notes:
        slide_ooxml.set_notes(slide, spec.notes)
    if spec.transition:
        slide_ooxml.set_transition(slide, spec.transition)
    if spec.name:
        slide_ooxml.set_name(slide, spec.name)
    return slide


def _resolve_layout(ctx: E.RenderContext, spec: S.Slide):
    if spec.layout_name is not None:
        try:
            return slide_ooxml.find_layout(ctx.presentation, spec.layout_name)
        except KeyError as exc:
            raise InvalidSpec(str(exc)) from exc
    if isinstance(spec, S.TemplateSlide):
        raise InvalidSpec(
            "A 'template' slide needs layout_name — the layout of your corporate template "
            "to build on (the layout names are shown by pptx_read on a deck made from it)."
        )
    return slide_ooxml.blank_layout(ctx.presentation)


def _apply_background(ctx: E.RenderContext, slide, spec: S.SlideBase) -> None:
    theme = ctx.theme
    if spec.background_image:
        picture = E.place_image(
            ctx, slide, ctx.geometry.full(), spec.background_image, fit="cover",
        )
        shape_ooxml.send_to_back(picture)
        if spec.background_image_dim:
            veil = E.add_card(ctx, slide, ctx.geometry.full(), radius=0, shape="rect")
            shape_ooxml.solid_fill(veil, "000000", alpha=spec.background_image_dim)
        return
    if spec.background_gradient:
        slide_ooxml.set_background(slide, gradient=[parse_color(c) for c in spec.background_gradient])
        return
    if spec.background:
        slide_ooxml.set_background(slide, color=spec.background)
        return
    if theme.background.upper() != "FFFFFF":
        slide_ooxml.set_background(slide, color=theme.background)


# ---------------------------------------------------------------------------
# Shared slide furniture
# ---------------------------------------------------------------------------


def _header(ctx: E.RenderContext, slide, title: str | None, kicker: str | None = None,
            accent_rule: bool = True) -> Frame:
    """Draw the title band and return the content area left underneath."""
    geometry = ctx.geometry
    theme = ctx.theme
    has_kicker = bool(kicker)
    if has_kicker:
        E.add_kicker(ctx, slide, geometry.kicker(), kicker)
    if not title:
        base = geometry.content(with_title=False)
        if not has_kicker:
            return base
        top = geometry.kicker().bottom + geometry.gap
        return Frame(base.x, top, base.w, max(0, base.bottom - top))

    band = geometry.title(with_kicker=has_kicker)
    E.add_title(ctx, slide, band, title, valign="bottom")
    if accent_rule:
        E.add_rule(
            ctx, slide,
            Frame(band.x, band.bottom + int(geometry.title_gap * 0.34), cm(2.6), 0),
            theme.accent, thickness_cm=0.09,
        )
    return geometry.content(with_title=True, with_kicker=has_kicker).pad(top=int(geometry.gap * 0.5))


def _fit_or_warn(ctx: E.RenderContext, items: Sequence[tuple[str, int]], frame: Frame,
                 base_pt: float, min_pt: float, what: str) -> float:
    """Choose a font size, and say so in the notes when the content simply does not fit."""
    size = fit_font_size(items, frame, base_pt=base_pt, min_pt=min_pt, line_spacing=1.16,
                         space_before_pt=base_pt * 0.42)
    if size <= min_pt:
        needed = text_height_cm(items, max(1.0, frame.w_cm - 0.5), min_pt,
                                line_spacing=1.16, space_before_pt=min_pt * 0.42)
        if needed > frame.h_cm:
            ctx.note(
                f"{what} holds more text than its box: it would need {needed:.1f} cm at the "
                f"smallest readable size and has {frame.h_cm:.1f} cm. Split it across two "
                "slides, or cut words."
            )
    return size


def _takeaway_band(ctx: E.RenderContext, content: Frame, text: str | None) -> tuple[Frame, Frame | None]:
    """Reserve the bottom strip for a takeaway. Returns (body, takeaway frame)."""
    if not text:
        return content, None
    height = max(cm(1.15), min(cm(1.9), int(content.h * 0.19)))
    band = content.bottom_slice(height)
    return content.pad(bottom=height + ctx.gap), band


def _apply_footer(ctx: E.RenderContext, slide, spec: S.SlideBase) -> None:
    footer = ctx.footer
    if footer is None or spec.hide_footer:
        return
    layout_name = getattr(spec, "layout", "")
    if layout_name in CHROMELESS and not footer.show_on_title:
        return
    if _has_own_furniture(slide):
        # The template's own footer placeholders are already on this slide;
        # drawing ours on top would print the company name twice.
        return

    theme = ctx.theme
    geometry = ctx.geometry
    band = geometry.footer()
    color = footer.color or theme.muted
    size = theme.pt("footer")

    if footer.rule:
        E.add_rule(ctx, slide, Frame(band.x, band.y - int(geometry.gap * 0.4), band.w, 0),
                   theme.border, thickness_cm=0.03)

    logo_w = 0
    if footer.logo:
        logo_h = int(band.h * 0.92)
        logo_frame = Frame(band.right - cm(3.0), band.y, cm(3.0), logo_h)
        picture = E.place_image(ctx, slide, logo_frame, footer.logo, fit="contain",
                                align="right", valign="middle")
        logo_w = int(picture.width) + geometry.gap

    number_w = cm(1.6) if footer.slide_numbers else 0
    if footer.slide_numbers:
        _slide_number(ctx, slide,
                      Frame(band.right - logo_w - number_w, band.y, number_w, band.h),
                      color, size)

    if footer.text:
        E.add_plain_text(
            ctx, slide, Frame(band.x, band.y, int(band.w * 0.5), band.h), footer.text,
            size_pt=size, color=color, valign="middle", align="left", autofit=False,
        )
    if footer.confidentiality:
        E.add_plain_text(
            ctx, slide, Frame(band.x + int(band.w * 0.25), band.y, int(band.w * 0.5), band.h),
            footer.confidentiality,
            size_pt=size, color=color, valign="middle", align="center", autofit=False,
        )


def _has_own_furniture(slide) -> bool:
    try:
        placeholders = list(slide.placeholders)
    except (AttributeError, KeyError):  # pragma: no cover — defensive
        return False
    return any(int(p.placeholder_format.type) in {13, 15} for p in placeholders)


def _slide_number(ctx: E.RenderContext, slide, frame: Frame, color: str, size_pt: float):
    box = slide.shapes.add_textbox(Emu(frame.x), Emu(frame.y), Emu(frame.w), Emu(frame.h))
    text_frame = box.text_frame
    text_frame.word_wrap = False
    text_frame.vertical_anchor = anchor("middle")
    text_ooxml.set_insets(text_frame, 0, 0, 0, 0)
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = E.alignment("right")
    field = text_ooxml.add_field(paragraph, "slidenum")
    text_ooxml.field_run_properties(
        field, size_pt=size_pt, color=color, font=ctx.theme.body_font
    )
    return box


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------


def _title_slide(ctx: E.RenderContext, slide, spec: S.TitleSlide) -> Frame | None:
    theme = ctx.theme
    geometry = ctx.geometry
    dark_bg = bool(spec.background or spec.background_gradient or spec.background_image)
    on_dark = bool(spec.background_image) or (
        spec.background is not None and theme.on(parse_color(spec.background)) == "FFFFFF"
    )
    text_color = "FFFFFF" if on_dark else theme.text
    muted_color = "E8ECF2" if on_dark else theme.muted

    if spec.variant == "band":
        panel = Frame(0, 0, int(geometry.width * 0.38), geometry.height)
        E.add_card(ctx, slide, panel, fill=theme.accent, radius=0, shape="rect")
        if spec.logo:
            E.place_image(ctx, slide, panel.pad(left=cm(1.6), right=cm(1.6),
                                                top=cm(1.4), bottom=int(geometry.height * 0.78)),
                          spec.logo, fit="contain", align="left")
        body = Frame(panel.right + geometry.margin_x, 0,
                     geometry.width - panel.right - 2 * geometry.margin_x, geometry.height)
        align = "left"
    elif spec.variant == "centered":
        body = geometry.full().pad(left=int(geometry.width * 0.12), right=int(geometry.width * 0.12))
        align = "center"
    else:
        body = geometry.full().pad(left=geometry.margin_x, right=int(geometry.width * 0.16))
        align = "left"

    block_top = int(geometry.height * (0.36 if spec.variant != "centered" else 0.30))
    if spec.variant != "band" and spec.logo:
        E.place_image(
            ctx, slide, Frame(body.x, int(geometry.height * 0.13), cm(4.2), cm(1.7)),
            spec.logo, fit="contain", align="left" if align == "left" else "center",
        )

    rule_y = block_top - int(geometry.height * 0.045)
    if align == "left":
        E.add_rule(ctx, slide, Frame(body.x, rule_y, cm(3.4), 0), theme.accent2 if dark_bg
                   else theme.accent, thickness_cm=0.11)

    title_h = int(geometry.height * 0.17)
    E.add_plain_text(
        ctx, slide, Frame(body.x, block_top, body.w, title_h), spec.title,
        base_size_pt=theme.pt("deck_title"),
        min_size_pt=theme.pt("deck_title") * 0.5,
        color=text_color, font=theme.heading_font, bold=True,
        align=align, valign="top", line_spacing=1.06,
    )
    cursor = block_top + title_h + int(geometry.height * 0.02)
    if spec.subtitle:
        E.add_plain_text(
            ctx, slide, Frame(body.x, cursor, body.w, int(geometry.height * 0.11)), spec.subtitle,
            size_pt=theme.pt("deck_subtitle"), color=muted_color, align=align, valign="top",
            line_spacing=1.15,
        )

    meta = " · ".join(part for part in (spec.author, spec.date) if part)
    if meta:
        E.add_plain_text(
            ctx, slide,
            Frame(body.x, geometry.height - geometry.margin_bottom - cm(1.4), body.w, cm(1.0)),
            meta,
            size_pt=theme.pt("small"), color=muted_color, align=align, valign="bottom",
            autofit=False,
        )


def _section_slide(ctx: E.RenderContext, slide, spec: S.SectionSlide) -> Frame | None:
    theme = ctx.theme
    geometry = ctx.geometry

    if spec.variant == "full":
        if not (spec.background or spec.background_gradient or spec.background_image):
            slide_ooxml.set_background(slide, color=theme.accent)
        text_color = theme.on_accent
        number_color = mix(theme.on_accent, theme.accent, 0.62)
        body = geometry.full().pad(left=geometry.margin_x, right=geometry.margin_x)
    elif spec.variant == "left":
        panel = Frame(0, 0, int(geometry.width * 0.42), geometry.height)
        E.add_card(ctx, slide, panel, fill=theme.accent, radius=0, shape="rect")
        text_color = theme.text
        number_color = mix(theme.accent, theme.background, 0.72)
        body = Frame(panel.right + geometry.margin_x, 0,
                     geometry.width - panel.right - 2 * geometry.margin_x, geometry.height)
    else:  # band
        band_h = int(geometry.height * 0.42)
        band = Frame(0, int(geometry.height * 0.29), geometry.width, band_h)
        E.add_card(ctx, slide, band, fill=theme.accent, radius=0, shape="rect")
        text_color = theme.on_accent
        number_color = mix(theme.on_accent, theme.accent, 0.62)
        body = band.pad(left=geometry.margin_x, right=geometry.margin_x)

    if spec.number:
        E.add_plain_text(
            ctx, slide,
            Frame(body.x, body.y + int(body.h * 0.16), body.w, int(geometry.height * 0.13)),
            spec.number,
            size_pt=theme.pt("section_title") * 0.62, color=number_color, bold=True,
            font=theme.heading_font, valign="bottom", autofit=False,
        )

    title_top = body.y + int(body.h * (0.34 if spec.number else 0.26))
    E.add_plain_text(
        ctx, slide, Frame(body.x, title_top, body.w, int(geometry.height * 0.22)), spec.title,
        size_pt=theme.pt("section_title"),
        min_size_pt=theme.pt("section_title") * 0.5,
        color=text_color, font=theme.heading_font, bold=True, valign="top", line_spacing=1.05,
    )
    if spec.subtitle:
        E.add_plain_text(
            ctx, slide,
            Frame(body.x, title_top + int(geometry.height * 0.22), body.w, int(geometry.height * 0.12)),
            spec.subtitle,
            base_size_pt=theme.pt("body"), color=mix(text_color, theme.accent, 0.28), valign="top",
            line_spacing=1.15,
        )


def _bullets_slide(ctx: E.RenderContext, slide, spec: S.BulletsSlide) -> Frame | None:
    theme = ctx.theme
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)

    if spec.image:
        image_w = int(body.w * spec.image_ratio)
        if spec.image_position == "left":
            image_frame = Frame(body.x, body.y, image_w, body.h)
            body = Frame(image_frame.right + ctx.gap, body.y,
                         body.w - image_w - ctx.gap, body.h)
        else:
            image_frame = Frame(body.right - image_w, body.y, image_w, body.h)
            body = Frame(body.x, body.y, body.w - image_w - ctx.gap, body.h)
        E.place_image(ctx, slide, image_frame, spec.image, fit="cover")

    bullets = E.normalise_bullets(spec.bullets)
    if spec.text and not bullets:
        E.add_plain_text(
            ctx, slide, body, spec.text,
            base_size_pt=theme.pt("body"), min_size_pt=theme.pt("body") * 0.68,
            color=theme.text, valign="top", line_spacing=1.28,
        )
    elif bullets:
        items = [(E.bullet_text(item), item.level) for item in bullets]
        size = _fit_or_warn(ctx, items, body, theme.pt("body"), theme.pt("body") * 0.62,
                            f"Slide “{spec.title or 'sans titre'}”")
        E.add_textbox(
            ctx, slide, body, bullets,
            size_pt=size, color=theme.text, valign="top",
            bullet_kind="none" if spec.numbered else "char",
            numbered=spec.numbered,
            line_spacing=1.18, space_before_pt=size * 0.5,
            columns=spec.columns,
        )

    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _two_content_slide(ctx: E.RenderContext, slide, spec: S.TwoContentSlide) -> Frame | None:
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)
    left, right = body.split_h(spec.ratio, ctx.gap)
    E.render_content(ctx, slide, spec.left, left, title=spec.left_title)
    E.render_content(ctx, slide, spec.right, right, title=spec.right_title)
    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _image_slide(ctx: E.RenderContext, slide, spec: S.ImageSlide) -> Frame | None:
    theme = ctx.theme
    geometry = ctx.geometry

    if spec.variant == "full_bleed":
        picture = E.place_image(ctx, slide, geometry.full(), spec.image, fit="cover")
        shape_ooxml.send_to_back(picture)
        if spec.title or spec.caption:
            band_h = int(geometry.height * 0.26)
            band = Frame(0, geometry.height - band_h, geometry.width, band_h)
            veil = E.add_card(ctx, slide, band, fill="000000", radius=0, shape="rect")
            shape_ooxml.solid_fill(veil, "000000", alpha=0.55)
            inner = band.pad(left=geometry.margin_x, right=geometry.margin_x,
                             top=int(band.h * 0.2), bottom=int(band.h * 0.18))
            if spec.title:
                E.add_plain_text(
                    ctx, slide, inner.top_slice(int(inner.h * 0.62)), spec.title,
                    size_pt=theme.pt("slide_title") * 0.85, color="FFFFFF",
                    font=theme.heading_font, bold=True, valign="bottom",
                )
            if spec.caption:
                E.add_plain_text(
                    ctx, slide, inner.bottom_slice(int(inner.h * 0.4)), spec.caption,
                    size_pt=theme.pt("small"), color="E8ECF2", valign="top",
                )
        return

    content = _header(ctx, slide, spec.title, spec.kicker)

    if spec.variant == "split" and spec.text:
        image_frame, text_frame = content.split_h(0.55, ctx.gap)
        E.place_image(ctx, slide, image_frame, spec.image, fit=spec.fit)
        E.add_plain_text(
            ctx, slide, text_frame, spec.text,
            base_size_pt=theme.pt("body"), min_size_pt=theme.pt("body") * 0.7,
            color=theme.text, valign="middle", line_spacing=1.3,
        )
        if spec.caption:
            E.add_plain_text(
                ctx, slide, Frame(image_frame.x, image_frame.bottom - cm(0.8), image_frame.w, cm(0.8)),
                spec.caption, size_pt=theme.pt("caption"), color=theme.muted, valign="bottom",
                autofit=False,
            )
        return

    caption_h = cm(0.9) if spec.caption else 0
    text_h = 0
    if spec.text:
        text_h = min(int(content.h * 0.28), cm(2.6))
    image_frame = content.pad(bottom=caption_h + text_h)
    E.place_image(ctx, slide, image_frame, spec.image, fit=spec.fit)
    cursor = image_frame.bottom
    if spec.caption:
        E.add_plain_text(
            ctx, slide, Frame(content.x, cursor, content.w, caption_h), spec.caption,
            size_pt=theme.pt("caption"), color=theme.muted, align="center", valign="top",
            autofit=False,
        )
        cursor += caption_h
    if spec.text:
        E.add_plain_text(
            ctx, slide, Frame(content.x, cursor, content.w, text_h), spec.text,
            size_pt=theme.pt("small"), color=theme.text, valign="top", line_spacing=1.25,
        )


def _chart_slide(ctx: E.RenderContext, slide, spec: S.ChartSlide) -> Frame | None:
    if spec.chart is None:
        raise InvalidSpec("A 'chart' slide needs a `chart`: the ref of a chart_<type> tool.")
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)

    bullets = E.normalise_bullets(spec.bullets)
    if bullets and spec.bullets_position in ("left", "right"):
        if spec.bullets_position == "left":
            side, chart_frame = body.split_h(0.34, ctx.gap)
        else:
            chart_frame, side = body.split_h(0.64, ctx.gap)
        E.add_textbox(
            ctx, slide, side, bullets,
            color=ctx.theme.text, valign="middle", bullet_kind="char", line_spacing=1.2,
            base_size_pt=ctx.theme.pt("small") + 1, min_size_pt=ctx.theme.pt("small") * 0.85,
        )
    else:
        chart_frame = body

    E.render_element(ctx, slide, spec.chart, chart_frame)
    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _table_slide(ctx: E.RenderContext, slide, spec: S.TableSlide) -> Frame | None:
    if spec.table is None:
        raise InvalidSpec("A 'table' slide needs a `table`: {header, rows}.")
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)
    E.render_element(ctx, slide, spec.table, body)
    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _kpi_slide(ctx: E.RenderContext, slide, spec: S.KpiSlide) -> Frame | None:
    if not spec.items:
        raise InvalidSpec("A 'kpi' slide needs at least one item in `items`.")
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)

    bullets = E.normalise_bullets(spec.bullets)
    extras = bool(spec.chart) or bool(bullets)
    tiles_h = int(body.h * (0.44 if extras else 1.0))
    if not extras:
        tiles_h = min(body.h, max(cm(3.4), int(body.h * 0.62)))
    tiles = Frame(body.x, body.y, body.w, tiles_h)

    E.render_kpi_element(
        ctx, slide,
        S.KpiElement(type="kpi", items=spec.items, columns=spec.columns, style=spec.style),
        tiles,
    )

    rest = body.remaining_below(tiles, ctx.gap)
    if spec.chart and bullets:
        chart_frame, side = rest.split_h(0.62, ctx.gap)
        E.render_element(ctx, slide, spec.chart, chart_frame)
        E.add_textbox(ctx, slide, side, bullets, color=ctx.theme.text, valign="top",
                      bullet_kind="char", line_spacing=1.18)
    elif spec.chart:
        E.render_element(ctx, slide, spec.chart, rest)
    elif bullets:
        E.add_textbox(ctx, slide, rest, bullets, color=ctx.theme.text, valign="top",
                      bullet_kind="char", line_spacing=1.2)

    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return rest


def _comparison_slide(ctx: E.RenderContext, slide, spec: S.ComparisonSlide) -> Frame | None:
    theme = ctx.theme
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)
    left, right = body.split_h(0.5, ctx.gap)

    for column, frame, fallback in (
        (spec.left, left, theme.accent),
        (spec.right, right, theme.accent2),
    ):
        accent = parse_color(column.accent) if column.accent else fallback
        head_h = cm(1.15)
        head = frame.top_slice(head_h)
        E.add_card(ctx, slide, head, fill=accent, radius=0.09)
        label = f"{column.icon}  {column.title}" if (column.icon and column.title) else (
            column.title or column.icon or ""
        )
        if label:
            E.add_plain_text(
                ctx, slide, head, label,
                base_size_pt=theme.pt("small") + 2, min_size_pt=theme.pt("caption"),
                color=theme.on(accent), bold=True, font=theme.heading_font,
                align="center", valign="middle", padding_cm=0.2, max_lines=1,
            )
        panel = Frame(frame.x, head.bottom, frame.w, frame.h - head_h)
        E.add_card(ctx, slide, panel, fill=theme.tinted(accent, 0.93), radius=0.045)
        items = E.normalise_bullets(column.items)
        if items:
            E.add_textbox(
                ctx, slide, panel.inset(cm(0.42)), items,
                color=theme.text, valign="top", bullet_kind="char", bullet_color=accent,
                line_spacing=1.2, min_size_pt=theme.pt("small") * 0.8,
            )

    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _timeline_slide(ctx: E.RenderContext, slide, spec: S.TimelineSlide) -> Frame | None:
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)
    if spec.orientation == "horizontal":
        # A horizontal timeline is a band, not a block: centre it rather than
        # letting it float at the top of a tall content area.
        band_h = min(body.h, cm(6.2))
        body = Frame(body.x, body.y + (body.h - band_h) // 2, body.w, band_h)
    E.render_timeline_element(
        ctx, slide,
        S.TimelineElement(type="timeline", milestones=spec.milestones,
                          orientation=spec.orientation),
        body,
    )
    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _process_slide(ctx: E.RenderContext, slide, spec: S.ProcessSlide) -> Frame | None:
    content = _header(ctx, slide, spec.title, spec.kicker)
    body, takeaway_frame = _takeaway_band(ctx, content, spec.takeaway)
    band_h = min(body.h, cm(5.4 if spec.style == "numbered" else 4.6))
    body = Frame(body.x, body.y + (body.h - band_h) // 2, body.w, band_h)
    E.render_process_element(
        ctx, slide,
        S.ProcessElement(type="process", steps=spec.steps, style=spec.style),
        body,
    )
    if takeaway_frame is not None:
        E.add_takeaway(ctx, slide, takeaway_frame, spec.takeaway)
    return body


def _agenda_slide(ctx: E.RenderContext, slide, spec: S.AgendaSlide) -> Frame | None:
    theme = ctx.theme
    content = _header(ctx, slide, spec.title, spec.kicker)
    items = E.normalise_bullets(spec.items)
    if not items:
        return content

    columns = 2 if len(items) > 6 else 1
    per_column = (len(items) + columns - 1) // columns
    frames = content.columns(columns, ctx.gap * 2) if columns > 1 else [content]

    # Rows take the height they need and stack from the top; spreading four items
    # over 12 cm reads as a mistake, not as breathing room.
    gap = int(ctx.gap * 0.45)
    row_h = min(cm(1.45), max(cm(0.95), (content.h - gap * (per_column - 1)) // max(per_column, 1)))
    badge_w = cm(1.15)

    index = 0
    for column_index, column_frame in enumerate(frames):
        chunk = items[column_index * per_column:(column_index + 1) * per_column]
        for offset, item in enumerate(chunk):
            row = Frame(column_frame.x, column_frame.y + offset * (row_h + gap),
                        column_frame.w, row_h)
            active = spec.active is not None and index == spec.active
            accent = theme.accent if active else theme.muted
            if active:
                E.add_card(ctx, slide, row, fill=theme.tinted(theme.accent, 0.9), radius=0.08)
            E.add_plain_text(
                ctx, slide, Frame(row.x + cm(0.25), row.y, badge_w, row.h), f"{index + 1:02d}",
                size_pt=theme.pt("body"), color=accent, bold=True, font=theme.heading_font,
                valign="middle", autofit=False,
            )
            E.add_textbox(
                ctx, slide,
                Frame(row.x + badge_w + cm(0.3), row.y, row.w - badge_w - cm(0.6), row.h),
                [item],
                base_size_pt=theme.pt("body") * 0.95,
                min_size_pt=theme.pt("small") * 0.85,
                color=theme.text,
                bold=active,
                valign="middle",
                max_lines=2,
            )
            index += 1
    return content


def _quote_slide(ctx: E.RenderContext, slide, spec: S.QuoteSlide) -> Frame | None:
    theme = ctx.theme
    geometry = ctx.geometry
    centred = spec.variant == "centered"
    body = geometry.safe().pad(
        left=int(geometry.width * (0.10 if centred else 0.04)),
        right=int(geometry.width * (0.10 if centred else 0.20)),
    )

    mark_h = cm(2.2)
    mark = slide.shapes.add_textbox(Emu(body.x), Emu(body.y), cm(3.0), mark_h)
    paragraph = mark.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "“"
    run.font.size = Pt(theme.pt("quote") * 3.0)
    run.font.bold = True
    run.font.color.rgb = rgb(theme.tinted(theme.accent, 0.55))
    text_ooxml.set_typeface(run, theme.heading_font)
    if centred:
        paragraph.alignment = E.alignment("center")
        mark.left = Emu(body.cx - cm(1.5))

    # The quotation and its attribution are one block, vertically centred: a quote
    # marooned at the top of the slide with the name at the bottom reads as two
    # unrelated things.
    available = max(cm(2.0), body.h - mark_h)
    block_h = min(available, cm(7.2))
    block = Frame(body.x, body.y + mark_h + (available - block_h) // 2, body.w, block_h)
    attribution = " — ".join(part for part in (spec.author, spec.role) if part)
    attribution_h = cm(1.1) if attribution else 0
    quote_frame = block.pad(bottom=attribution_h)

    E.add_plain_text(
        ctx, slide, quote_frame, spec.text,
        base_size_pt=theme.pt("quote"), min_size_pt=theme.pt("quote") * 0.45,
        color=theme.text, font=theme.heading_font,
        align="center" if centred else "left", valign="middle", line_spacing=1.26,
    )
    if attribution:
        E.add_plain_text(
            ctx, slide, Frame(block.x, quote_frame.bottom, block.w, attribution_h), attribution,
            base_size_pt=theme.pt("small") + 1, color=theme.muted,
            align="center" if centred else "left", valign="top", max_lines=1,
        )


def _closing_slide(ctx: E.RenderContext, slide, spec: S.ClosingSlide) -> Frame | None:
    theme = ctx.theme
    geometry = ctx.geometry
    body = geometry.full().pad(left=int(geometry.width * 0.12), right=int(geometry.width * 0.12))

    top = int(geometry.height * 0.34)
    E.add_plain_text(
        ctx, slide, Frame(body.x, top, body.w, int(geometry.height * 0.18)), spec.title,
        size_pt=theme.pt("deck_title") * 0.92, min_size_pt=theme.pt("section_title") * 0.6,
        color=theme.text, font=theme.heading_font, bold=True, align="center", valign="bottom",
    )
    E.add_rule(
        ctx, slide,
        Frame(body.cx - cm(1.7), top + int(geometry.height * 0.20), cm(3.4), 0),
        theme.accent, thickness_cm=0.1,
    )
    cursor = top + int(geometry.height * 0.24)
    if spec.subtitle:
        E.add_plain_text(
            ctx, slide, Frame(body.x, cursor, body.w, cm(1.4)), spec.subtitle,
            base_size_pt=theme.pt("body"), color=theme.muted, align="center", valign="top",
        )
        cursor += cm(1.5)
    if spec.contact:
        E.add_textbox(
            ctx, slide, Frame(body.x, cursor, body.w, cm(3.0)),
            [S.BulletSpec(text=line) for line in spec.contact],
            size_pt=theme.pt("small") + 1, color=theme.text, align="center", valign="top",
            line_spacing=1.35, autofit=False,
        )
    if spec.logo:
        E.place_image(
            ctx, slide,
            Frame(body.cx - cm(2.1), geometry.height - geometry.margin_bottom - cm(2.2),
                  cm(4.2), cm(1.6)),
            spec.logo, fit="contain",
        )


def _blank_slide(ctx: E.RenderContext, slide, spec: S.BlankSlide) -> Frame | None:
    if spec.title or spec.kicker:
        return _header(ctx, slide, spec.title, spec.kicker)
    return ctx.geometry.safe()


def _template_slide(ctx: E.RenderContext, slide, spec: S.TemplateSlide) -> Frame | None:
    slide_ooxml.clone_footer_placeholders(slide)
    for key, value in (spec.placeholders or {}).items():
        placeholder = None
        if str(key).strip().isdigit():
            placeholder = E.find_placeholder(slide, idx=int(key))
        if placeholder is None:
            placeholder = E.find_placeholder(slide, name=str(key))
        if placeholder is None:
            available = ", ".join(
                f"{p.placeholder_format.idx}:{p.name!r}" for p in slide.placeholders
            )
            raise InvalidSpec(
                f"This layout has no placeholder {key!r}. Available: {available or '(none)'}."
            )
        E.fill_placeholder(ctx, slide, placeholder, value)


_HANDLERS = {
    S.TitleSlide: _title_slide,
    S.SectionSlide: _section_slide,
    S.BulletsSlide: _bullets_slide,
    S.TwoContentSlide: _two_content_slide,
    S.ImageSlide: _image_slide,
    S.ChartSlide: _chart_slide,
    S.TableSlide: _table_slide,
    S.KpiSlide: _kpi_slide,
    S.ComparisonSlide: _comparison_slide,
    S.TimelineSlide: _timeline_slide,
    S.ProcessSlide: _process_slide,
    S.AgendaSlide: _agenda_slide,
    S.QuoteSlide: _quote_slide,
    S.ClosingSlide: _closing_slide,
    S.BlankSlide: _blank_slide,
    S.TemplateSlide: _template_slide,
}
