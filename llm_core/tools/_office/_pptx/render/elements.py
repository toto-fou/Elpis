# SPDX-License-Identifier: MIT
"""Drawing primitives: one function per element type, all taking a computed frame.

Everything above this module decides *where* something goes; this module puts it
there. Keeping the two apart is what lets ``pptx_build`` and the granular
``pptx_add_*`` tools produce byte-identical output from the same spec.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Sequence

from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

from ..errors import InvalidSpec
from ..ooxml import charts as chart_ooxml, shapes as shape_ooxml, tables as table_ooxml, text as text_ooxml
from ..ooxml.util import (
    RPR_ORDER,
    alignment,
    anchor,
    drop,
    frag,
    insert_in_order,
    mix,
    parse_color,
    rgb,
)
from . import images as image_io, spec as S
from .layout import Frame, Geometry, cm, fit_font_size
from .theme import Theme

SHAPE_TYPES: dict[str, MSO_SHAPE] = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "circle": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON,
    "arrow": MSO_SHAPE.RIGHT_ARROW,
    "pentagon": MSO_SHAPE.PENTAGON,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "star": MSO_SHAPE.STAR_5_POINT,
    "plaque": MSO_SHAPE.PLAQUE,
    "callout": MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
    "line": MSO_SHAPE.RECTANGLE,
    "hexagon": MSO_SHAPE.HEXAGON,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
}

_NUMERIC = re.compile(r"^[\s€$£%+\-]*[\d][\d\s.,']*\s*[%€$£kKmMbB]{0,3}\s*$")

TREND_GLYPHS = {"up": "▲", "down": "▼", "flat": "▬"}


@dataclass
class RenderContext:
    """Everything a builder needs that is not the slide itself."""

    presentation: Any
    theme: Theme
    geometry: Geometry
    notes: list[str] = field(default_factory=list)
    footer: Any = None
    #: Where the notes list stood when the current tool call began, so a tool
    #: reports only what it caused rather than everything the deck ever noted.
    mark: int = 0

    def note(self, message: str) -> None:
        if message not in self.notes:
            self.notes.append(message)

    @property
    def gap(self) -> int:
        return self.geometry.gap


# ---------------------------------------------------------------------------
# Frames
# ---------------------------------------------------------------------------


def resolve_frame(ctx: RenderContext, spec: S.FrameSpec | None, default: Frame) -> Frame:
    """Turn a ``FrameSpec`` into a rectangle, falling back to the layout's own."""
    if spec is None:
        return default
    if spec.col is not None or spec.row is not None:
        # The grid is laid over the area this element would otherwise have taken,
        # so a slide with a kicker does not push its own elements under the title.
        base = ctx.geometry.area(spec.area) if spec.area else default
        return base.grid_cell(
            col=spec.col or 0, row=spec.row or 0, cols=spec.cols, rows_=spec.rows,
            gap=int(ctx.gap * 0.5),
        )
    frame = ctx.geometry.area(spec.area) if spec.area else default
    x = cm(spec.x_cm) if spec.x_cm is not None else frame.x
    y = cm(spec.y_cm) if spec.y_cm is not None else frame.y
    w = cm(spec.w_cm) if spec.w_cm is not None else frame.w
    h = cm(spec.h_cm) if spec.h_cm is not None else frame.h
    return Frame(x, y, w, h)


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------


def bullet_text(bullet: S.BulletSpec) -> str:
    if bullet.text:
        return bullet.text
    return "".join(run.text or "" for run in bullet.runs)


def normalise_bullets(items: Sequence[Any]) -> list[S.BulletSpec]:
    out: list[S.BulletSpec] = []
    for item in items or []:
        if isinstance(item, S.BulletSpec):
            out.append(item)
        elif isinstance(item, dict):
            out.append(S.BulletSpec.model_validate(item))
        else:
            out.append(S.BulletSpec(text=str(item)))
    return out


def apply_run_format(run, theme: Theme, spec: S.RunSpec, *, default_size=None,
                     default_color=None, default_font=None) -> None:
    font = run.font
    size = spec.size_pt or default_size
    if size:
        font.size = Pt(float(size))
    if spec.bold is not None:
        font.bold = bool(spec.bold)
    if spec.italic is not None:
        font.italic = bool(spec.italic)
    if spec.underline is not None:
        font.underline = bool(spec.underline)
    color = spec.color or default_color
    if color:
        font.color.rgb = rgb(color)
    typeface = spec.font or (theme.mono_font if spec.code else None) or default_font
    if typeface:
        text_ooxml.set_typeface(run, typeface)
    if spec.strike:
        text_ooxml.set_strike(run, True)
    if spec.caps:
        text_ooxml.set_caps(run, spec.caps)
    if spec.baseline:
        text_ooxml.set_baseline(run, spec.baseline)
    if spec.letter_spacing_pt:
        text_ooxml.set_char_spacing(run, spec.letter_spacing_pt)
    if spec.highlight:
        text_ooxml.set_highlight(run, spec.highlight)


def set_link(ctx: RenderContext, slide, run, target: str) -> None:
    """External URL, or ``#3`` to jump to slide 3 (0-based)."""
    text = str(target).strip()
    if not text.startswith("#"):
        run.hyperlink.address = text
        return
    try:
        index = int(text[1:])
    except ValueError:
        raise InvalidSpec(
            f"Internal link {target!r} must be '#N' with N a 0-based slide index."
        ) from None
    slides = ctx.presentation.slides
    if not 0 <= index < len(slides):
        raise InvalidSpec(
            f"Internal link {target!r} points at slide {index}, but the deck has {len(slides)}."
        )
    rel_id = slide.part.relate_to(slides[index].part, RT.SLIDE)
    rpr = run._r.get_or_add_rPr()
    drop(rpr, "a:hlinkClick")
    insert_in_order(rpr, frag(
        '<a:hlinkClick xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{rel_id}" action="ppaction://hlinksldjump"/>'
    ), RPR_ORDER)


def write_paragraph(
    ctx: RenderContext,
    slide,
    paragraph,
    bullet: S.BulletSpec,
    *,
    size_pt: float,
    color: str,
    font: str | None,
    bold: bool | None = None,
    italic: bool | None = None,
    align: str | None = None,
    line_spacing: float | None = None,
    space_before_pt: float | None = None,
    bullet_kind: str = "none",
    bullet_color: str | None = None,
    numbered: bool = False,
) -> None:
    paragraph.level = min(max(int(bullet.level or 0), 0), 4)
    if align:
        paragraph.alignment = alignment(align)
    if line_spacing:
        paragraph.line_spacing = float(line_spacing)
    if space_before_pt is not None:
        paragraph.space_before = Pt(float(space_before_pt))
    paragraph.space_after = Pt(0)

    size = float(bullet.size_pt or size_pt)
    text_color = bullet.color or color
    runs = bullet.runs or [S.RunSpec(text=bullet_text(bullet))]
    for run_spec in runs:
        run = paragraph.add_run()
        run.text = run_spec.text or ""
        apply_run_format(
            run, ctx.theme, run_spec,
            default_size=size,
            default_color=text_color,
            default_font=font,
        )
        if bold is not None and run_spec.bold is None:
            run.font.bold = bool(bold)
        if bullet.bold is not None and run_spec.bold is None:
            run.font.bold = bool(bullet.bold)
        if italic is not None and run_spec.italic is None:
            run.font.italic = bool(italic)
        if bullet.italic is not None and run_spec.italic is None:
            run.font.italic = bool(bullet.italic)
        if run_spec.link:
            set_link(ctx, slide, run, run_spec.link)

    explicit = bullet.bullet
    if explicit == "":
        text_ooxml.set_bullet(paragraph, "none")
    elif explicit:
        text_ooxml.set_bullet(paragraph, "char", char=explicit,
                              color=bullet_color or ctx.theme.accent, size_pct=95)
    elif numbered:
        text_ooxml.set_bullet(paragraph, "number", color=bullet_color or ctx.theme.accent)
    elif bullet_kind == "char":
        text_ooxml.set_bullet(paragraph, "char", color=bullet_color or ctx.theme.accent,
                              size_pct=72)
    else:
        text_ooxml.set_bullet(paragraph, "none")
    if bullet.space_before_pt is not None:
        paragraph.space_before = Pt(float(bullet.space_before_pt))


def add_textbox(
    ctx: RenderContext,
    slide,
    frame: Frame,
    bullets: Sequence[S.BulletSpec],
    *,
    size_pt: float | None = None,
    base_size_pt: float | None = None,
    min_size_pt: float | None = None,
    color: str | None = None,
    font: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    align: str | None = None,
    valign: str | None = "top",
    line_spacing: float = 1.12,
    space_before_pt: float | None = None,
    bullet_kind: str = "none",
    bullet_color: str | None = None,
    numbered: bool = False,
    padding_cm: float = 0.0,
    wrap: bool = True,
    autofit: bool = True,
    columns: int = 1,
    max_lines: int | None = None,
    char_ratio: float | None = None,
):
    """The one place a text box is created. Everything else calls through here.

    ``size_pt`` fixes the size; ``base_size_pt`` gives the fitter a starting point
    and lets it shrink from there. Passing neither starts from the body size.
    """
    box = slide.shapes.add_textbox(Emu(frame.x), Emu(frame.y), Emu(frame.w), Emu(frame.h))
    text_frame = box.text_frame
    text_frame.word_wrap = wrap
    inset = cm(padding_cm)
    text_ooxml.set_insets(text_frame, inset, inset, inset, inset)
    if valign:
        text_frame.vertical_anchor = anchor(valign)
    if columns > 1:
        text_ooxml.set_columns(text_frame, columns, int(ctx.gap * 0.8))

    items = normalise_bullets(bullets)
    if not items:
        items = [S.BulletSpec(text="")]

    base = size_pt or base_size_pt or ctx.theme.pt("body")
    gap_pt = space_before_pt if space_before_pt is not None else base * 0.42
    # Bold and heading faces run wider than the average the fitter assumes.
    ratio = char_ratio if char_ratio is not None else (0.56 if bold else 0.5)
    if autofit and size_pt is None:
        base = fit_font_size(
            [(bullet_text(item), item.level) for item in items],
            frame,
            base_pt=base,
            min_pt=min_size_pt or max(9.0, base * 0.62),
            line_spacing=line_spacing,
            space_before_pt=gap_pt,
            inset_cm=padding_cm * 2 + 0.3,
            max_lines=max_lines,
            char_ratio=ratio,
        )
        gap_pt = base * 0.42 if space_before_pt is None else space_before_pt

    text_frame.clear()
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        write_paragraph(
            ctx, slide, paragraph, item,
            size_pt=base,
            color=color or ctx.theme.text,
            font=font or ctx.theme.body_font,
            bold=bold,
            italic=italic,
            align=align,
            line_spacing=line_spacing,
            space_before_pt=0 if index == 0 else gap_pt,
            bullet_kind=bullet_kind,
            bullet_color=bullet_color,
            numbered=numbered,
        )
    text_ooxml.set_autofit(text_frame, "shrink" if autofit else "none")
    return box


def add_plain_text(ctx: RenderContext, slide, frame: Frame, text: str, **kwargs):
    return add_textbox(ctx, slide, frame, [S.BulletSpec(text=text or "")], **kwargs)


# ---------------------------------------------------------------------------
# Furniture
# ---------------------------------------------------------------------------


def new_shape(slide, kind, frame: Frame):
    """Every autoshape this module draws is created here, flat and unstyled."""
    shape = slide.shapes.add_shape(kind, Emu(frame.x), Emu(frame.y), Emu(frame.w), Emu(frame.h))
    shape_ooxml.strip_style(shape)
    return shape


def add_card(
    ctx: RenderContext,
    slide,
    frame: Frame,
    *,
    fill: str | None = None,
    gradient: Sequence[str] | None = None,
    gradient_angle: float = 90.0,
    border: str | None = None,
    border_width_pt: float = 1.0,
    radius: float | None = 0.055,
    shadow: bool = False,
    shape: str = "rounded_rect",
):
    """A filled rectangle behind content — the tile, the callout, the banner."""
    kind = SHAPE_TYPES.get(str(shape).lower(), MSO_SHAPE.ROUNDED_RECTANGLE)
    box = new_shape(slide, kind, frame)
    box.text_frame.word_wrap = True
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        shape_ooxml.set_adjustment(box, 0, float(radius))
    if gradient:
        shape_ooxml.gradient_fill(box, [parse_color(c) for c in gradient], angle_deg=gradient_angle)
    elif fill:
        shape_ooxml.solid_fill(box, fill)
    else:
        shape_ooxml.clear_fill(box)
    if border:
        shape_ooxml.outline(box, border, border_width_pt)
    else:
        shape_ooxml.no_outline(box)
    if shadow:
        shape_ooxml.soft_shadow(box)
    else:
        shape_ooxml.no_shadow(box)
    return box


def add_rule(ctx: RenderContext, slide, frame: Frame, color: str | None = None,
             thickness_cm: float = 0.05):
    """A hairline. A rectangle, not a connector — connectors move when text reflows."""
    line = new_shape(slide, MSO_SHAPE.RECTANGLE,
                     Frame(frame.x, frame.y, frame.w, cm(thickness_cm)))
    shape_ooxml.solid_fill(line, color or ctx.theme.border)
    shape_ooxml.no_outline(line)
    shape_ooxml.no_shadow(line)
    line.text_frame.word_wrap = False
    return line


def add_accent_bar(ctx: RenderContext, slide, frame: Frame, color: str | None = None,
                   width_cm: float = 0.11):
    bar = new_shape(slide, MSO_SHAPE.RECTANGLE,
                    Frame(frame.x, frame.y, cm(width_cm), frame.h))
    shape_ooxml.solid_fill(bar, color or ctx.theme.accent)
    shape_ooxml.no_outline(bar)
    shape_ooxml.no_shadow(bar)
    return bar


def add_title(ctx: RenderContext, slide, frame: Frame, text: str, *, size_pt: float | None = None,
              color: str | None = None, align: str | None = None, valign: str = "bottom",
              font: str | None = None, max_lines: int | None = 2):
    """A title, shrunk to fit its band rather than allowed to run into the content."""
    base = size_pt or ctx.theme.pt("slide_title")
    return add_plain_text(
        ctx, slide, frame, text,
        base_size_pt=base,
        min_size_pt=base * 0.55,
        color=color or ctx.theme.text,
        font=font or ctx.theme.heading_font,
        bold=True,
        align=align,
        valign=valign,
        line_spacing=1.06,
        max_lines=max_lines,
    )


def add_kicker(ctx: RenderContext, slide, frame: Frame, text: str, color: str | None = None,
               align: str | None = None):
    box = add_textbox(
        ctx, slide, frame,
        [S.BulletSpec(runs=[S.RunSpec(
            text=text, caps="all", letter_spacing_pt=1.1, bold=True,
        )])],
        size_pt=ctx.theme.pt("kicker"),
        color=color or ctx.theme.accent,
        font=ctx.theme.body_font,
        align=align,
        valign="bottom",
        autofit=False,
    )
    return box


def add_takeaway(ctx: RenderContext, slide, frame: Frame, text: str, color: str | None = None):
    """The claim under a chart or table. Boxed, accented, and never optional-looking."""
    theme = ctx.theme
    accent = color or theme.accent
    card = add_card(ctx, slide, frame, fill=theme.tinted(accent, 0.9), radius=0.05)
    add_accent_bar(ctx, slide, frame, accent, width_cm=0.1)
    inner = frame.pad(left=cm(0.42), top=cm(0.1), right=cm(0.25), bottom=cm(0.1))
    add_plain_text(
        ctx, slide, inner, text,
        size_pt=theme.pt("small"),
        color=theme.on(theme.tinted(accent, 0.9)),
        font=theme.body_font,
        bold=True,
        valign="middle",
        line_spacing=1.15,
        padding_cm=0.12,
    )
    return card


# ---------------------------------------------------------------------------
# Element renderers
# ---------------------------------------------------------------------------


def render_text_element(ctx: RenderContext, slide, element: S.TextElement, frame: Frame):
    theme = ctx.theme
    fill = element.fill
    card = None
    if fill or element.gradient or element.border:
        card = add_card(
            ctx, slide, frame,
            fill=fill,
            gradient=element.gradient,
            border=element.border.color if element.border else None,
            border_width_pt=element.border.width_pt if element.border else 1.0,
            radius=element.radius if element.radius is not None else 0.055,
            shadow=element.shadow,
        )
    if element.accent_bar:
        add_accent_bar(ctx, slide, frame, element.accent_bar)

    padding = element.padding_cm
    if padding is None:
        padding = 0.32 if (fill or element.gradient or element.border) else 0.0
    inner = frame
    if element.accent_bar:
        inner = inner.pad(left=cm(0.34))

    bullets = normalise_bullets(element.bullets)
    if not bullets:
        if element.runs:
            bullets = [S.BulletSpec(runs=list(element.runs))]
        else:
            bullets = [S.BulletSpec(text=element.text or "")]

    default_color = element.color or (theme.on(fill) if fill else theme.text)
    box = add_textbox(
        ctx, slide, inner, bullets,
        size_pt=element.font_size_pt,
        min_size_pt=element.min_font_size_pt,
        color=default_color,
        font=element.font or theme.body_font,
        bold=element.bold,
        italic=element.italic,
        align=element.align,
        valign=element.valign or "top",
        line_spacing=element.line_spacing or 1.14,
        space_before_pt=element.space_before_pt,
        bullet_kind="char" if (element.bullets and not element.numbered) else "none",
        numbered=element.numbered,
        padding_cm=padding,
        autofit=element.autofit,
        columns=element.columns,
    )
    return card or box


def render_image_element(ctx: RenderContext, slide, element: S.ImageElement, frame: Frame):
    caption_h = cm(0.85) if element.caption else 0
    box = frame.pad(bottom=caption_h)
    picture = place_image(
        ctx, slide, box, element.source,
        fit=element.fit, align=element.align, valign=element.valign,
    )
    if element.border and element.border.color:
        shape_ooxml.outline(picture, element.border.color, element.border.width_pt)
    if element.shadow:
        shape_ooxml.soft_shadow(picture)
    if element.alt_text:
        shape_ooxml.set_alt_text(picture, element.alt_text)
    if element.caption:
        add_plain_text(
            ctx, slide, Frame(frame.x, box.bottom, frame.w, caption_h), element.caption,
            size_pt=ctx.theme.pt("caption"),
            color=ctx.theme.muted,
            align=element.align or "center",
            valign="top",
            autofit=False,
        )
    return picture


def place_image(ctx: RenderContext, slide, frame: Frame, source: str, *,
                fit: str = "contain", align: str | None = "center",
                valign: str | None = "middle"):
    """Insert a picture sized to the frame under the requested fit rule."""
    resolved = image_io.resolve_image(source)
    if fit == "stretch":
        width, height = frame.w, frame.h
    elif fit == "cover":
        width_cm, height_cm = image_io.cover(resolved.aspect, frame.w_cm, frame.h_cm)
        width, height = cm(width_cm), cm(height_cm)
    else:
        width_cm, height_cm = image_io.contain(resolved.aspect, frame.w_cm, frame.h_cm)
        width, height = cm(width_cm), cm(height_cm)

    if align in (None, "center", "centre"):
        left = frame.x + (frame.w - width) // 2
    elif align == "right":
        left = frame.right - width
    else:
        left = frame.x
    if valign in (None, "middle", "center", "centre"):
        top = frame.y + (frame.h - height) // 2
    elif valign == "bottom":
        top = frame.bottom - height
    else:
        top = frame.y

    picture = slide.shapes.add_picture(
        resolved.rewound(), Emu(left), Emu(top), Emu(width), Emu(height)
    )
    if fit == "cover":
        picture.left, picture.top = Emu(frame.x), Emu(frame.y)
        image_io.crop_to_box(picture, frame.w, frame.h)
    return picture


def render_chart_element(ctx: RenderContext, slide, element: S.ChartElement, frame: Frame):
    key = chart_ooxml.normalise_type(element.chart_type)
    native = key in chart_ooxml.NATIVE_TYPES

    # Graphiques natifs seulement : les autres types arrivent déjà en image
    # (rendu ECharts côté serveur, voir ``_office/graphiques.py``).
    if not native:
        raise InvalidSpec(
            f"Unknown chart type {element.chart_type!r}. Native types: "
            f"{', '.join(chart_ooxml.native_chart_types())}."
        )

    request = chart_ooxml.ChartRequest(
        chart_type=key,
        categories=list(element.categories),
        series=[
            chart_ooxml.SeriesData(name=s.name, values=s.values, points=s.points, color=s.color)
            for s in element.series
        ],
        title=element.title,
        x_title=element.x_title,
        y_title=element.y_title,
        legend=element.legend,
        data_labels=element.data_labels,
        data_label_position=element.data_label_position,
        number_format=element.number_format,
        gridlines=element.gridlines,
        y_min=element.y_min,
        y_max=element.y_max,
        palette=element.palette or ctx.theme.palette(),
        gap_width=element.gap_width,
        overlap=element.overlap,
        hole_size=element.hole_size,
        smooth=element.smooth,
        font=ctx.theme.body_font,
        text_color=ctx.theme.muted if ctx.theme.is_dark else chart_ooxml.TEXT_COLOR,
        font_size_pt=ctx.theme.pt("caption") + 0.5,
    )
    graphic = chart_ooxml.add_chart(slide, request, frame.x, frame.y, frame.w, frame.h)
    if len(element.series) > len(ctx.theme.palette()):
        ctx.note(
            f"{len(element.series)} series exceeds the {len(ctx.theme.palette())}-colour palette, "
            "so colours repeat. Group the tail into an 'Other' series, or split the chart."
        )
    if element.alt_text:
        shape_ooxml.set_alt_text(graphic, element.alt_text)
    return graphic


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, S.TableCellSpec):
        return value.text or ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def render_table_element(ctx: RenderContext, slide, element: S.TableElement, frame: Frame):
    theme = ctx.theme
    grid: list[list[Any]] = []
    if element.header:
        grid.append(list(element.header))
    grid.extend([list(row) for row in element.rows])
    if not grid:
        raise InvalidSpec("A table needs 'header' and/or 'rows'.")
    columns = max(len(row) for row in grid)

    caption_h = cm(0.8) if element.caption else 0
    box = frame.pad(bottom=caption_h)

    size_pt = element.font_size_pt or max(9.0, theme.pt("small") - (0.5 if len(grid) > 8 else 0))

    # Rows get the height their text needs, not an equal share of the box: a
    # three-row table stretched over 12 cm reads as a layout accident.
    header_h = cm(1.0) if element.header else 0
    body_rows = max(1, len(grid) - (1 if element.header else 0))
    natural = cm(size_pt * 2.05 / 28.3465 + 0.2)
    if element.row_height_cm:
        body_h = cm(element.row_height_cm)
    else:
        available = max(0, box.h - header_h)
        body_h = max(cm(0.72), min(natural, available // body_rows))
    total_h = header_h + body_h * body_rows

    graphic = slide.shapes.add_table(
        len(grid), columns, Emu(box.x), Emu(box.y), Emu(box.w), Emu(min(total_h, box.h))
    )
    table = graphic.table
    table_ooxml.set_table_style(table, "none")
    table_ooxml.set_banding(table, first_row=bool(element.header), band_rows=element.style == "banded")

    if element.col_widths:
        weights = [max(0.0001, float(w)) for w in element.col_widths][:columns]
        weights += [1.0] * (columns - len(weights))
        total = sum(weights)
        for index, weight in enumerate(weights):
            table.columns[index].width = Emu(int(box.w * weight / total))

    for index, row in enumerate(table.rows):
        row.height = Emu(header_h if (element.header and index == 0) else body_h)

    header_fill = element.header_fill or (
        theme.accent if element.style == "dark_header" else theme.tinted(theme.accent, 0.9)
    )
    header_color = element.header_color or theme.on(header_fill)
    column_align = _column_alignments(grid, columns, element)

    for row_index, row in enumerate(grid):
        is_header = bool(element.header) and row_index == 0
        for col_index in range(columns):
            value = row[col_index] if col_index < len(row) else None
            cell = table.cell(row_index, col_index)
            _fill_table_cell(
                ctx, cell, value,
                header=is_header,
                element=element,
                size_pt=size_pt,
                header_fill=header_fill,
                header_color=header_color,
                first_column=col_index == 0,
                striped=element.style == "banded" and not is_header and row_index % 2 == 0,
                column_align=column_align[col_index],
            )
            _cell_borders(ctx, cell, element, is_header)

    for merge in element.merges:
        table_ooxml.merge_region(
            table, merge.row_start, merge.col_start, merge.row_end, merge.col_end
        )

    if element.caption:
        add_plain_text(
            ctx, slide, Frame(frame.x, box.y + min(total_h, box.h) + cm(0.1), frame.w, caption_h),
            element.caption,
            size_pt=theme.pt("caption"), color=theme.muted, valign="top", autofit=False,
        )
    return graphic


def _column_alignments(grid, columns: int, element: S.TableElement) -> list[str | None]:
    """One alignment per column, decided by the body rows.

    A header must sit over its column, not over its own content type: "Var." above
    a column of percentages belongs on the right with them.
    """
    body = grid[1:] if element.header else grid
    aligns: list[str | None] = []
    for index in range(columns):
        cells = [row[index] for row in body if index < len(row)]
        texts = [_cell_text(cell) for cell in cells]
        filled = [text for text in texts if text.strip()]
        numeric = sum(1 for text in filled if _NUMERIC.match(text))
        if filled and numeric >= max(1, len(filled) * 0.6):
            aligns.append(element.numeric_align)
        else:
            aligns.append(element.align)
    return aligns


def _fill_table_cell(ctx, cell, value, *, header, element, size_pt, header_fill,
                     header_color, first_column, striped, column_align=None) -> None:
    theme = ctx.theme
    spec = value if isinstance(value, S.TableCellSpec) else S.TableCellSpec(text=_cell_text(value))
    text = spec.text or ""

    table_ooxml.set_cell_margins(cell, left=0.22, right=0.22, top=0.1, bottom=0.1)
    cell.vertical_anchor = anchor("middle")

    if header:
        table_ooxml.set_cell_fill(cell, spec.fill or header_fill)
    elif spec.fill:
        table_ooxml.set_cell_fill(cell, spec.fill)
    elif striped:
        table_ooxml.set_cell_fill(cell, theme.surface)
    else:
        table_ooxml.set_cell_fill(cell, None)

    frame = cell.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    # A cell that carries a newline means it, so break the line rather than
    # letting the character vanish into the XML.
    segments = text.split("\n")
    run = paragraph.add_run()
    run.text = segments[0]
    for segment in segments[1:]:
        paragraph.add_line_break()
        extra = paragraph.add_run()
        extra.text = segment

    align = spec.align or column_align
    if align:
        paragraph.alignment = alignment(align)

    font = run.font
    font.size = Pt(float(spec.size_pt or size_pt))
    font.bold = spec.bold if spec.bold is not None else (
        True if header or (element.first_column_bold and first_column) else False
    )
    if spec.italic is not None:
        font.italic = bool(spec.italic)
    font.color.rgb = rgb(spec.color or (header_color if header else theme.text))
    text_ooxml.set_typeface(run, theme.body_font)


def _cell_borders(ctx, cell, element, is_header: bool) -> None:
    theme = ctx.theme
    style = element.style
    if style == "bordered":
        table_ooxml.set_cell_border(cell, color=theme.border, width_pt=0.75)
        return
    table_ooxml.set_cell_border(cell, ("left", "right"), color=None)
    if style == "minimal":
        table_ooxml.set_cell_border(cell, ("top",), color=None)
        table_ooxml.set_cell_border(
            cell, ("bottom",),
            color=theme.border if not is_header else theme.accent,
            width_pt=1.5 if is_header else 0.75,
        )
        return
    # "clean" / "banded": a rule under the header, hairlines between rows.
    table_ooxml.set_cell_border(cell, ("top",), color=None)
    table_ooxml.set_cell_border(
        cell, ("bottom",),
        color=theme.accent if is_header else theme.border,
        width_pt=1.75 if is_header else 0.75,
    )


def render_shape_element(ctx: RenderContext, slide, element: S.ShapeElement, frame: Frame):
    theme = ctx.theme
    if str(element.shape).lower() == "line":
        border_color = element.border.color if element.border else None
        return add_rule(ctx, slide, frame, element.fill or border_color or theme.border)
    kind = SHAPE_TYPES.get(str(element.shape).lower())
    if kind is None:
        raise InvalidSpec(
            f"Unknown shape {element.shape!r}. Use one of: {', '.join(sorted(SHAPE_TYPES))}."
        )
    box = new_shape(slide, kind, frame)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape_ooxml.set_adjustment(box, 0, element.radius if element.radius is not None else 0.055)
    fill = element.fill or theme.accent
    if element.gradient:
        shape_ooxml.gradient_fill(box, element.gradient, angle_deg=element.gradient_angle)
    else:
        shape_ooxml.solid_fill(box, fill)
    if element.border and element.border.color:
        shape_ooxml.outline(box, element.border.color, element.border.width_pt, element.border.dash)
    else:
        shape_ooxml.no_outline(box)
    if element.shadow:
        shape_ooxml.soft_shadow(box)
    else:
        shape_ooxml.no_shadow(box)
    if element.rotation:
        box.rotation = float(element.rotation)

    if element.text or element.runs:
        text_frame = box.text_frame
        text_frame.word_wrap = True
        text_ooxml.set_insets(text_frame, cm(0.2), cm(0.12), cm(0.2), cm(0.12))
        text_frame.vertical_anchor = anchor(element.valign or "middle")
        bullet = S.BulletSpec(runs=list(element.runs)) if element.runs else S.BulletSpec(text=element.text)
        write_paragraph(
            ctx, slide, text_frame.paragraphs[0], bullet,
            size_pt=element.font_size_pt or theme.pt("small"),
            color=element.text_color or theme.on(element.gradient[0] if element.gradient else fill),
            font=theme.body_font,
            bold=element.bold if element.bold is not None else True,
            align=element.align or "center",
        )
        text_ooxml.set_autofit(text_frame, "shrink")
    return box


def render_kpi_element(ctx: RenderContext, slide, element: S.KpiElement, frame: Frame):
    items = list(element.items)
    if not items:
        raise InvalidSpec("A kpi element needs at least one item.")
    count = element.columns or len(items)
    count = max(1, min(count, len(items)))
    rows = (len(items) + count - 1) // count
    row_frames = frame.rows(rows, ctx.gap) if rows > 1 else [frame]

    shapes = []
    for row_index in range(rows):
        chunk = items[row_index * count:(row_index + 1) * count]
        if not chunk:
            continue
        cells = row_frames[row_index].columns(count, ctx.gap)
        for item, cell in zip(chunk, cells):
            shapes.append(_render_kpi_tile(ctx, slide, cell, item, element.style))
    return shapes[0] if shapes else None


def _render_kpi_tile(ctx: RenderContext, slide, frame: Frame, item: S.KpiSpec, style: str):
    theme = ctx.theme
    accent = parse_color(item.color) if item.color else theme.accent
    if style == "tile":
        card = add_card(ctx, slide, frame, fill=theme.surface, radius=0.045)
        # A colour rule along the top edge is what makes a tile look deliberate.
        add_rule(ctx, slide, frame.top_slice(0), accent, thickness_cm=0.09)
        on_fill = theme.on(theme.surface)
    elif style == "outline":
        card = add_card(ctx, slide, frame, fill=None, border=theme.border, radius=0.045)
        on_fill = theme.text
    else:
        card = None
        on_fill = theme.text

    inner = frame.pad(left=cm(0.42), right=cm(0.42), top=cm(0.36), bottom=cm(0.3))
    has_caption = bool(item.caption)
    weights = [1.55, 0.8] + ([0.7] if has_caption else [])
    bands = inner.rows(len(weights), 0, weights=weights)

    # The delta gets its own column so a long one shrinks itself rather than
    # forcing the number it annotates onto a second line.
    value_frame, delta_frame = bands[0], None
    if item.delta:
        value_frame, delta_frame = bands[0].split_h(0.68, cm(0.12))

    add_textbox(
        ctx, slide, value_frame, [S.BulletSpec(text=item.value)],
        base_size_pt=theme.pt("kpi_value"),
        min_size_pt=theme.pt("kpi_value") * 0.42,
        char_ratio=0.64,
        color=accent if style != "plain" else theme.text,
        font=theme.heading_font,
        bold=True,
        valign="bottom",
        line_spacing=1.0,
        max_lines=1,
    )

    if delta_frame is not None:
        good = (item.trend == "up" and item.good_is == "up") or (
            item.trend == "down" and item.good_is == "down"
        )
        delta_color = (
            theme.muted if item.trend in (None, "flat")
            else ("1E8E5A" if good else "C8453F")
        )
        glyph = TREND_GLYPHS.get(item.trend or "", "")
        add_textbox(
            ctx, slide, delta_frame,
            [S.BulletSpec(text=f"{glyph} {item.delta}".strip())],
            base_size_pt=theme.pt("kpi_label") * 1.05,
            min_size_pt=theme.pt("caption") * 0.8,
            color=delta_color,
            bold=True,
            align="right",
            valign="bottom",
            max_lines=1,
        )

    if item.label:
        add_plain_text(
            ctx, slide, bands[1], item.label,
            base_size_pt=theme.pt("kpi_label"),
            min_size_pt=theme.pt("caption") * 0.85,
            color=on_fill,
            bold=True,
            valign="top",
            line_spacing=1.1,
            max_lines=2,
        )
    if has_caption:
        add_plain_text(
            ctx, slide, bands[2], item.caption,
            size_pt=theme.pt("caption"),
            color=theme.muted,
            valign="top",
            line_spacing=1.1,
        )
    return card or slide.shapes[-1]


def render_timeline_element(ctx: RenderContext, slide, element: S.TimelineElement, frame: Frame):
    milestones = list(element.milestones)
    if not milestones:
        raise InvalidSpec("A timeline needs at least one milestone.")
    if element.orientation == "vertical":
        return _timeline_vertical(ctx, slide, frame, milestones)
    return _timeline_horizontal(ctx, slide, frame, milestones)


def _timeline_horizontal(ctx: RenderContext, slide, frame: Frame, milestones):
    theme = ctx.theme
    axis_y = frame.y + int(frame.h * 0.34)
    add_rule(ctx, slide, Frame(frame.x, axis_y, frame.w, 0), theme.border, thickness_cm=0.055)

    columns = frame.columns(len(milestones), int(ctx.gap * 0.6))
    dot = cm(0.34)
    for milestone, column in zip(milestones, columns):
        centre = column.cx
        marker = new_shape(
            slide, MSO_SHAPE.OVAL,
            Frame(centre - dot // 2, axis_y - dot // 2 + cm(0.028), dot, dot),
        )
        shape_ooxml.solid_fill(marker, theme.accent if milestone.done else theme.background)
        shape_ooxml.outline(marker, theme.accent, 1.75)
        shape_ooxml.no_shadow(marker)

        if milestone.date:
            add_plain_text(
                ctx, slide, Frame(column.x, frame.y, column.w, axis_y - frame.y - cm(0.34)),
                milestone.date,
                size_pt=theme.pt("caption"), color=theme.accent, bold=True,
                align="center", valign="bottom", autofit=False,
            )
        body_top = axis_y + cm(0.42)
        bullets = [S.BulletSpec(runs=[S.RunSpec(text=milestone.title, bold=True)])]
        if milestone.text:
            bullets.append(S.BulletSpec(
                runs=[S.RunSpec(text=milestone.text, color=theme.muted)],
                space_before_pt=3,
            ))
        add_textbox(
            ctx, slide, Frame(column.x, body_top, column.w, max(cm(1.0), frame.bottom - body_top)),
            bullets,
            size_pt=theme.pt("small"),
            min_size_pt=9.0,
            color=theme.text,
            align="center",
            valign="top",
            line_spacing=1.12,
        )
    return slide.shapes[-1]


def _timeline_vertical(ctx: RenderContext, slide, frame: Frame, milestones):
    theme = ctx.theme
    axis_x = frame.x + cm(0.95)
    spine = new_shape(slide, MSO_SHAPE.RECTANGLE,
                      Frame(axis_x, frame.y, cm(0.055), frame.h))
    shape_ooxml.solid_fill(spine, theme.border)
    shape_ooxml.no_outline(spine)
    shape_ooxml.no_shadow(spine)

    rows = frame.rows(len(milestones), int(ctx.gap * 0.5))
    dot = cm(0.34)
    for milestone, row in zip(milestones, rows):
        marker = new_shape(
            slide, MSO_SHAPE.OVAL,
            Frame(axis_x - dot // 2 + cm(0.028), row.y + cm(0.18), dot, dot),
        )
        shape_ooxml.solid_fill(marker, theme.accent if milestone.done else theme.background)
        shape_ooxml.outline(marker, theme.accent, 1.75)
        shape_ooxml.no_shadow(marker)

        bullets = []
        if milestone.date:
            bullets.append(S.BulletSpec(runs=[S.RunSpec(
                text=milestone.date, color=theme.accent, bold=True,
                size_pt=theme.pt("caption"), caps="all", letter_spacing_pt=0.8,
            )]))
        bullets.append(S.BulletSpec(runs=[S.RunSpec(text=milestone.title, bold=True)],
                                    space_before_pt=2 if bullets else 0))
        if milestone.text:
            bullets.append(S.BulletSpec(
                runs=[S.RunSpec(text=milestone.text, color=theme.muted)], space_before_pt=2
            ))
        add_textbox(
            ctx, slide, Frame(axis_x + cm(0.65), row.y, frame.right - axis_x - cm(0.65), row.h),
            bullets,
            size_pt=theme.pt("small"),
            color=theme.text,
            valign="middle",
            line_spacing=1.12,
        )
    return spine


def render_process_element(ctx: RenderContext, slide, element: S.ProcessElement, frame: Frame):
    steps = list(element.steps)
    if not steps:
        raise InvalidSpec("A process needs at least one step.")
    theme = ctx.theme

    if element.style == "numbered":
        return _process_numbered(ctx, slide, frame, steps)

    overlap = cm(0.28) if element.style == "chevron" else int(ctx.gap * 0.6)
    kind = MSO_SHAPE.CHEVRON if element.style == "chevron" else MSO_SHAPE.PENTAGON
    band_h = min(frame.h, cm(2.35))
    band = Frame(frame.x, frame.y, frame.w, band_h)
    count = len(steps)
    width = (band.w + overlap * (count - 1)) // count

    for index, step in enumerate(steps):
        left = band.x + index * (width - overlap)
        weight = 0.16 + 0.62 * (index / max(count - 1, 1))
        fill = mix(theme.accent, theme.background, 0.78 - weight * 0.78)
        shape = new_shape(slide, kind, Frame(left, band.y, width, band.h))
        shape_ooxml.solid_fill(shape, fill)
        shape_ooxml.no_outline(shape)
        shape_ooxml.no_shadow(shape)
        text_frame = shape.text_frame
        text_frame.word_wrap = True
        text_ooxml.set_insets(text_frame, cm(0.55), cm(0.18), cm(0.5), cm(0.18))
        text_frame.vertical_anchor = anchor("middle")
        write_paragraph(
            ctx, slide, text_frame.paragraphs[0],
            S.BulletSpec(runs=[S.RunSpec(text=step.title, bold=True)]),
            size_pt=theme.pt("small"), color=theme.on(fill), font=theme.body_font, align="center",
        )
        text_ooxml.set_autofit(text_frame, "shrink")

        if step.text:
            column_w = band.w // count
            add_plain_text(
                ctx, slide,
                Frame(band.x + index * column_w, band.bottom + int(ctx.gap * 0.6),
                      column_w, max(cm(0.8), frame.bottom - band.bottom - ctx.gap)),
                step.text,
                size_pt=theme.pt("caption") + 0.5, color=theme.muted, align="center",
                valign="top", line_spacing=1.12, padding_cm=0.18,
            )
    return slide.shapes[-1]


def _process_numbered(ctx: RenderContext, slide, frame: Frame, steps):
    theme = ctx.theme
    columns = frame.columns(len(steps), ctx.gap)
    badge = cm(1.05)
    for index, (step, column) in enumerate(zip(steps, columns), start=1):
        circle = new_shape(slide, MSO_SHAPE.OVAL, Frame(column.x, column.y, badge, badge))
        shape_ooxml.solid_fill(circle, theme.accent)
        shape_ooxml.no_outline(circle)
        shape_ooxml.no_shadow(circle)
        text_frame = circle.text_frame
        text_frame.vertical_anchor = anchor("middle")
        text_ooxml.set_insets(text_frame, 0, 0, 0, 0)
        write_paragraph(
            ctx, slide, text_frame.paragraphs[0],
            S.BulletSpec(runs=[S.RunSpec(text=str(index), bold=True)]),
            size_pt=theme.pt("small") + 1, color=theme.on_accent, font=theme.heading_font,
            align="center",
        )
        bullets = [S.BulletSpec(runs=[S.RunSpec(text=step.title, bold=True)])]
        if step.text:
            bullets.append(S.BulletSpec(
                runs=[S.RunSpec(text=step.text, color=theme.muted)], space_before_pt=3
            ))
        top = column.y + badge + int(ctx.gap * 0.5)
        add_textbox(
            ctx, slide, Frame(column.x, top, column.w, max(cm(1.0), column.bottom - top)),
            bullets,
            size_pt=theme.pt("small"), color=theme.text, valign="top", line_spacing=1.12,
        )
    return slide.shapes[-1]


def render_placeholder_element(ctx: RenderContext, slide, element: S.PlaceholderElement,
                               frame: Frame):
    placeholder = find_placeholder(slide, idx=element.idx, name=element.name)
    if placeholder is None:
        raise InvalidSpec(
            f"No placeholder {element.idx if element.idx is not None else element.name!r} on this "
            "slide: use a layout from pptx_create's list instead."
        )
    if element.image:
        return fill_placeholder(ctx, slide, placeholder, {"image": element.image})
    value: Any = element.text if element.text is not None else [
        b.model_dump() if isinstance(b, S.BulletSpec) else b for b in element.bullets
    ]
    return fill_placeholder(ctx, slide, placeholder, value)


def find_placeholder(slide, idx: int | None = None, name: str | None = None):
    for placeholder in slide.placeholders:
        if idx is not None and placeholder.placeholder_format.idx == idx:
            return placeholder
        if name and (placeholder.name or "").strip().lower() == str(name).strip().lower():
            return placeholder
    return None


def fill_placeholder(ctx: RenderContext, slide, placeholder, value):
    """Put a string, a bullet list or an image into a template placeholder."""
    if isinstance(value, dict) and value.get("image"):
        resolved = image_io.resolve_image(value["image"])
        return placeholder.insert_picture(resolved.rewound())
    if isinstance(value, str):
        items = [S.BulletSpec(text=value)]
        bullet_kind = "none"
    else:
        items = normalise_bullets(value or [])
        bullet_kind = "inherit"

    text_frame = placeholder.text_frame
    text_frame.clear()
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.level = min(max(int(item.level or 0), 0), 4)
        runs = item.runs or [S.RunSpec(text=bullet_text(item))]
        for run_spec in runs:
            run = paragraph.add_run()
            run.text = run_spec.text or ""
            # The layout's own text styles decide size and colour; only explicit
            # overrides are applied, so the corporate look survives.
            apply_run_format(run, ctx.theme, run_spec)
            if run_spec.link:
                set_link(ctx, slide, run, run_spec.link)
        if bullet_kind == "none":
            text_ooxml.set_bullet(paragraph, "none")
    return placeholder


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

_RENDERERS = {
    S.TextElement: render_text_element,
    S.ImageElement: render_image_element,
    S.ChartElement: render_chart_element,
    S.TableElement: render_table_element,
    S.ShapeElement: render_shape_element,
    S.KpiElement: render_kpi_element,
    S.TimelineElement: render_timeline_element,
    S.ProcessElement: render_process_element,
    S.PlaceholderElement: render_placeholder_element,
}


def render_element(ctx: RenderContext, slide, element, default_frame: Frame):
    handler = _RENDERERS.get(type(element))
    if handler is None:  # pragma: no cover — the discriminated union prevents this
        raise InvalidSpec(f"No renderer for element type {type(element).__name__}.")
    frame = resolve_frame(ctx, element.frame, default_frame)
    shape = handler(ctx, slide, element, frame)
    if element.name and shape is not None and hasattr(shape, "name"):
        shape.name = element.name
    if element.alt_text and shape is not None:
        try:
            shape_ooxml.set_alt_text(shape, element.alt_text)
        except Exception:  # pragma: no cover — a few shape kinds have no cNvPr
            pass
    return shape


def render_content(ctx: RenderContext, slide, content, frame: Frame, *, title: str | None = None):
    """A column of a two-content slide: prose, bullets, or a full element."""
    theme = ctx.theme
    body = frame
    if title:
        head = frame.top_slice(cm(0.95))
        add_plain_text(
            ctx, slide, head, title,
            size_pt=theme.pt("small") + 1.5, color=theme.text, bold=True,
            font=theme.heading_font, valign="middle", autofit=False,
        )
        add_rule(ctx, slide, Frame(frame.x, head.bottom - cm(0.1), frame.w, 0), theme.accent,
                 thickness_cm=0.06)
        body = frame.remaining_below(head, int(ctx.gap * 0.5))

    if content is None:
        return None
    if isinstance(content, str):
        return add_plain_text(
            ctx, slide, body, content,
            size_pt=None, color=theme.text, valign="top", line_spacing=1.2,
        )
    if isinstance(content, list):
        return add_textbox(
            ctx, slide, body, normalise_bullets(content),
            color=theme.text, valign="top", bullet_kind="char", line_spacing=1.16,
        )
    return render_element(ctx, slide, content, body)
