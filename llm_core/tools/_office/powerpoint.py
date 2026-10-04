# SPDX-License-Identifier: MIT
"""llm_core/tools/_office/powerpoint.py — outils PowerPoint : ``pptx_create``, ``pptx_read``, ``pptx_edit``.

Sans état, comme les outils Word. Une diapo est un objet ``{"layout": …}`` lu
avec tolérance (synonymes français, champs voisins, mise en page déduite des
champs quand elle manque) ; un champ que la mise en page n'utilise pas est
ignoré (signalé). Les graphiques viennent des outils chart_* : un
``!id`` dans ``chart`` (ou dans ``left``/``right``) devient un graphique natif,
une image, un tableau ou des tuiles.

Numérotation : diapos à partir de 1 (comme PowerPoint), formes à partir de 0
(ordre de ``pptx_read``). Dans ``pptx_edit``, ``slide`` désigne la diapo telle
qu'elle était AVANT l'appel ; ``to`` est la position voulue dans le deck au
moment de l'opération.
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple

from lxml import etree
from pptx import Presentation
from pydantic import ValidationError

from .._chart.normalise import Notes, fold
from . import graphiques as G
from ._pptx import deckio
from ._pptx.errors import PptxMcpError
from ._pptx.ooxml import slides as slide_ooxml, theme as theme_ooxml
from ._pptx.render import builder as B, spec as S
from ._pptx.render.elements import RenderContext
from ._pptx.render.layout import Geometry
from ._pptx.render.mdconv import markdown_to_slides
from ._pptx.render.theme import PRESETS, resolve_theme
from .commun import (
    Env,
    OfficeError,
    booleen,
    canon,
    chemin_sortie,
    choisir,
    ecrire,
    entier,
    lire_fichier,
    texte,
)
from .ops import lire_ops, requis
from .paquet import PaquetInvalide, normaliser, retirer_lecteur

LAYOUTS = {
    "title": ["cover", "titre", "title_slide", "couverture", "page_de_titre", "intro", "opening",
              "accueil", "first"],
    "section": ["divider", "separator", "section_header", "intercalaire", "chapitre", "chapter",
                "partie", "transition"],
    "agenda": ["toc", "sommaire", "plan", "outline", "programme", "ordre_du_jour", "summary"],
    "bullets": ["content", "text", "list", "bullet", "bullet_points", "puces", "liste", "texte",
                "points", "contenu", "title_and_content", "bulleted", "body"],
    "two_content": ["two_column", "two_columns", "split", "deux_colonnes", "columns",
                    "side_by_side", "colonnes"],
    "kpi": ["metrics", "dashboard", "chiffres_cles", "indicateurs", "key_figures", "stats",
            "numbers", "kpis", "chiffres", "figures"],
    "chart": ["graph", "graphique", "diagram", "diagramme", "chart_slide", "plot", "courbe"],
    "table": ["tableau", "grid", "data_table", "tabular"],
    "comparison": ["versus", "vs", "pros_cons", "comparaison", "compare",
                   "avantages_inconvenients", "before_after", "avant_apres"],
    "timeline": ["roadmap", "feuille_de_route", "jalons", "milestones", "chronologie", "planning",
                 "calendrier"],
    "process": ["steps", "etapes", "processus", "workflow", "flow", "demarche", "procedure"],
    "image": ["picture", "visual", "photo", "illustration", "img"],
    "quote": ["statement", "citation", "quotation", "message", "verbatim"],
    "closing": ["end", "thanks", "merci", "fin", "questions", "q_a", "contact", "closing_slide"],
    "blank": ["free", "custom", "vide", "empty"],
}
CHAMPS_DIAPO = {
    "title": ["titre", "heading", "headline", "slide_title", "header", "titre_diapo"],
    "subtitle": ["sous_titre", "subheading", "tagline", "baseline", "sub_title"],
    "bullets": ["points", "puces", "bullet_points", "lines", "list", "liste", "key_points",
                "contenu_liste"],
    "text": ["paragraph", "prose", "texte", "body_text", "description"],
    "notes": ["speaker_notes", "notes_orateur", "commentaire", "presenter_notes", "script",
              "remarks"],
    "takeaway": ["key_message", "conclusion", "a_retenir", "so_what", "insight", "message_cle",
                 "key_takeaway"],
    "chart": ["graph", "graphique", "figure", "chart_ref", "ref", "chart_id", "visual_ref"],
    "items": ["kpis", "metrics", "indicators", "indicateurs", "chiffres", "values", "agenda",
              "entries", "elements"],
    "milestones": ["jalons", "events", "evenements", "dates", "phases_dates"],
    "steps": ["etapes", "stages", "phases"],
    "table": ["tableau", "rows", "data", "donnees", "grid"],
    "left": ["gauche", "col1", "column1", "left_column", "before", "avant", "option_a", "pros",
             "avantages"],
    "right": ["droite", "col2", "column2", "right_column", "after", "apres", "option_b", "cons",
              "inconvenients"],
    "image": ["picture", "photo", "img", "src", "image_path", "illustration"],
    "author": ["auteur", "by", "speaker", "orateur", "presenter"],
    "role": ["fonction", "position", "job_title"],
    "contact": ["contacts", "email", "emails", "coordonnees"],
    "date": ["le", "when"],
    "caption": ["legende", "source"],
    "kicker": ["eyebrow", "surtitre", "section_label"],
}
ALIAS_CREATE = {
    "path": ["file", "filename", "file_name", "file_path", "filepath", "output", "output_path",
             "fichier", "chemin", "name", "nom", "save_as", "destination", "dest", "target"],
    "slides": ["diapos", "diapositives", "pages", "deck", "content", "contenu", "slide_list"],
    "markdown": ["outline", "md", "plan", "text", "texte", "slides_markdown"],
    "title": ["titre", "deck_title", "presentation_title"],
    "author": ["auteur", "by", "presenter", "orateur"],
    "theme": ["style", "design", "preset", "look", "colors", "couleurs", "palette"],
    "template": ["modele", "base", "template_path", "from_template", "model", "gabarit", "master"],
    "size": ["format", "ratio", "aspect", "aspect_ratio", "slide_size", "taille"],
    "footer": ["pied", "pied_de_page", "footer_text"],
    "slide_numbers": ["numbering", "page_numbers", "numeros", "numerotation", "numbers"],
}
ALIAS_READ = {
    "path": ALIAS_CREATE["path"],
    "slide": ["slide_number", "diapo", "page", "n", "index", "slide_index", "numero"],
    "find": ["search", "query", "pattern", "chercher", "rechercher", "q", "grep"],
}
ALIAS_EDIT = {
    "path": [a for a in ALIAS_CREATE["path"] if a not in ("output", "output_path", "save_as",
                                                           "dest")],
    "ops": ["operations", "edits", "changes", "actions", "modifications", "steps"],
    "save_as": ["output", "output_path", "new_path", "copy_to", "enregistrer_sous", "dest"],
}
OPS = {
    "replace": ["replace_text", "substitute", "remplacer", "find_replace", "search_replace",
                "find_and_replace"],
    "fill": ["fill_fields", "fill_placeholders", "merge", "remplir", "fields", "variables"],
    "set_text": ["set", "rewrite", "update", "edit", "modifier", "set_shape_text", "change"],
    "add_slides": ["add", "add_slide", "insert", "insert_slides", "insert_slide", "append",
                   "append_slides", "ajouter", "new_slide", "new_slides"],
    "delete_slide": ["delete", "remove", "remove_slide", "delete_slides", "supprimer", "del"],
    "move_slide": ["move", "reorder", "deplacer", "move_to"],
    "duplicate_slide": ["duplicate", "copy", "copy_slide", "dupliquer", "clone"],
    "notes": ["set_notes", "speaker_notes", "notes_orateur", "add_notes"],
    "meta": ["metadata", "properties", "set_metadata", "proprietes"],
}
_S = ["slide_number", "diapo", "slide_index", "n", "index", "page", "numero"]
CHAMPS_OPS = {
    "replace": {"find": ["search", "old", "from", "pattern", "text", "chercher", "old_text"],
                "replace": ["replacement", "new", "with", "to", "by", "value", "new_text"],
                "match_case": ["case_sensitive", "case"], "regex": ["is_regex", "use_regex"]},
    "fill": {"values": ["fields", "data", "valeurs", "placeholders", "variables", "context",
                        "mapping"]},
    "set_text": {"slide": _S, "shape": ["shape_index", "forme", "shape_id", "element", "box"],
                 "text": ["new_text", "value", "texte", "content", "bullets"]},
    "add_slides": {"slides": ["slide", "diapos", "content", "new_slides"],
                   "markdown": ["md", "outline", "text"],
                   "after": ["after_slide", "apres", "position", "at", "where"],
                   "before": ["avant", "before_slide"]},
    "delete_slide": {"slide": _S + ["slides"]},
    "move_slide": {"slide": _S + ["from"], "to": ["position", "new_position", "vers", "target"]},
    "duplicate_slide": {"slide": _S, "to": ["position", "after", "vers"]},
    "notes": {"slide": _S, "text": ["notes", "texte", "content", "value"]},
    "meta": {"title": ["titre"], "author": ["auteur"], "subject": ["sujet"],
             "keywords": ["mots_cles", "tags"]},
}
EXEMPLE_SLIDES = [
    {"layout": "title", "title": "Comité stratégique", "subtitle": "Revue du T4 2026"},
    {"layout": "bullets", "title": "Points clés", "bullets": ["Croissance de 12 %", "Marge stable"]},
    {"layout": "chart", "title": "Ventes par trimestre", "chart": "!a1b2c3d4e5f6",
     "takeaway": "Le T4 tire l'année."},
    {"layout": "kpi", "title": "Chiffres clés",
     "items": [{"value": "18,4 M€", "label": "Chiffre d'affaires", "delta": "+12 %"}]},
]
EXEMPLE_OPS = [
    {"op": "replace", "find": "T3", "replace": "T4"},
    {"op": "set_text", "slide": 2, "shape": 1, "text": "Nouveau texte"},
    {"op": "add_slides", "after": 3, "slides": [EXEMPLE_SLIDES[1]]},
    {"op": "delete_slide", "slide": 5},
    {"op": "notes", "slide": 1, "text": "Insister sur la régularité."},
]


def _deduire(f: Dict[str, Any]) -> Optional[str]:
    if "find" in f or "search" in f:
        return "replace"
    if "values" in f or "fields" in f:
        return "fill"
    if "slides" in f or "markdown" in f:
        return "add_slides"
    if "shape" in f and "text" in f:
        return "set_text"
    if "notes" in f:
        return "notes"
    if "to" in f and "slide" in f:
        return "move_slide"
    return None


# ── Normalisation d'une diapo ───────────────────────────────────────────────
def _puces(v: Any) -> List[Any]:
    """Puces tolérantes : liste, texte multi-lignes (« - », indentation), objets."""
    if v is None:
        return []
    if isinstance(v, str):
        lignes = [ln for ln in v.replace("\r\n", "\n").split("\n") if ln.strip()]
        puces: List[Any] = []
        for ln in lignes:
            m = re.match(r"^(\s*)(?:[-*•+]|\d+[.)])?\s*(.*)$", ln)
            niveau = min(4, len(m.group(1).expandtabs(2)) // 2) if m else 0
            t = (m.group(2) if m else ln).strip()
            if t:
                puces.append({"text": t, "level": niveau} if niveau else t)
        return puces
    if isinstance(v, dict):
        v = [v]
    out: List[Any] = []
    for x in v if isinstance(v, list) else [v]:
        if isinstance(x, str):
            out.extend(_puces(x) or [])          # « - point », « ··sous-point »
            continue
        if isinstance(x, list):
            out.extend({"text": y if isinstance(y, str) else str(y.get("text", "")), "level": 1}
                       for y in x)
        elif isinstance(x, dict):
            d = canon(x, {"text": ["texte", "label", "content", "title"],
                          "level": ["niveau", "indent", "depth"]}, Notes())
            out.append({k: d[k] for k in ("text", "level", "bold", "italic", "color") if k in d})
        elif x is not None:
            out.append(str(x))
    return out


def _kpi(x: Any, notes: Notes) -> Dict[str, Any]:
    if not isinstance(x, dict):
        return {"value": str(x)}
    d = canon(x, {"value": ["valeur", "number", "nombre", "figure", "chiffre", "amount", "montant"],
                  "label": ["libelle", "name", "title", "titre", "indicateur", "metric", "nom"],
                  "delta": ["evolution", "change", "variation", "diff", "ecart", "progression"],
                  "trend": ["tendance", "direction", "sens"],
                  "caption": ["legende", "period", "periode", "context", "sub"],
                  "good_is": ["better", "mieux"]}, notes, "kpi.")
    v = d.get("value")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        unit = str(d.pop("unit", "") or "")
        d["value"] = G.cellule(float(v)) + (f" {unit}" if unit else "")
    elif v is None:
        d["value"] = "—"
    else:
        d["value"] = str(v)
    if isinstance(d.get("delta"), (int, float)) and not isinstance(d["delta"], bool):
        dv = float(d["delta"])
        d["delta"] = ("+" if dv >= 0 else "−") + G.cellule(abs(dv))
    elif d.get("delta") is not None and not isinstance(d["delta"], str):
        d["delta"] = str(d["delta"])
    if d.get("trend") is None and isinstance(d.get("delta"), str):
        s = d["delta"].strip()
        d["trend"] = "up" if s.startswith("+") else ("down" if s[:1] in "-−" else None)
    if d.get("trend") is not None:
        d["trend"] = {"hausse": "up", "baisse": "down", "stable": "flat", "up": "up",
                      "down": "down", "flat": "flat"}.get(fold(d["trend"]), None)
    return {k: d[k] for k in ("value", "label", "caption", "delta", "trend", "good_is", "color")
            if d.get(k) is not None}


def _jalon(x: Any, notes: Notes) -> Dict[str, Any]:
    if isinstance(x, str):
        m = re.match(r"^\s*([^:–—-]{1,24}?)\s*[:–—-]\s*(.+)$", x)
        return {"date": m.group(1), "title": m.group(2)} if m else {"title": x}
    d = canon(x if isinstance(x, dict) else {}, {
        "date": ["when", "quand", "period", "periode", "label_date", "time", "mois", "trimestre"],
        "title": ["titre", "label", "name", "event", "milestone", "jalon", "libelle"],
        "text": ["description", "detail", "details", "texte"],
        "done": ["completed", "fait", "termine", "achieved", "past"]}, notes, "milestone.")
    out = {k: d[k] for k in ("date", "title", "text") if d.get(k) is not None}
    out["title"] = str(out.get("title", ""))
    if "done" in d:
        out["done"] = booleen(d["done"])
    return out


def _etape(x: Any, notes: Notes) -> Dict[str, Any]:
    if isinstance(x, str):
        return {"title": x}
    d = canon(x if isinstance(x, dict) else {}, {
        "title": ["titre", "label", "name", "step", "etape", "libelle"],
        "text": ["description", "detail", "texte"]}, notes, "step.")
    return {"title": str(d.get("title", "")), **({"text": str(d["text"])} if d.get("text") else {})}


def _colonne(x: Any, notes: Notes) -> Dict[str, Any]:
    if isinstance(x, (list, str)):
        return {"items": _puces(x)}
    d = canon(x if isinstance(x, dict) else {}, {
        "title": ["titre", "heading", "label", "name"],
        "items": ["bullets", "points", "puces", "list", "content", "lines"],
        "icon": ["icone", "symbol"]}, notes, "column.")
    out: Dict[str, Any] = {"items": _puces(d.get("items"))}
    if d.get("title"):
        out["title"] = str(d["title"])
    if d.get("icon"):
        out["icon"] = str(d["icon"])[:2]
    return out


def _table_element(v: Any, notes: Notes) -> Optional[Dict[str, Any]]:
    """Tableau : {header, rows}, liste de lignes-objets, liste de listes, ou Markdown."""
    if isinstance(v, dict):
        d = canon(v, {"header": ["headers", "columns", "colonnes", "entetes"],
                      "rows": ["lignes", "data", "body"]}, notes, "table.")
        rows = d.get("rows") or []
        if rows and all(isinstance(r, dict) for r in rows):
            return _table_element(rows, notes)
        return {"type": "table", "header": [str(h) for h in d.get("header") or []] or None,
                "rows": [[G.cellule(c) for c in r] for r in rows if isinstance(r, list)]}
    if isinstance(v, str) and "|" in v:
        lignes = [ln.strip().strip("|") for ln in v.strip().split("\n") if ln.strip()]
        lignes = [ln for ln in lignes if not re.fullmatch(r"[\s:|-]+", ln)]
        cells = [[c.strip() for c in ln.split("|")] for ln in lignes]
        return {"type": "table", "header": cells[0], "rows": cells[1:]} if cells else None
    if isinstance(v, list) and v:
        if all(isinstance(r, dict) for r in v):
            cles = list(dict.fromkeys(k for r in v for k in r))
            return {"type": "table", "header": cles,
                    "rows": [[G.cellule(r.get(k)) for k in cles] for r in v]}
        if all(isinstance(r, list) for r in v):
            return {"type": "table", "header": [str(c) for c in v[0]],
                    "rows": [[G.cellule(c) for c in r] for r in v[1:]]}
    return None


def _element_graphique(r: G.Rendu) -> Dict[str, Any]:
    if r.mode == "natif":
        n = r.natif
        return {"type": "chart", "chart_type": n["chart_type"],
                "categories": n.get("categories") or [], "series": n["series"],
                "title": None, "x_title": n.get("x_title"), "y_title": n.get("y_title"),
                "number_format": n.get("number_format"), "legend": n.get("legend", "auto"),
                **({"y_max": n["y_max"]} if n.get("y_max") is not None else {}),
                **({"smooth": True} if n.get("smooth") else {})}
    if r.mode == "image":
        return {"type": "image", "source": G.data_uri(r.png), "fit": "contain",
                "alt_text": r.titre or r.kind}
    if r.mode == "kpi":
        return {"type": "kpi", "items": [_kpi({"value": t.get("value"), "label": t.get("label"),
                                               "delta": t.get("delta"),
                                               "trend": t.get("sign")}, Notes())
                                         for t in r.tuiles]}
    return {"type": "table", "header": r.colonnes,
            "rows": [[G.cellule(c) for c in ln] for ln in r.lignes[:40]]}


def _refs_diapo(d: Dict[str, Any]) -> List[str]:
    """Graphiques cités par une diapo : champs de graphique APRÈS lecture des
    synonymes (même lecture que ``normaliser_diapo``, l'identifiant nu y est
    admis) et toute référence ``!id`` ailleurs."""
    c = canon(d, CHAMPS_DIAPO, Notes())
    out = [r for r in (G.ref_de(c.get(k)) for k in ("chart", "table", "left", "right")) if r]
    return out + G.refs_dans(d)


def normaliser_diapo(i: int, brut: Any, rendus: Dict[str, G.Rendu], env: Env) -> Dict[str, Any]:
    notes = env.notes
    if isinstance(brut, str):
        brut = {"title": brut}
        notes.fix(f"slides[{i}] given as text: read as a title")
    if not isinstance(brut, dict):
        raise OfficeError(f"slides[{i}] is not an object", fix='each slide is {"layout": ..., ...}',
                          example=EXEMPLE_SLIDES, code="invalid_slide")
    cle_layout = next((k for k in brut if fold(k) in ("layout", "type", "kind", "template",
                                                      "mise_en_page", "slide_type", "style")), None)
    d = canon({k: v for k, v in brut.items() if k != cle_layout}, CHAMPS_DIAPO, notes,
              f"slides[{i}].")
    layout = choisir(brut.get(cle_layout), LAYOUTS, notes, f"slides[{i}].layout") \
        if cle_layout else None
    if cle_layout and layout is None:
        notes.warn(f"slides[{i}]: unknown layout '{brut.get(cle_layout)}', chosen from the fields")
    if layout is None:
        layout = _deviner(i, d)
        notes.fix(f"slides[{i}]: layout '{layout}' chosen from its fields")
    out: Dict[str, Any] = {"layout": layout}
    for k in ("title", "subtitle", "notes", "takeaway", "kicker", "author", "date", "role",
              "caption"):
        if d.get(k) not in (None, "", []):
            out[k] = texte(d[k]) if k != "notes" else texte(d[k])
    # Graphique cité : il décide de la forme réelle de la diapo.
    ref = G.ref_de(d.get("chart")) or (G.ref_de(d.get("table")) if layout == "table" else None)
    if ref:
        r = G.rendu(rendus, ref)
        el = _element_graphique(r)
        if layout not in ("chart", "kpi", "table", "two_content"):
            notes.fix(f"slides[{i}]: a chart reference makes it a 'chart' slide")
            layout = out["layout"] = "chart"
        if layout == "table" or (r.mode == "table" and layout == "chart"):
            out["layout"] = "table"
            out["table"] = el if el["type"] == "table" else _table_element(
                [dict(zip(r.colonnes, ln)) for ln in r.lignes], notes)
        elif layout == "kpi" and el["type"] == "kpi":
            out["items"] = el["items"]
        else:
            out["layout"] = "chart"
            out["chart"] = el
        if not out.get("title") and r.titre:
            out["title"] = r.titre
    lay = out["layout"]
    if lay in ("bullets",):
        if d.get("bullets") is not None:
            out["bullets"] = _puces(d["bullets"])
        elif isinstance(d.get("items"), (list, str)):
            out["bullets"] = _puces(d["items"])
        if d.get("text") and not out.get("bullets"):
            t = texte(d["text"])
            if re.search(r"^\s*[-*•]\s", t, re.M):
                out["bullets"] = _puces(t)
            else:
                out["text"] = t
        if d.get("image"):
            out["image"] = str(d["image"])
        if not out.get("bullets") and not out.get("text"):
            raise OfficeError(f"slides[{i}] (bullets): no bullets or text", code="invalid_slide",
                              fix="add bullets = [\"point 1\", \"point 2\"] or text",
                              example=[EXEMPLE_SLIDES[1]])
    elif lay == "agenda":
        out["items"] = _puces(d.get("items") if d.get("items") is not None else d.get("bullets"))
        out.setdefault("title", "Sommaire")
    elif lay == "kpi" and "items" not in out:
        items = d.get("items") or d.get("bullets")
        if not isinstance(items, list) or not items:
            raise OfficeError(f"slides[{i}] (kpi): items is required", code="invalid_slide",
                              fix='items = [{"value": "18,4 M€", "label": "CA", "delta": "+12 %"}]',
                              example=[EXEMPLE_SLIDES[3]])
        out["items"] = [_kpi(x, notes) for x in items[:8]]
    elif lay == "chart" and "chart" not in out:
        raise OfficeError(f"slides[{i}] (chart): chart is required",
                          fix="create the chart with a chart_<type> tool, then put its ref here: "
                              "chart = \"!a1b2c3d4e5f6\"", example=[EXEMPLE_SLIDES[2]],
                          code="invalid_slide")
    elif lay == "table" and "table" not in out:
        tab = _table_element(d.get("table") if d.get("table") is not None else d.get("items"), notes)
        if not tab:
            raise OfficeError(f"slides[{i}] (table): table is required", code="invalid_slide",
                              fix='table = {"header": ["A", "B"], "rows": [["1", "2"]]}, a list '
                                  'of row objects, or a chart_table ref "!id"')
        out["table"] = tab
    elif lay == "comparison":
        out["left"] = _colonne(d.get("left"), notes)
        out["right"] = _colonne(d.get("right"), notes)
    elif lay == "two_content":
        for cote in ("left", "right"):
            v = d.get(cote)
            ref_cote = G.ref_de(v)
            if ref_cote:
                out[cote] = _element_graphique(G.rendu(rendus, ref_cote))
            elif isinstance(v, dict) and fold(v.get("type", "")) in ("chart", "image", "table", "kpi"):
                out[cote] = v
            elif v is not None:
                out[cote] = _puces(v) if not (isinstance(v, str) and "\n" not in v and
                                               not v.lstrip().startswith(("-", "*", "•"))) else v
        for cote in ("left_title", "right_title"):
            if brut.get(cote):
                out[cote] = str(brut[cote])
    elif lay == "timeline":
        src = d.get("milestones") or d.get("items") or d.get("steps") or d.get("bullets") or []
        out["milestones"] = [_jalon(x, notes) for x in (src if isinstance(src, list) else [src])]
    elif lay == "process":
        src = d.get("steps") or d.get("items") or d.get("bullets") or []
        out["steps"] = [_etape(x, notes) for x in (src if isinstance(src, list) else _puces(src))]
    elif lay == "image":
        img = d.get("image")
        if not img:
            raise OfficeError(f"slides[{i}] (image): image is required (a sandbox path)",
                              code="invalid_slide", fix="image = 'images/photo.png'")
        out["image"] = str(img)
        out["fit"] = "contain"
        if d.get("text"):
            out["text"] = texte(d["text"])
    elif lay == "quote":
        out["text"] = texte(d.get("text") or d.get("title") or d.get("bullets") or "")
        out.pop("title", None)
    elif lay == "closing":
        c = d.get("contact")
        if c:
            out["contact"] = [str(x) for x in (c if isinstance(c, list) else [c])]
        out.setdefault("title", "Merci")
    elif lay == "title":
        out.setdefault("title", "")
    elif lay == "section":
        out.setdefault("title", "")
    # Champs que la mise en page n'a pas : ignorés (le moteur les refuserait).
    modele = _MODELES.get(out["layout"])
    if modele is not None:
        for k in [k for k in out if k not in modele.model_fields]:
            out.pop(k)
            notes.fix(f"slides[{i}]: '{k}' is not used by a '{out['layout']}' slide: ignored")
    return out


_MODELES = {"title": S.TitleSlide, "section": S.SectionSlide, "agenda": S.AgendaSlide,
            "bullets": S.BulletsSlide, "two_content": S.TwoContentSlide, "kpi": S.KpiSlide,
            "chart": S.ChartSlide, "table": S.TableSlide, "comparison": S.ComparisonSlide,
            "timeline": S.TimelineSlide, "process": S.ProcessSlide, "image": S.ImageSlide,
            "quote": S.QuoteSlide, "closing": S.ClosingSlide, "blank": S.BlankSlide}


def _deviner(i: int, d: Dict[str, Any]) -> str:
    if d.get("chart"):
        return "chart"
    items = d.get("items")
    if isinstance(items, list) and items and all(isinstance(x, dict) for x in items) and \
            any(fold(k) in ("value", "valeur", "number", "chiffre") for x in items for k in x):
        return "kpi"
    if d.get("milestones"):
        return "timeline"
    if d.get("steps"):
        return "process"
    if d.get("table"):
        return "table"
    if d.get("left") is not None and d.get("right") is not None:
        return "comparison" if all(isinstance(d[c], dict) for c in ("left", "right")) else "two_content"
    if d.get("bullets") or d.get("text") or d.get("items"):
        return "bullets"
    if d.get("image"):
        return "image"
    if d.get("author") and not d.get("title"):
        return "quote"
    return "title" if i == 0 else "section"


# ── Ouverture / écriture ────────────────────────────────────────────────────
def _ouvrir(env: Env, path: Any) -> Tuple[str, Any, str]:
    """→ (chemin, présentation, empreinte des octets lus : verrou de l'écriture)."""
    import hashlib
    rel, brut = lire_fichier(env, path, "pptx", "PowerPoint deck")
    sha = hashlib.sha256(brut).hexdigest()
    try:
        blob, notes = normaliser(brut, "pptx", f"file '{rel}'")
    except PaquetInvalide as e:
        raise OfficeError(str(e), code="bad_file",
                          fix="give the path of a .pptx file (PowerPoint 2007 or later)")
    for n in notes:
        env.notes.fix(n)
    try:
        return rel, Presentation(io.BytesIO(blob)), sha
    except Exception as e:                       # noqa: BLE001 — fichier corrompu
        raise OfficeError(f"'{rel}' could not be opened as a PowerPoint deck: {e}",
                          code="bad_file", fix="check that the file is a valid .pptx")


def _enregistrer(env: Env, prs: Any, rel: str, attendu: Optional[str] = None
                 ) -> Dict[str, Any]:
    sortie = io.BytesIO()
    prs.save(sortie)
    return ecrire(env, rel, sortie.getvalue(), attendu)


def _moteur(fn, *a, **k):
    try:
        return fn(*a, **k)
    except OfficeError:
        raise
    except ValidationError as e:
        errs = e.errors()
        err = errs[0] if errs else {}
        lieu = ".".join(str(x) for x in err.get("loc", ()))
        raise OfficeError(f"invalid slide content at {lieu or 'root'}: {err.get('msg', e)}",
                          code="invalid_slide", example=EXEMPLE_SLIDES)
    except (PptxMcpError, PaquetInvalide) as e:
        raise OfficeError(str(e), code="invalid_slide")
    except re.error as e:
        raise OfficeError(f"invalid regular expression: {e}", code="invalid_op",
                          fix="fix the pattern, or set regex=false to search the text literally")
    except (etree.XMLSyntaxError, zipfile.BadZipFile) as e:
        raise OfficeError(f"the file or template is damaged: {e}", code="bad_file",
                          fix="open and re-save it with PowerPoint or LibreOffice, or use "
                              "another file")


def _lecteur(env: Env):
    from .word import _avec_lecteur
    return _avec_lecteur(env)


def _diapos(env: Env, slides: Any, markdown: str) -> Tuple[List[Dict[str, Any]], Dict[str, G.Rendu]]:
    if isinstance(slides, str) and not markdown:
        markdown, slides = slides, None
        env.notes.fix("slides given as text: read as a Markdown outline")
    if isinstance(slides, dict):
        slides = [slides]
    if markdown and not slides:
        diapos = _moteur(markdown_to_slides, markdown)
        rendus = G.preparer(env, G.REF.findall(markdown), titre_image=False)
        return [_md_refs(s, rendus) for s in diapos], rendus
    if not isinstance(slides, list) or not slides:
        raise OfficeError("slides is empty", code="empty_content",
                          fix="slides = a list of slide objects, each with a layout",
                          example=EXEMPLE_SLIDES)
    ids: List[str] = []
    for s in slides:
        if isinstance(s, dict):
            ids.extend(_refs_diapo(s))
    rendus = G.preparer(env, ids, titre_image=False)
    return [normaliser_diapo(i, s, rendus, env) for i, s in enumerate(slides)], rendus


def _md_refs(s: Dict[str, Any], rendus: Dict[str, G.Rendu]) -> Dict[str, Any]:
    """Plan Markdown : une puce « !id » devient le graphique de la diapo."""
    puces = s.get("bullets") or []
    for p in list(puces):
        r = G.ref_de(p if isinstance(p, str) else (p or {}).get("text"))
        if r:
            puces.remove(p)
            s = {**s, "layout": "chart", "chart": _element_graphique(rendus[r])}
            s.pop("bullets", None) if not puces else None
            break
    return s


# ── pptx_create ─────────────────────────────────────────────────────────────
def pptx_create(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_CREATE, env.notes)
    rel = chemin_sortie(a.get("path"), "pptx", env.notes)
    jeton = _lecteur(env)
    try:
        diapos, rendus = _diapos(env, a.get("slides"), texte(a.get("markdown")))
        pied = a.get("footer")
        spec = _moteur(S.DeckSpec, filename=rel.rsplit("/", 1)[-1],
                       meta=S.MetaSpec(title=a.get("title") or None, author=a.get("author") or None),
                       theme=_theme(a.get("theme"), env),
                       size=_taille(a.get("size"), env),
                       template=str(a["template"]) if a.get("template") else None,
                       footer=S.FooterSpec(text=str(pied) if pied else None,
                                           slide_numbers=booleen(a.get("slide_numbers"), True)),
                       slides=diapos)
        prs, ctx = _moteur(B.build_deck, spec)
    finally:
        retirer_lecteur(jeton)
    for n in ctx.notes:
        env.notes.warn(n)
    res = _enregistrer(env, prs, rel)
    plan = _plan(prs)
    out = {**res, "slides": len(plan), "outline": plan}
    if rendus:
        out["charts"] = [{"ref": f"!{r.id}", "inserted_as": {"natif": "native chart",
                                                             "image": "image", "table": "table",
                                                             "kpi": "kpi tiles"}[r.mode]}
                         for r in rendus.values()]
    out["summary"] = (f"{'Replaced' if res['old_sha256'] else 'Created'} {res['path']}: "
                      f"{len(plan)} slides ({', '.join(dict.fromkeys(d['layout'] for d in diapos))})")
    return out


def _theme(v: Any, env: Env) -> S.ThemeSpec:
    if v in (None, "", {}):
        return S.ThemeSpec()
    if isinstance(v, str):
        if re.fullmatch(r"#?[0-9a-fA-F]{6}", v.strip()):
            return S.ThemeSpec(accent=v.strip())
        p = choisir(v, {k: [] for k in PRESETS}, env.notes, "theme")
        if p:
            return S.ThemeSpec(preset=p)
        env.notes.warn(f"theme '{v}' unknown: one of {', '.join(PRESETS)}, or a colour '#1F4E79'")
        return S.ThemeSpec()
    if isinstance(v, dict):
        t = canon(v, {"preset": ["name", "theme", "nom", "style"],
                      "accent": ["color", "colour", "couleur", "primary", "brand", "main_color"],
                      "body_font": ["font", "police", "font_family"],
                      "heading_font": ["title_font", "police_titres"]}, env.notes, "theme.")
        out: Dict[str, Any] = {}
        if t.get("preset"):
            p = choisir(t["preset"], {k: [] for k in PRESETS}, env.notes, "theme.preset")
            if p:
                out["preset"] = p
        for k in ("accent", "body_font", "heading_font"):
            if t.get(k):
                out[k] = str(t[k])
        return _moteur(S.ThemeSpec, **out)
    return S.ThemeSpec()


def _taille(v: Any, env: Env) -> str:
    if not v:
        return "16:9"
    f = fold(v)
    for cle, al in {"16:9": ["16_9", "169", "wide", "widescreen", "large"],
                    "4:3": ["4_3", "43", "standard", "classic"],
                    "16:10": ["16_10", "1610"], "a4": ["a4", "a4_portrait"],
                    "a4_landscape": ["a4_landscape", "a4_paysage"]}.items():
        if f in al or f == fold(cle):
            return cle
    env.notes.warn(f"size '{v}' unknown: 16:9 used")
    return "16:9"


def _plan(prs: Any) -> List[Dict[str, Any]]:
    out = []
    for k, slide in enumerate(prs.slides):
        out.append({"slide": k + 1, "title": (deckio.slide_title(slide) or "")[:100]})
    return out


# ── pptx_read ───────────────────────────────────────────────────────────────
def pptx_read(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_READ, env.notes)
    rel, prs, _ = _ouvrir(env, a.get("path"))
    n = len(prs.slides)
    voulue = entier(a.get("slide"))
    if voulue is not None and not 1 <= voulue <= n:
        raise OfficeError(f"slide {voulue} does not exist (the deck has slides 1 to {n})",
                          code="bad_index")
    diapos = []
    for k, slide in enumerate(prs.slides):
        if voulue is not None and k + 1 != voulue:
            continue
        d: Dict[str, Any] = {"slide": k + 1, "layout": slide.slide_layout.name,
                             "title": deckio.slide_title(slide)}
        formes = []
        for j, shape in enumerate(slide.shapes):
            if _numero_diapo(shape):
                continue
            f: Dict[str, Any] = {"shape": j, "kind": deckio.shape_kind(shape)}
            if shape.has_text_frame and (shape.text_frame.text or "").strip():
                f["text"] = shape.text_frame.text.strip()[:1500]
            if getattr(shape, "has_table", False):
                f["rows"] = [[c.text.strip()[:120] for c in row.cells] for row in shape.table.rows][:20]
            if getattr(shape, "has_chart", False):
                f["series"] = [s.name for s in shape.chart.series]
            if f.get("text") or f["kind"] in ("chart", "table", "picture"):
                formes.append(f)
        d["shapes"] = formes
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                d["notes"] = notes[:1500]
        diapos.append(d)
    res: Dict[str, Any] = {"path": env.espace.afficher(rel), "slides_total": n,
                           "size": f"{prs.slide_width / 360000:.1f} x {prs.slide_height / 360000:.1f} cm",
                           "slides": diapos}
    tout = " ".join(f.get("text", "") for d in diapos for f in d["shapes"])
    champs = sorted(set(m.strip() for m in re.findall(r"\{\{\s*([^{}]{1,60}?)\s*\}\}", tout)))
    if champs:
        res["placeholders"] = champs
    if a.get("find"):
        motif = re.compile(re.escape(str(a["find"])), re.I)
        res["matches"] = [{"slide": d["slide"], "shape": f["shape"], "text": f.get("text", "")[:160]}
                          for d in diapos for f in d["shapes"] if motif.search(f.get("text", ""))][:50]
        if not res["matches"]:
            res["matches_note"] = f"'{a['find']}' not found (case ignored)"
    return res


def _numero_diapo(shape: Any) -> bool:
    """Numéro de diapo (champ slidenum) : bruit pour le modèle, non listé."""
    el = getattr(shape, "_element", None)
    if el is None:
        return False
    return any(f.get("type") == "slidenum" for f in el.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}fld"))


# ── pptx_edit ───────────────────────────────────────────────────────────────
def _contexte(prs: Any) -> RenderContext:
    """Contexte de rendu d'un deck existant : thème lu dans sa partie theme."""
    lu: Dict[str, Any] = theme_ooxml.read_theme(prs) or {}
    couleurs = lu.get("colors") or {}
    geo = Geometry.from_presentation(prs)
    spec = S.ThemeSpec(accent=couleurs.get("accent1") or None,
                       accent2=couleurs.get("accent2") or None,
                       heading_font=lu.get("major_font") or None,
                       body_font=lu.get("minor_font") or None)
    return RenderContext(presentation=prs, theme=resolve_theme(spec, slide_height_cm=geo.height_cm),
                         geometry=geo, footer=S.FooterSpec(slide_numbers=True))


def pptx_edit(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_EDIT, env.notes)
    ops = lire_ops(a.get("ops"), OPS, CHAMPS_OPS, _deduire, env.notes, EXEMPLE_OPS)
    rel, prs, sha_lu = _ouvrir(env, a.get("path"))
    cible = rel
    if a.get("save_as"):
        cible = chemin_sortie(a["save_as"], "pptx", env.notes)
    elif not rel.lower().endswith(".pptx"):
        cible = chemin_sortie(rel, "pptx", env.notes)
        env.notes.fix(f"a {rel.rsplit('.', 1)[-1]} file is saved as a new .pptx: {cible}")
    # Diapos d'origine, figées : les numéros de l'appel ne glissent pas.
    origine = list(prs.slides)
    supprimees: set = set()
    bilan: List[Dict[str, Any]] = []
    jeton = _lecteur(env)
    ctx: Optional[RenderContext] = None
    change = False
    try:
        for i, (nom, c) in enumerate(ops):
            if nom == "add_slides" and ctx is None:
                ctx = _contexte(prs)
            fait = _appliquer(i, nom, c, prs, origine, supprimees, ctx, env)
            bilan.append({"op": nom, **fait})
            change = change or fait.get("changed", True)
    finally:
        retirer_lecteur(jeton)
    if not change:
        raise OfficeError("no operation changed the deck: nothing was saved",
                          code="nothing_changed",
                          fix="check the exact text, slide and shape numbers with pptx_read, then "
                              "retry with matching values", example=EXEMPLE_OPS[:2])
    if ctx is not None:
        for n_ in ctx.notes:
            env.notes.warn(n_)
    # Même fichier : écrit seulement s'il n'a pas changé depuis la lecture.
    res = _enregistrer(env, prs, cible, sha_lu if cible == rel else None)
    plan = _plan(prs)
    faits = [o["op"] + ("" if o.get("changed", True) else " (no change)") for o in bilan]
    return {**res, "applied": bilan, "slides": len(plan), "outline": plan,
            "summary": f"Saved {res['path']}: {len(ops)} operation(s) ({', '.join(faits)}), "
                       f"{len(plan)} slides"}


def _position(prs: Any, slide: Any) -> int:
    return list(prs.slides).index(slide)


def _diapo(i: int, nom: str, v: Any, prs, origine, supprimees) -> Any:
    k = entier(v)
    if k is None or not 1 <= k <= len(origine):
        raise OfficeError(f"ops[{i}] ({nom}): slide {v!r} does not exist (slides 1 to "
                          f"{len(origine)}, numbered from 1)", code="bad_index",
                          fix="take slide numbers from pptx_read")
    s = origine[k - 1]
    if id(s) in supprimees:
        raise OfficeError(f"ops[{i}] ({nom}): slide {k} was deleted by an earlier operation",
                          code="bad_index")
    return s


def _appliquer(i: int, nom: str, c: Dict[str, Any], prs, origine, supprimees,
               ctx: Optional[RenderContext], env: Env) -> Dict[str, Any]:
    if nom == "replace":
        cherche = str(requis(i, nom, c, "find", EXEMPLE_OPS[0]))
        nb = _moteur(deckio.replace_text, prs, cherche, texte(c.get("replace")),
                     regex=booleen(c.get("regex")), ignore_case=not booleen(c.get("match_case"), True))
        if not nb:
            env.notes.warn(f"ops[{i}] replace: '{cherche}' not found")
        return {"replaced": nb, "changed": bool(nb)}
    if nom == "fill":
        valeurs = requis(i, nom, c, "values", {"op": "fill", "values": {"client": "ACME"}})
        if not isinstance(valeurs, dict):
            raise OfficeError(f"ops[{i}] (fill): values must be an object {{field: value}}",
                              code="invalid_op")
        total, absents = 0, []
        for cle, val in valeurs.items():
            motif = re.escape(str(cle).strip())
            nb = _moteur(deckio.replace_text, prs,
                         r"(\{\{\s*" + motif + r"\s*\}\}|<<\s*" + motif + r"\s*>>)",
                         "" if val is None else str(val), regex=True, ignore_case=True)
            total += nb
            if not nb:
                absents.append(str(cle))
        if absents:
            env.notes.warn(f"ops[{i}] fill: no placeholder for {', '.join(absents)}")
        return {"filled": total, "changed": bool(total)}
    if nom == "set_text":
        s = _diapo(i, nom, requis(i, nom, c, "slide", EXEMPLE_OPS[1]), prs, origine, supprimees)
        j = entier(requis(i, nom, c, "shape", EXEMPLE_OPS[1]))
        formes = list(s.shapes)
        if j is None or not 0 <= j < len(formes) or not formes[j].has_text_frame:
            raise OfficeError(f"ops[{i}] (set_text): shape {c.get('shape')!r} is not a text shape "
                              f"of slide {c.get('slide')}", code="bad_index",
                              fix="take the shape number of a shape with text from "
                                  "pptx_read(slide=N)")
        tf = formes[j].text_frame
        v = c.get("text")
        lignes = [str(x) for x in v] if isinstance(v, list) else texte(v).split("\n")
        paras = tf.paragraphs
        for k, ln in enumerate(lignes):
            if k < len(paras):
                runs = paras[k].runs
                if runs:
                    runs[0].text = ln
                    for r in runs[1:]:
                        r._r.getparent().remove(r._r)
                else:
                    paras[k].add_run().text = ln
            else:
                p = tf.add_paragraph()
                p.level = paras[-1].level if paras else 0
                p.text = ln
        for p in list(tf.paragraphs)[len(lignes):]:
            p._p.getparent().remove(p._p)
        return {"slide": entier(c.get("slide")), "shape": j}
    if nom == "add_slides":
        slides = c.get("slides")
        md = texte(c.get("markdown"))
        diapos, _ = _diapos(env, slides, md)
        parsed = _moteur(S.parse_slides, diapos)
        avant = len(prs.slides)
        _moteur(B.add_slides, ctx, parsed)
        nouvelles = list(prs.slides)[avant:]
        apres, avant_ = c.get("after"), c.get("before")
        pos: Optional[int] = None
        if avant_ is not None:
            pos = _position(prs, _diapo(i, nom, avant_, prs, origine, supprimees))
        elif apres is not None and fold(apres) not in ("end", "fin", "last"):
            if fold(apres) in ("start", "debut", "first", "0"):
                pos = 0
            else:
                pos = _position(prs, _diapo(i, nom, apres, prs, origine, supprimees)) + 1
        if pos is not None:
            for k, s in enumerate(nouvelles):
                slide_ooxml.move_slide(prs, _position(prs, s), pos + k)
        return {"added": len(nouvelles)}
    if nom == "delete_slide":
        cible = requis(i, nom, c, "slide", EXEMPLE_OPS[3])
        liste = cible if isinstance(cible, list) else [cible]
        diapos = list({id(x): x for x in (_diapo(i, nom, v, prs, origine, supprimees)
                                          for v in liste)}.values())
        for s in diapos:
            supprimees.add(id(s))
            slide_ooxml.delete_slide(prs, _position(prs, s))
        return {"deleted": len(diapos)}
    if nom == "move_slide":
        s = _diapo(i, nom, requis(i, nom, c, "slide", {"op": "move_slide", "slide": 4, "to": 2}),
                   prs, origine, supprimees)
        to = entier(requis(i, nom, c, "to", {"op": "move_slide", "slide": 4, "to": 2}))
        n = len(prs.slides)
        if to is None or not 1 <= to <= n:
            raise OfficeError(f"ops[{i}] (move_slide): position {c.get('to')!r} is outside 1..{n}",
                              code="bad_index")
        slide_ooxml.move_slide(prs, _position(prs, s), to - 1)
        return {"slide": entier(c.get("slide")), "to": to}
    if nom == "duplicate_slide":
        s = _diapo(i, nom, requis(i, nom, c, "slide", {"op": "duplicate_slide", "slide": 2}),
                   prs, origine, supprimees)
        copie = slide_ooxml.duplicate_slide(prs, _position(prs, s))
        to = entier(c.get("to"))
        if to is not None and 1 <= to <= len(prs.slides):
            slide_ooxml.move_slide(prs, _position(prs, copie), to - 1)
        return {"copy_of": entier(c.get("slide")), "at": _position(prs, copie) + 1}
    if nom == "notes":
        s = _diapo(i, nom, requis(i, nom, c, "slide", EXEMPLE_OPS[4]), prs, origine, supprimees)
        slide_ooxml.set_notes(s, texte(c.get("text")))
        return {"slide": entier(c.get("slide"))}
    if nom == "meta":
        valeurs = {k: c[k] for k in ("title", "author", "subject", "keywords") if c.get(k)}
        if not valeurs:
            raise OfficeError(f"ops[{i}] (meta): give title, author, subject or keywords",
                              code="invalid_op")
        deckio.set_metadata(prs, **valeurs)
        return {"set": sorted(valeurs)}
    raise OfficeError(f"ops[{i}]: operation '{nom}' not supported", code="unknown_op")
