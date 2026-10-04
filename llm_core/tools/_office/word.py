# SPDX-License-Identifier: MIT
"""llm_core/tools/_office/word.py — outils Word : ``docx_create``, ``docx_read``, ``docx_edit``.

Sans état : chaque appel lit le fichier dans la sandbox, le transforme en
mémoire et le réécrit (aucun identifiant de document, rien en mémoire entre
deux appels). Le contenu s'écrit en Markdown, avec quatre ajouts :

  ``!a1b2c3d4e5f6`` seul sur sa ligne  graphique ou tableau d'un outil chart_*
  ``> [!NOTE]`` (TIP, IMPORTANT, WARNING, CAUTION)  encadré
  ``[TOC]``                             table des matières
  ``\\newpage``                          saut de page

Numéros (ceux de ``docx_read``) : paragraphes du corps ``p``, tableaux,
lignes et colonnes, tous à partir de 0 (ligne 0 = en-tête). Une liste
d'opérations s'applique au document tel qu'il était AVANT l'appel : les
numéros ne glissent pas après une insertion ou une suppression du même lot.
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.text.paragraph import Paragraph
from lxml import etree
from pydantic import ValidationError

from .._chart.normalise import fold
from . import graphiques as G
from ._docx import docio
from ._docx.errors import DocxMcpError
from ._docx.ooxml.util import qn
from ._docx.render import builder as B, spec as S
from ._docx.render.mdconv import markdown_to_blocks
from .commun import (
    Env,
    OfficeError,
    booleen,
    canon,
    chemin_entree,
    chemin_sortie,
    ecrire,
    entier,
    lire,
    lire_fichier,
    texte,
)
from .ops import lire_ops, requis
from .paquet import PaquetInvalide, installer_lecteur, normaliser, retirer_lecteur

# ── Arguments ───────────────────────────────────────────────────────────────
ALIAS_CREATE = {
    "path": ["file", "filename", "file_name", "file_path", "filepath", "output", "output_path",
             "fichier", "chemin", "name", "nom", "save_as", "destination", "dest", "target"],
    "content": ["markdown", "text", "body", "md", "contenu", "texte", "content_markdown",
                "document", "blocks", "sections", "corps"],
    "title": ["titre", "document_title", "doc_title"],
    "subtitle": ["sous_titre", "subheading"],
    "author": ["auteur", "by", "redacteur", "creator"],
    "date": ["le", "dated"],
    "cover": ["page_de_garde", "cover_page", "title_page", "couverture"],
    "toc": ["table_of_contents", "sommaire", "contents", "table_des_matieres", "summary_table"],
    "template": ["modele", "base", "template_path", "from_template", "model", "gabarit"],
    "theme": ["style", "styles", "design", "look", "colors", "couleurs"],
    "page": ["page_setup", "paper", "format", "mise_en_page"],
    "header": ["en_tete", "entete", "headers"],
    "footer": ["pied", "pied_de_page", "footers"],
    "page_numbers": ["numbering", "page_number", "numeros_de_page", "numerotation",
                     "page_numbering", "numbers"],
    "watermark": ["filigrane"],
}
ALIAS_READ = {
    "path": ALIAS_CREATE["path"],
    "find": ["search", "query", "pattern", "chercher", "rechercher", "q", "grep", "text"],
    "start": ["from", "offset", "from_paragraph", "debut", "start_paragraph", "p"],
    "limit": ["max", "max_paragraphs", "count", "n", "max_items"],
}
ALIAS_EDIT = {
    "path": ALIAS_CREATE["path"],
    "ops": ["operations", "edits", "changes", "actions", "modifications", "steps"],
    "save_as": ["output", "output_path", "new_path", "copy_to", "enregistrer_sous", "dest"],
}
# save_as ne doit pas être pris pour path : dans docx_edit, « output » = save_as.
ALIAS_EDIT["path"] = [a for a in ALIAS_EDIT["path"] if a not in ("output", "output_path",
                                                                  "save_as", "dest")]

OPS = {
    "replace": ["replace_text", "substitute", "remplacer", "find_replace", "search_replace",
                "find_and_replace", "sub"],
    "fill": ["fill_fields", "fill_placeholders", "merge", "remplir", "placeholders", "fields",
             "variables", "mail_merge", "fill_in"],
    "set": ["set_text", "set_paragraph", "rewrite", "update", "edit", "modifier",
            "replace_paragraph", "set_paragraph_text", "change"],
    "delete": ["remove", "delete_paragraph", "supprimer", "del", "delete_paragraphs",
               "remove_paragraph", "delete_table"],
    "insert": ["insert_after", "insert_before", "add_after", "inserer", "insert_blocks",
               "insert_markdown", "add_before"],
    "append": ["add", "add_end", "append_blocks", "ajouter", "add_content", "append_markdown",
               "add_section", "append_section"],
    "set_cell": ["cell", "update_cell", "edit_cell", "cellule", "set_table_cell"],
    "add_row": ["append_row", "insert_row", "ajouter_ligne", "add_table_row", "new_row"],
    "delete_row": ["remove_row", "supprimer_ligne", "delete_table_row"],
    "meta": ["metadata", "properties", "set_metadata", "proprietes", "info", "set_properties"],
}
_F_TEXTE = ["new_text", "value", "texte", "with", "content", "replacement"]
CHAMPS_OPS = {
    "replace": {"find": ["search", "old", "from", "pattern", "text", "chercher", "old_text",
                         "find_text", "source"],
                "replace": ["replacement", "new", "with", "to", "by", "remplacer_par", "value",
                            "new_text", "replace_with"],
                "match_case": ["case_sensitive", "case", "respect_case"],
                "regex": ["is_regex", "use_regex", "regexp"]},
    "fill": {"values": ["fields", "data", "valeurs", "placeholders", "variables", "context",
                        "mapping", "champs"]},
    "set": {"p": ["paragraph", "index", "paragraph_index", "para", "n", "i", "id"],
            "text": _F_TEXTE},
    "delete": {"p": ["paragraph", "index", "paragraph_index", "para", "n", "i", "paragraphs",
                     "indexes", "indices"],
               "to": ["until", "end", "jusqu_a", "to_p", "last"],
               "table": ["table_index", "tableau"]},
    "insert": {"after": ["apres", "after_p", "after_paragraph", "after_heading", "below",
                         "position", "at", "where"],
               "before": ["avant", "before_p", "before_paragraph", "before_heading", "above"],
               "at_end_of": ["end_of", "end_of_section", "fin_de", "section_end", "in_section"],
               "content": ["markdown", "text", "md", "contenu", "texte", "blocks", "body"]},
    "append": {"content": ["markdown", "text", "md", "contenu", "texte", "blocks", "body"]},
    "set_cell": {"table": ["table_index", "tableau", "t"], "row": ["r", "ligne", "row_index"],
                 "col": ["column", "c", "colonne", "col_index", "column_index"],
                 "text": _F_TEXTE},
    "add_row": {"table": ["table_index", "tableau", "t"],
                "values": ["cells", "row", "data", "valeurs", "cellules"],
                "after": ["after_row", "apres", "position", "at"]},
    "delete_row": {"table": ["table_index", "tableau", "t"], "row": ["r", "ligne", "row_index"]},
    "meta": {"title": ["titre"], "author": ["auteur"], "subject": ["sujet"],
             "keywords": ["mots_cles", "tags"]},
}
EXEMPLE_OPS = [
    {"op": "replace", "find": "T3 2026", "replace": "T4 2026"},
    {"op": "fill", "values": {"client": "ACME", "date": "4 octobre 2026"}},
    {"op": "set", "p": 12, "text": "Nouveau texte du paragraphe."},
    {"op": "insert", "after": "Résultats", "content": "Paragraphe ajouté.\n\n!a1b2c3d4e5f6"},
    {"op": "append", "content": "## Conclusion\n\nTexte de conclusion."},
    {"op": "set_cell", "table": 0, "row": 2, "col": 1, "text": "18,4 M€"},
]


def _deduire(f: Dict[str, Any]) -> Optional[str]:
    if "find" in f or "search" in f or ("replace" in f and "p" not in f):
        return "replace"
    if "values" in f or "fields" in f:
        return "fill"
    if "row" in f and ("col" in f or "column" in f):
        return "set_cell"
    if ("after" in f or "before" in f or "at_end_of" in f) and ("content" in f or "markdown" in f):
        return "insert"
    if "content" in f or "markdown" in f:
        return "append"
    if ("p" in f or "paragraph" in f) and "text" in f:
        return "set"
    return None


# ── Markdown → blocs ────────────────────────────────────────────────────────
_FENCE = re.compile(r"^\s*(```|~~~)")
_TOC = re.compile(r"^\s*(\[\[?\s*(toc|_toc_|sommaire|table des mati[eè]res)\s*\]?\]|\{:toc\}|"
                  r"<!--\s*toc\s*-->)\s*$", re.I)
_SAUT = re.compile(r"^\s*(\\newpage|\\pagebreak|\\clearpage|<!--\s*(page\s*-?\s*break|new\s*page|"
                   r"saut\s*de\s*page)\s*-->|\[\s*(page\s*-?\s*break|new\s*page|saut de page)\s*\]|"
                   r"<div[^>]*page-break-(after|before)\s*:\s*always[^>]*>\s*(</div>)?|"
                   r"<br[^>]*page-break[^>]*>)\s*$", re.I)
_ALERTE = re.compile(r"^\s*>\s*\[!(?P<k>[A-Za-zÀ-ÿ]+)\]\s*(?P<t>.*)$")
_VARIANTES = {"note": "note", "info": "info", "tip": "success", "astuce": "success",
              "conseil": "success", "important": "info", "warning": "warning",
              "attention": "warning", "avertissement": "warning", "caution": "danger",
              "danger": "danger", "remarque": "note", "success": "success", "succes": "success",
              "error": "danger", "erreur": "danger"}
_INLINE_MD = re.compile(r"(\*\*|__|\*|_|`)")


def refs_de(contenu: str) -> List[str]:
    return G.REF.findall(contenu or "")


def blocs_markdown(contenu: str, rendus: Dict[str, G.Rendu], env: Env,
                   toc_titre: str = "Sommaire") -> List[Any]:
    """Markdown étendu → blocs du moteur Word (graphiques déjà préparés)."""
    blocs: List[Any] = []
    tampon: List[str] = []

    def vider() -> None:
        if tampon:
            txt = "\n".join(tampon)
            if txt.strip():
                blocs.extend(markdown_to_blocks(txt))
            tampon.clear()

    lignes = (contenu or "").replace("\r\n", "\n").split("\n")
    i, dans_code = 0, False
    while i < len(lignes):
        ligne = lignes[i]
        if _FENCE.match(ligne):
            dans_code = not dans_code
            tampon.append(ligne)
            i += 1
            continue
        if dans_code:
            tampon.append(ligne)
            i += 1
            continue
        m = G.REF_SEULE.match(ligne)
        if m:
            vider()
            blocs.extend(blocs_graphique(G.rendu(rendus, m.group(1))))
            i += 1
            continue
        if G.REF.search(ligne):
            # Référence au milieu d'une phrase : la phrase, puis le graphique.
            ids = G.REF.findall(ligne)
            tampon.append(G.REF.sub("", ligne).rstrip())
            vider()
            for gid in ids:
                blocs.extend(blocs_graphique(G.rendu(rendus, gid)))
            env.notes.fix("chart reference inside a sentence: the chart is placed after it "
                          "(write !id alone on its own line)")
            i += 1
            continue
        if _TOC.match(ligne):
            vider()
            blocs.append(S.TocBlock(type="toc", title=toc_titre, page_break_after=False))
            i += 1
            continue
        if _SAUT.match(ligne):
            vider()
            blocs.append(S.PageBreakBlock(type="page_break"))
            i += 1
            continue
        m = _ALERTE.match(ligne)
        if m:
            vider()
            variante = _VARIANTES.get(fold(m.group("k")), "info")
            corps: List[str] = []
            i += 1
            while i < len(lignes) and lignes[i].lstrip().startswith(">"):
                corps.append(lignes[i].lstrip()[1:].strip())
                i += 1
            titre = m.group("t").strip() or None
            txt = " ".join(c for c in corps if c)
            blocs.append(S.CalloutBlock(type="callout", variant=variante, title=titre,
                                        text=_INLINE_MD.sub("", txt)))
            continue
        tampon.append(ligne)
        i += 1
    vider()
    return blocs


def blocs_graphique(r: G.Rendu) -> List[Any]:
    if r.mode == "natif":
        n = r.natif
        return [S.ChartBlock(type="chart", chart_type=n["chart_type"],
                             categories=n.get("categories") or [],
                             series=[S.SeriesSpec(**s) for s in n["series"]],
                             title=n.get("title"), x_title=n.get("x_title"),
                             y_title=n.get("y_title"), number_format=n.get("number_format"),
                             legend=n.get("legend", "auto"), y_max=n.get("y_max"),
                             smooth=n.get("smooth"))]
    if r.mode == "image":
        return [S.ImageBlock(type="image", source=G.data_uri(r.png), width_cm=16.0,
                             alt_text=r.titre or r.kind)]
    if r.mode == "kpi":
        lignes = [[t.get("label", ""), t.get("value", ""),
                   " ".join(x for x in (t.get("delta"), t.get("ref")) if x)] for t in r.tuiles]
        avec_delta = any(x[2] for x in lignes)
        entete = ["Indicateur", "Valeur"] + (["Évolution"] if avec_delta else [])
        rows = [x if avec_delta else x[:2] for x in lignes]
        return _table_bloc(entete, rows, r.titre)
    return _table_bloc(r.colonnes, [[G.cellule(v) for v in ln] for ln in r.lignes], r.titre)


def _table_bloc(entete: List[str], lignes: List[List[Any]], titre: str) -> List[Any]:
    return [S.TableBlock(type="table", header=[str(h) for h in entete],
                         rows=[[str(c) for c in ln] for ln in lignes],
                         caption=titre or None)]


# ── Ouverture / écriture ────────────────────────────────────────────────────
def _ouvrir(env: Env, path: Any) -> Tuple[str, Any, str]:
    """→ (chemin, document, empreinte des octets lus : verrou de l'écriture)."""
    import hashlib
    rel, brut = lire_fichier(env, path, "docx", "Word document")
    sha = hashlib.sha256(brut).hexdigest()
    try:
        blob, notes = normaliser(brut, "docx", f"file '{rel}'")
    except PaquetInvalide as e:
        raise OfficeError(str(e), code="bad_file",
                          fix="give the path of a .docx file (Word 2007 or later)")
    for n in notes:
        env.notes.fix(n)
    try:
        return rel, docx.Document(io.BytesIO(blob)), sha
    except Exception as e:                       # noqa: BLE001 — fichier corrompu
        raise OfficeError(f"'{rel}' could not be opened as a Word document: {e}",
                          code="bad_file", fix="check that the file is a valid .docx")


def _enregistrer(env: Env, document: Any, rel: str, attendu: Optional[str] = None
                 ) -> Dict[str, Any]:
    B.finalize(document)
    sortie = io.BytesIO()
    document.save(sortie)
    return ecrire(env, rel, sortie.getvalue(), attendu)


def _avec_lecteur(env: Env):
    """Les images et modèles cités par chemin se lisent dans la sandbox de l'appelant."""

    def lecteur(source: str, quoi: str) -> bytes:
        rel = chemin_entree(source)
        try:
            return lire(env, rel, quoi)
        except FileNotFoundError:
            raise OfficeError(f"{quoi} not found in the sandbox: {env.espace.afficher(rel)}",
                              code="not_found",
                              fix=f"check the {quoi} path (list the folder first)")
    return installer_lecteur(lecteur)


def _moteur(fn, *a, **k):
    """Erreurs du moteur → refus guidé."""
    try:
        return fn(*a, **k)
    except OfficeError:
        raise
    except ValidationError as e:
        err = e.errors()[0] if e.errors() else {}
        lieu = ".".join(str(x) for x in err.get("loc", ()))
        raise OfficeError(f"invalid content at {lieu or 'root'}: {err.get('msg', e)}",
                          code="invalid_content")
    except (DocxMcpError, PaquetInvalide) as e:
        raise OfficeError(str(e), code="invalid_content")
    except re.error as e:
        raise OfficeError(f"invalid regular expression: {e}", code="invalid_op",
                          fix="fix the pattern, or set regex=false to search the text literally")
    except (etree.XMLSyntaxError, zipfile.BadZipFile) as e:
        # Modèle ou fichier dont le XML est cassé (passé le contrôle du paquet).
        raise OfficeError(f"the file or template is damaged: {e}", code="bad_file",
                          fix="open and re-save it with Word or LibreOffice, or use another file")


# ── docx_create ─────────────────────────────────────────────────────────────
def docx_create(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_CREATE, env.notes)
    rel = chemin_sortie(a.get("path"), "docx", env.notes)
    contenu = a.get("content")
    if isinstance(contenu, list) and contenu and all(isinstance(x, dict) for x in contenu):
        blocs_bruts: Optional[List[Any]] = contenu
        contenu_md = ""
    else:
        blocs_bruts = None
        contenu_md = texte(contenu)
    if not contenu_md.strip() and not blocs_bruts:
        raise OfficeError("content is empty", code="empty_content",
                          fix="content = the document text in Markdown (# headings, lists, "
                              "tables, **bold**…)",
                          example={"path": "rapports/bilan.docx", "title": "Bilan 2026",
                                   "content": "# Synthèse\n\nLe chiffre d'affaires progresse de "
                                              "**12 %**.\n\n## Détail\n\n- Point 1\n- Point 2"})
    titre = str(a.get("title") or "").strip()
    jeton = _avec_lecteur(env)
    try:
        rendus = G.preparer(env, refs_de(contenu_md))
        blocs: List[Any] = []
        couverture = a.get("cover")
        if couverture:
            c = couverture if isinstance(couverture, dict) else {}
            c = canon(c, {"title": ["titre"], "subtitle": ["sous_titre"], "author": ["auteur"],
                          "date": [], "logo": ["image"]}, env.notes, "cover.")
            blocs.append(S.CoverBlock(type="cover", title=str(c.get("title") or titre or ""),
                                      subtitle=c.get("subtitle") or a.get("subtitle"),
                                      author=c.get("author") or a.get("author"),
                                      date=c.get("date") or a.get("date"), logo=c.get("logo")))
        elif titre and not re.match(r"^\s*#\s", contenu_md or ""):
            blocs.append(S.HeadingBlock(type="heading", level=0, text=titre))
            if a.get("subtitle"):
                blocs.append(S.ParagraphBlock(type="paragraph", text=str(a["subtitle"]),
                                              style="Subtitle"))
        if booleen(a.get("toc")) and not _TOC.search(contenu_md or ""):
            blocs.append(S.TocBlock(type="toc", title="Sommaire"))
        if blocs_bruts is not None:
            env.notes.fix("content given as blocks: built as-is (Markdown is simpler)")
            blocs.extend(_moteur(S.parse_blocks, blocs_bruts))
        else:
            blocs.extend(_moteur(blocs_markdown, contenu_md, rendus, env))
        # Avec un modèle, sa mise en page et son pied font foi : on ne les
        # remplace que sur demande explicite (sinon A4 portrait et « Page X / Y »
        # par défaut effaceraient la charte).
        modele = bool(a.get("template"))
        page = _page(a.get("page"), env) if (a.get("page") or not modele) else None
        pied = _pied(a.get("footer"), booleen(a.get("page_numbers"), not modele)) \
            if (a.get("footer") or a.get("page_numbers") is not None or not modele) else None
        spec = _moteur(S.DocumentSpec,
                       filename=rel.rsplit("/", 1)[-1],
                       meta=S.MetaSpec(title=titre or None, author=a.get("author") or None),
                       theme=_theme(a.get("theme"), env), page=page,
                       header=_entete(a.get("header")),
                       footer=pied,
                       watermark=S.WatermarkSpec(text=str(a["watermark"]))
                       if a.get("watermark") else None,
                       blocks=blocs)
        base = None
        if a.get("template"):
            from ._docx.templates import load
            base, notes_m, _ = _moteur(load, str(a["template"]))
            for n in notes_m:
                env.notes.fix(n)
            try:
                docx.Document(io.BytesIO(base))
            except Exception as e:               # noqa: BLE001 — modèle illisible
                raise OfficeError(f"the template '{a['template']}' cannot be opened: {e}",
                                  code="bad_file", fix="use a valid .dotx or .docx template")
        document, ctx = _moteur(B.build_document, spec, base)
    finally:
        retirer_lecteur(jeton)
    _notes_moteur(ctx.notes, env)
    res = _enregistrer(env, document, rel)
    bilan = _bilan(document, rendus)
    return {**res, **bilan,
            "summary": f"{'Replaced' if res['old_sha256'] else 'Created'} {res['path']}: "
                       f"{_phrase(bilan)}"}


# Notes du moteur sans action possible pour le modèle : non relayées.
_NOTES_MUETTES = ("Table of contents:",)


def _notes_moteur(notes: List[str], env: Env) -> None:
    for n in notes:
        if not n.startswith(_NOTES_MUETTES):
            env.notes.warn(n)


def _theme(v: Any, env: Env) -> S.ThemeSpec:
    if isinstance(v, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", v.strip()):
        return S.ThemeSpec(accent=v.strip())
    if not isinstance(v, dict):
        if v:
            env.notes.warn(f"theme '{v}' not understood: default look kept "
                           "(theme = {\"accent\": \"#1F4E79\", \"font\": \"Calibri\"})")
        return S.ThemeSpec()
    t = canon(v, {"accent": ["color", "colour", "couleur", "primary", "main_color"],
                  "body_font": ["font", "police", "font_family", "text_font"],
                  "heading_font": ["title_font", "headings_font", "police_titres"],
                  "base_font_size_pt": ["font_size", "size", "taille", "font_size_pt"]},
              env.notes, "theme.")
    garde = {k: t[k] for k in ("accent", "body_font", "heading_font", "base_font_size_pt")
             if t.get(k) not in (None, "")}
    return _moteur(S.ThemeSpec, **garde)


def _page(v: Any, env: Env) -> S.PageSpec:
    if v in (None, "", {}):
        return S.PageSpec()
    if isinstance(v, str):
        f = fold(v)
        orient = "landscape" if any(x in f for x in ("landscape", "paysage", "horizontal")) \
            else "portrait"
        taille = next((t for t in ("a3", "a4", "a5", "letter", "legal") if t in f), "a4")
        return S.PageSpec(size=taille.upper() if taille.startswith("a") else taille.capitalize(),
                          orientation=orient)
    if isinstance(v, dict):
        p = canon(v, {"size": ["format", "paper", "taille"],
                      "orientation": ["orient", "sens"],
                      "margins_cm": ["margins", "margin", "marges", "margin_cm"]}, env.notes, "page.")
        o = fold(p.get("orientation") or "portrait")
        marges = p.get("margins_cm")
        if isinstance(marges, (int, float)):
            marges = {k: float(marges) for k in ("top", "bottom", "left", "right")}
        return _moteur(S.PageSpec, size=str(p.get("size") or "A4"),
                       orientation="landscape" if o in ("landscape", "paysage") else "portrait",
                       margins_cm=marges if isinstance(marges, dict) else None)
    return S.PageSpec()


def _entete(v: Any) -> Optional[S.HeaderSpec]:
    if not v:
        return None
    if isinstance(v, dict):
        d = {k: v.get(k) for k in ("text", "left", "center", "right") if v.get(k)}
        return S.HeaderSpec(**d, rule=True) if d else None
    return S.HeaderSpec(text=str(v), rule=True)


def _pied(v: Any, numeros: bool) -> Optional[S.FooterSpec]:
    if isinstance(v, dict):
        d = {k: v.get(k) for k in ("text", "left", "center", "right") if v.get(k)}
        return S.FooterSpec(**d, page_numbers=numeros, page_number_format="Page {PAGE} / {NUMPAGES}")
    if v:
        return S.FooterSpec(left=str(v), align="left", page_numbers=numeros,
                            page_number_format="Page {PAGE} / {NUMPAGES}")
    if numeros:
        return S.FooterSpec(page_numbers=True, page_number_format="Page {PAGE} / {NUMPAGES}")
    return None


def _bilan(document: Any, rendus: Optional[Dict[str, G.Rendu]] = None) -> Dict[str, Any]:
    plan = [{"p": x["index"], "level": x["level"], "text": x["text"][:120]}
            for x in docio.get_outline(document)]
    out: Dict[str, Any] = {"paragraphs": len(document.paragraphs), "tables": len(document.tables),
                           "images": _nb_images(document), "outline": plan[:60]}
    if rendus:
        out["charts"] = [{"ref": f"!{r.id}", "inserted_as": {"natif": "native chart",
                                                             "image": "image", "table": "table",
                                                             "kpi": "table"}[r.mode]}
                         for r in rendus.values()]
    return out


def _nb_images(document: Any) -> int:
    return sum(1 for s in document.inline_shapes
               if s._inline.find(".//" + qn("pic:pic")) is not None)


def _phrase(b: Dict[str, Any]) -> str:
    parts = [f"{len(b['outline'])} headings", f"{b['paragraphs']} paragraphs"]
    if b["tables"]:
        parts.append(f"{b['tables']} tables")
    if b.get("charts"):
        nat = sum(1 for c in b["charts"] if c["inserted_as"] == "native chart")
        parts.append(f"{len(b['charts'])} charts ({nat} native, editable)")
    elif b["images"]:
        parts.append(f"{b['images']} images")
    return ", ".join(parts)


# ── docx_read ───────────────────────────────────────────────────────────────
def docx_read(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_READ, env.notes)
    rel, document, _ = _ouvrir(env, a.get("path"))
    start = max(0, entier(a.get("start")) or 0)
    limit = min(max(1, entier(a.get("limit")) or 200), 1000)
    # Clés = éléments lxml eux-mêmes, pas leur id() : un mandataire lxml sans
    # référence est libéré et son id peut être réattribué à un autre élément.
    para_index = {p._p: i for i, p in enumerate(document.paragraphs)}
    tables = {t._tbl: (k, t) for k, t in enumerate(document.tables)}
    levels, default_id = docio.style_levels(document)
    contenu: List[Dict[str, Any]] = []
    # Page = paragraphes [start, suite) + les tableaux qui les SUIVENT (un
    # tableau appartient au paragraphe d'avant ; ceux d'avant le 1er paragraphe,
    # à la page 0). On ne coupe qu'à un paragraphe : un tableau n'est jamais
    # listé deux fois et ``next_start`` avance toujours.
    dernier_p, suite = -1, None
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p") and child in para_index:
            i = para_index[child]
            dernier_p = i
            if i < start:
                continue
            if len(contenu) >= limit and i > start:
                suite = i
                break
            par = Paragraph(child, document._body)
            item = _decrire_paragraphe(par, i, levels, default_id)
            if item:
                contenu.append(item)
        elif child.tag == qn("w:tbl") and child in tables and \
                (dernier_p >= start or (start == 0 and dernier_p == -1)):
            k, t = tables[child]
            lignes = [[c.text.strip()[:200] for c in row.cells] for row in t.rows]
            item = {"table": k, "after_p": dernier_p, "size": f"{len(lignes)}x{len(t.columns)}",
                    "rows": lignes[:40]}
            if len(lignes) > 40:
                item["rows_truncated"] = len(lignes) - 40
            contenu.append(item)
    res: Dict[str, Any] = {"path": env.espace.afficher(rel),
                           "paragraphs": len(document.paragraphs),
                           "tables": len(document.tables),
                           "sections": len(document.sections),
                           "title": document.core_properties.title or None,
                           "outline": [{"p": x["index"], "level": x["level"], "text": x["text"][:120]}
                                       for x in docio.get_outline(document)][:80],
                           "content": contenu}
    if suite is not None:
        res["next_start"] = suite
        res["note"] = f"more content: call docx_read again with start={suite}"
    entetes = docio.read_stories(document)
    for cle in ("headers", "footers"):
        textes = [x.get("text") for x in entetes.get(cle, []) if x.get("text")]
        if textes:
            res[cle[:-1]] = " | ".join(dict.fromkeys(textes))[:400]
    champs = sorted({m.strip() for _, p in docio.iter_located_paragraphs(document)
                     for m in re.findall(r"\{\{\s*([^{}]{1,60}?)\s*\}\}", p.text)})
    if champs:
        res["placeholders"] = champs
    verif = docio.check_document(document, check_leftovers=True)
    restes = [{"means": v.get("means"), "count": v.get("count"),
               "at": [_lieu(o) for o in (v.get("occurrences") or [])[:5]]}
              for v in verif.get("violations", [])
              if v.get("kind") == "leftover" and "Jinja" not in str(v.get("means"))]
    if restes:
        res["issues"] = restes[:20]
    if a.get("find"):
        trouves = docio.find_pattern(document, str(a["find"]), ignore_case=True, limit=50)
        res["matches"] = [_lieu(x) for x in trouves]
        if not trouves:
            res["matches_note"] = f"'{a['find']}' not found (case ignored)"
    return res


def _decrire_paragraphe(par: Paragraph, i: int, levels, default_id) -> Optional[Dict[str, Any]]:
    t = par.text.strip()
    el = par._p
    image = el.find(".//" + qn("w:drawing")) is not None
    graphique = image and el.find(".//" + qn("c:chart")) is not None
    if not t and not image:
        return None
    item: Dict[str, Any] = {"p": i}
    niveau = docio.heading_level_for(par, levels, default_id)
    if niveau is not None:
        item["h"] = niveau
    else:
        style = par.style.name if par.style is not None else ""
        if style and style not in ("Normal", "Body Text", "Corps de texte", "Default Paragraph Font"):
            item["style"] = style
    if t:
        item["text"] = t if len(t) <= 2000 else t[:2000] + "…"
    if niveau is None and el.find(".//" + qn("w:numPr")) is not None:
        item["list"] = True
    if image:
        item["object"] = "chart" if graphique else "image"
    return item


def _lieu(x: Dict[str, Any]) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    if x.get("where") == "body":
        out["p"] = x.get("paragraph_index")
    elif x.get("where") == "table":
        out.update({"table": x.get("table_index"), "row": x.get("row"), "col": x.get("column")})
    else:
        out["in"] = x.get("where")
    out["excerpt"] = x.get("excerpt")
    return out


# ── docx_edit ───────────────────────────────────────────────────────────────
class _Lot:
    """Le document tel qu'il était AVANT l'appel : paragraphes, tableaux et
    lignes figés au début (les lignes d'un tableau, à son premier usage), et ce
    que le lot a déjà supprimé. Sans cela, « supprimer le tableau 0 puis le
    tableau 1 » supprimerait l'ancien tableau 2."""

    def __init__(self, document: Any) -> None:
        self.document = document
        self.paragraphes = list(document.paragraphs)
        self.tables = list(document.tables)
        self._lignes: Dict[int, List[Any]] = {}
        self.supprimes: set = set()          # éléments lxml (gardés vivants)

    def p(self, i: int, nom: str, v: Any) -> Paragraph:
        k = entier(v)
        n = len(self.paragraphes)
        if k is None or not 0 <= k < n:
            raise OfficeError(f"ops[{i}] ({nom}): paragraph {v!r} does not exist "
                              f"(the document has paragraphs 0 to {n - 1})", code="bad_index",
                              fix="take the 'p' numbers from docx_read")
        par = self.paragraphes[k]
        if par._p in self.supprimes:
            raise OfficeError(f"ops[{i}] ({nom}): paragraph {k} was deleted by an earlier "
                              "operation of this call", code="bad_index",
                              fix="reorder or merge the operations")
        return par

    def table(self, i: int, nom: str, v: Any):
        k = entier(v)
        n = len(self.tables)
        if k is None or not 0 <= k < n:
            raise OfficeError(f"ops[{i}] ({nom}): table {v!r} does not exist (the document has "
                              f"{n} table(s), numbered from 0)", code="bad_index",
                              fix="take the table number from docx_read")
        t = self.tables[k]
        if t._tbl in self.supprimes:
            raise OfficeError(f"ops[{i}] ({nom}): table {k} was deleted by an earlier operation "
                              "of this call", code="bad_index")
        return k, t

    def lignes(self, k: int, t: Any) -> List[Any]:
        if k not in self._lignes:
            self._lignes[k] = list(t.rows)
        return self._lignes[k]

    def ligne(self, i: int, nom: str, k: int, t: Any, v: Any):
        lignes = self.lignes(k, t)
        r = entier(v)
        if r is None or not 0 <= r < len(lignes) or lignes[r]._tr in self.supprimes:
            raise OfficeError(f"ops[{i}] ({nom}): row {v!r} does not exist in table {k} "
                              f"(rows 0 to {len(lignes) - 1}, row 0 = header)", code="bad_index",
                              example=[EXEMPLE_OPS[5]])
        return r, lignes[r]


def docx_edit(args: Dict[str, Any], env: Env) -> Dict[str, Any]:
    a = canon(args, ALIAS_EDIT, env.notes)
    ops = lire_ops(a.get("ops"), OPS, CHAMPS_OPS, _deduire, env.notes, EXEMPLE_OPS)
    rel, document, sha_lu = _ouvrir(env, a.get("path"))
    cible = rel
    if a.get("save_as"):
        cible = chemin_sortie(a["save_as"], "docx", env.notes)
    elif not rel.lower().endswith(".docx"):
        cible = chemin_sortie(rel, "docx", env.notes)
        env.notes.fix(f"a {rel.rsplit('.', 1)[-1]} file is saved as a new .docx: {cible}")
    lot = _Lot(document)
    contenus = " ".join(texte(c.get("content")) for _, c in ops)
    jeton = _avec_lecteur(env)
    bilan_ops: List[Dict[str, Any]] = []
    try:
        rendus = G.preparer(env, refs_de(contenus))
        for i, (nom, c) in enumerate(ops):
            bilan_ops.append({"op": nom, **_appliquer(i, nom, c, lot, rendus, env)})
    finally:
        retirer_lecteur(jeton)
    if not any(o.get("changed", True) for o in bilan_ops):
        raise OfficeError("no operation changed the document: nothing was saved",
                          code="nothing_changed",
                          fix="check the exact text or numbers with docx_read (find=...), then "
                              "retry with matching values",
                          example=EXEMPLE_OPS[:2])
    # Même fichier : écrit seulement s'il n'a pas changé depuis la lecture.
    res = _enregistrer(env, document, cible, sha_lu if cible == rel else None)
    bilan = _bilan(document)
    faits = [o["op"] + ("" if o.get("changed", True) else " (no change)") for o in bilan_ops]
    return {**res, "applied": bilan_ops, "paragraphs": bilan["paragraphs"],
            "tables": bilan["tables"], "outline": bilan["outline"],
            "summary": f"Saved {res['path']}: {len(ops)} operation(s) ({', '.join(faits)})"}


def _ancre(i: int, nom: str, v: Any, lot: _Lot, fin_de_section: bool = False):
    """Ancre d'insertion → élément du corps APRÈS lequel insérer (ou avant, selon
    l'appelant) ; ``None`` = à la fin ; ``"start"`` = au début.

    ``v`` : numéro de paragraphe, « start » / « end », ou texte d'un titre (sinon
    d'un paragraphe). ``fin_de_section`` : dernier élément de la section du titre,
    tableau compris, juste avant le titre suivant de niveau égal ou supérieur."""
    if entier(v) is not None:
        return lot.p(i, nom, v)._p
    f = fold(v)
    if f in ("end", "fin", "bottom", "last"):
        return None
    if f in ("start", "debut", "top", "beginning", "first"):
        return "start"
    levels, default_id = docio.style_levels(lot.document)
    vivants = [(k, p) for k, p in enumerate(lot.paragraphes) if p._p not in lot.supprimes]
    titres = [(k, p, docio.heading_level_for(p, levels, default_id)) for k, p in vivants]
    titres = [(k, p, lv) for k, p, lv in titres if lv is not None]
    trouve = next(((k, p, lv) for k, p, lv in titres if fold(p.text) == f), None) or \
        next(((k, p, lv) for k, p, lv in titres if f and f in fold(p.text)), None)
    if trouve is None:
        tous = next(((k, p, None) for k, p in vivants if f and f in fold(p.text)), None)
        if tous is None or fin_de_section:
            raise OfficeError(f"ops[{i}] ({nom}): no heading matches '{v}'", code="bad_anchor",
                              fix="use a paragraph number 'p' or the exact heading text from "
                                  "docx_read (outline)")
        trouve = tous
    k, p, lv = trouve
    if not fin_de_section:
        return p._p
    niveau = lv if lv is not None else 9          # 0 = titre du document : tout le reste
    suivant = next((p2 for k2, p2, lv2 in titres if k2 > k and lv2 <= niveau), None)
    if suivant is None:
        return None
    return suivant._p.getprevious()


def _inserer(document, blocs: List[Any], ancre, avant: bool, env: Env) -> int:
    corps = document.element.body
    existants = set(corps)          # éléments gardés vivants (cf. docx_read)
    ctx = B.BuildContext(document=document)
    _moteur(B.add_blocks, ctx, blocs)
    _notes_moteur(ctx.notes, env)
    crees = [x for x in corps if x not in existants]
    if ancre not in (None, "start") and ancre.getparent() is None:
        # Ancre retirée du document (ne doit pas arriver : _ancre écarte les
        # supprimés) : refuser plutôt que d'insérer dans le vide.
        for el in crees:
            corps.remove(el)
        raise OfficeError("the insertion point was removed by an earlier operation of this call",
                          code="bad_anchor", fix="reorder the operations or use another anchor")
    if ancre is None:
        return len(crees)
    if ancre == "start":
        premier = next(iter(corps))          # jamais vide : w:sectPr reste en place
        for el in crees:
            corps.remove(el)
            premier.addprevious(el)
        return len(crees)
    cible = ancre
    for el in crees:
        corps.remove(el)
        if avant:
            ancre.addprevious(el)
        else:
            cible.addnext(el)
            cible = el
    return len(crees)


def _appliquer(i: int, nom: str, c: Dict[str, Any], lot: _Lot, rendus: Dict[str, G.Rendu],
               env: Env) -> Dict[str, Any]:
    document = lot.document
    if nom == "replace":
        cherche = str(requis(i, nom, c, "find", EXEMPLE_OPS[0]))
        rempl = texte(c.get("replace"))
        nb = _moteur(docio.replace_text, document, cherche, rempl,
                     regex=booleen(c.get("regex")), ignore_case=not booleen(c.get("match_case"), True))
        if not nb:
            env.notes.warn(f"ops[{i}] replace: '{cherche}' not found (case-sensitive; set "
                           "match_case=false to ignore case)")
        return {"replaced": nb, "changed": bool(nb)}
    if nom == "fill":
        valeurs = requis(i, nom, c, "values", EXEMPLE_OPS[1])
        if not isinstance(valeurs, dict):
            raise OfficeError(f"ops[{i}] (fill): values must be an object {{field: value}}",
                              example=[EXEMPLE_OPS[1]], code="invalid_op")
        total, absents = 0, []
        for cle, val in valeurs.items():
            k = re.escape(str(cle).strip())
            motif = r"(\{\{\s*" + k + r"\s*\}\}|<<\s*" + k + r"\s*>>|«\s*" + k + r"\s*»)"
            nb = _moteur(docio.replace_text, document, motif, "" if val is None else str(val),
                         regex=True, ignore_case=True)
            total += nb
            if not nb:
                absents.append(str(cle))
        if absents:
            env.notes.warn(f"ops[{i}] fill: no placeholder for {', '.join(absents)}")
        return {"filled": total, "changed": bool(total)}
    if nom == "set":
        par = lot.p(i, nom, requis(i, nom, c, "p", EXEMPLE_OPS[2]))
        _ecrire_paragraphe(par, _INLINE_MD.sub("", texte(c.get("text"))))
        return {"p": entier(c.get("p"))}
    if nom == "delete":
        if c.get("table") is not None:
            k, t = lot.table(i, nom, c["table"])
            lot.supprimes.add(t._tbl)
            t._tbl.getparent().remove(t._tbl)
            return {"table": k}
        cible = requis(i, nom, c, "p", {"op": "delete", "p": 7})
        liste = cible if isinstance(cible, list) else [cible]
        if c.get("to") is not None and not isinstance(cible, list):
            de, a_ = entier(cible), entier(c["to"])
            if de is None or a_ is None or a_ < de:
                raise OfficeError(f"ops[{i}] (delete): bad range {cible}..{c['to']}",
                                  code="bad_index", example=[{"op": "delete", "p": 4, "to": 6}])
            liste = list(range(de, a_ + 1))
        pars = list({id(x._p): x for x in (lot.p(i, nom, v) for v in liste)}.values())
        for par in pars:
            lot.supprimes.add(par._p)
            par._p.getparent().remove(par._p)
        return {"deleted_paragraphs": len(pars)}
    if nom in ("insert", "append"):
        contenu = texte(requis(i, nom, c, "content",
                               EXEMPLE_OPS[3] if nom == "insert" else EXEMPLE_OPS[4]))
        blocs = _moteur(blocs_markdown, contenu, rendus, env)
        if nom == "append" or all(c.get(k) is None for k in ("after", "before", "at_end_of")):
            if nom == "insert":
                env.notes.fix(f"ops[{i}] insert without 'after'/'before': appended at the end")
            return {"added_elements": _inserer(document, blocs, None, False, env)}
        if c.get("at_end_of") is not None:
            ancre = _ancre(i, nom, c["at_end_of"], lot, True)
            return {"added_elements": _inserer(document, blocs, ancre, False, env)}
        avant = c.get("after") is None
        ancre = _ancre(i, nom, c.get("before") if avant else c.get("after"), lot)
        if avant and ancre is None:              # « avant la fin » = à la fin
            avant = False
        return {"added_elements": _inserer(document, blocs, ancre, avant, env)}
    if nom in ("set_cell", "add_row", "delete_row"):
        k, t = lot.table(i, nom, requis(i, nom, c, "table", {"op": nom, "table": 0}))
        if nom == "set_cell":
            ln, ligne = lot.ligne(i, nom, k, t, c.get("row"))
            col = entier(c.get("col"))
            cells = ligne.cells
            if col is None or not 0 <= col < len(cells):
                raise OfficeError(f"ops[{i}] (set_cell): column {c.get('col')!r} does not exist "
                                  f"in table {k} (columns 0 to {len(cells) - 1})",
                                  code="bad_index", example=[EXEMPLE_OPS[5]])
            _ecrire_cellule(cells[col], texte(c.get("text")))
            return {"table": k, "row": ln, "col": col}
        if nom == "delete_row":
            ln, ligne = lot.ligne(i, nom, k, t, c.get("row"))
            lot.supprimes.add(ligne._tr)
            ligne._tr.getparent().remove(ligne._tr)
            return {"table": k, "row": ln}
        vals = requis(i, nom, c, "values", {"op": "add_row", "table": 0, "values": ["A", "12"]})
        lignes = [x for x in lot.lignes(k, t) if x._tr not in lot.supprimes]
        if not lignes:
            raise OfficeError(f"ops[{i}] (add_row): table {k} has no row left to copy",
                              code="bad_index")
        if isinstance(vals, dict):
            entete = [cel.text.strip() for cel in lignes[0].cells]
            dico = {fold(kk): vv for kk, vv in vals.items()}
            vals = [dico.get(fold(h), "") for h in entete]
        if not isinstance(vals, list):
            vals = [vals]
        import copy
        nouvelle = copy.deepcopy(lignes[-1]._tr)
        apres = entier(c.get("after"))
        if apres is not None:
            _, ref = lot.ligne(i, nom, k, t, apres)
            ref._tr.addnext(nouvelle)
        else:
            lignes[-1]._tr.addnext(nouvelle)
        from docx.table import _Row
        ligne = _Row(nouvelle, t)
        for j, cel in enumerate(ligne.cells):
            _ecrire_cellule(cel, G.cellule(vals[j]) if j < len(vals) else "")
        if len(vals) > len(ligne.cells):
            env.notes.warn(f"ops[{i}] add_row: {len(vals) - len(ligne.cells)} extra value(s) ignored")
        return {"table": k, "rows": len(t.rows)}
    if nom == "meta":
        valeurs = {k: c[k] for k in ("title", "author", "subject", "keywords") if c.get(k)}
        if not valeurs:
            raise OfficeError(f"ops[{i}] (meta): give title, author, subject or keywords",
                              code="invalid_op", example=[{"op": "meta", "title": "Bilan 2026"}])
        docio.set_metadata(document, **valeurs)
        return {"set": sorted(valeurs)}
    raise OfficeError(f"ops[{i}]: operation '{nom}' not supported", code="unknown_op")


def _ecrire_paragraphe(par: Paragraph, valeur: str) -> None:
    """Remplace le texte d'un paragraphe en gardant son style et la mise en
    forme de son premier fragment."""
    runs = par.runs
    if runs:
        runs[0].text = valeur
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(valeur)


def _ecrire_cellule(cel, valeur: str) -> None:
    pars = cel.paragraphs
    premier = pars[0]
    runs = premier.runs
    if runs:
        runs[0].text = valeur
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        premier.add_run(valeur)
    for p in pars[1:]:
        p._p.getparent().remove(p._p)
