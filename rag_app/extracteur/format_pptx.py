# SPDX-License-Identifier: MIT
# rag_app/extracteur/format_pptx.py — extraction PPTX (python-pptx) : une diapositive = une page, titre = heading.
from __future__ import annotations

from pathlib import Path
from typing import Any

from .blocs import Collecteur, cellule, compacter
from .erreurs import ErreurExtracteur

_EMU_PAR_POINT = 12700


def _boite(forme: Any) -> list[list[float]] | None:
    try:
        x, y, w, h = forme.left, forme.top, forme.width, forme.height
    except AttributeError:
        return None
    if None in (x, y, w, h) or w <= 0 or h <= 0:
        return None
    return [[round(x / _EMU_PAR_POINT, 1), round(y / _EMU_PAR_POINT, 1),
             round((x + w) / _EMU_PAR_POINT, 1), round((y + h) / _EMU_PAR_POINT, 1)]]


def _formes(formes: Any, groupe: bool = False):
    from pptx.shapes.group import GroupShape
    for forme in formes:
        if isinstance(forme, GroupShape):
            yield from _formes(forme.shapes, True)
        else:
            yield forme, groupe


def _texte_forme(forme: Any, col: Collecteur, n: int, boite: list[list[float]] | None, est_titre: bool) -> None:
    items: list[str] = []

    def vider() -> None:
        nonlocal items
        if items:
            col.ajouter("list", "\n".join(items), page=n, slide=n, boxes=boite)
        items = []

    for p in forme.text_frame.paragraphs:
        texte = compacter("".join(r.text for r in p.runs) or p.text)
        if not texte:
            continue
        if est_titre:
            col.ajouter("heading", texte, level=1, page=n, slide=n, boxes=boite)
        elif p.level > 0 or _a_puce(p):
            items.append("  " * p.level + "- " + texte)
        else:
            vider()
            col.ajouter("paragraph", texte, page=n, slide=n, boxes=boite)
    vider()


def _a_puce(paragraphe: Any) -> bool:
    ppr = paragraphe._p.pPr
    return ppr is not None and any(e.tag.endswith(("}buChar", "}buAutoNum")) for e in ppr)


def extraire(chemin: Path, col: Collecteur, meta: dict[str, Any]) -> None:
    from pptx import Presentation
    try:
        prs = Presentation(str(chemin))
    except MemoryError:
        raise
    except Exception as exc:  # noqa: BLE001 — toute erreur d'ouverture = fichier illisible
        raise ErreurExtracteur("corrupt", f"PPTX illisible : {type(exc).__name__}")
    meta["title"] = (prs.core_properties.title or "").strip() or None
    total = meta["pages"] = len(prs.slides)
    for n, diapo in enumerate(prs.slides, start=1):
        col.progresser(n - 1, total)
        titre_forme = diapo.shapes.title
        if titre_forme is not None and titre_forme.has_text_frame:
            _texte_forme(titre_forme, col, n, _boite(titre_forme), True)
            meta["title"] = meta["title"] or compacter(titre_forme.text_frame.text) or None
        for forme, dans_groupe in _formes(diapo.shapes):
            if titre_forme is not None and forme.shape_id == titre_forme.shape_id:
                continue
            boite = None if dans_groupe else _boite(forme)
            if getattr(forme, "has_table", False) and forme.has_table:
                col.tableau(([cellule(c.text) for c in rang.cells] for rang in forme.table.rows),
                            page=n, slide=n, boxes=boite)
            elif forme.has_text_frame:
                _texte_forme(forme, col, n, boite, False)
        if diapo.has_notes_slide:
            notes = compacter(diapo.notes_slide.notes_text_frame.text)
            if notes:
                col.ajouter("paragraph", "Notes : " + notes, page=n, slide=n)
    col.progresser(total, total)
