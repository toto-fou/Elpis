# SPDX-License-Identifier: MIT
# rag_app/extracteur/format_xlsx.py — extraction XLSX en lecture seule (openpyxl) : une feuille = tableaux Markdown.
from __future__ import annotations

import datetime as dt
from collections.abc import Iterator
from itertools import chain
from pathlib import Path
from typing import Any

from .blocs import Collecteur, cellule
from .erreurs import ErreurExtracteur

_LIGNES_MAX = 500_000


def _valeur(v: Any) -> str:
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return cellule(v)


def _lignes(feuille: Any, col: Collecteur) -> Iterator[list[str]]:
    for n, rang in enumerate(feuille.iter_rows(values_only=True), start=1):
        if n > _LIGNES_MAX:
            col.avertir("truncated", f"Feuille « {feuille.title} » tronquée à {_LIGNES_MAX} lignes.")
            return
        cellules = [_valeur(v) for v in rang]
        while cellules and not cellules[-1]:
            cellules.pop()
        if cellules:
            yield cellules


def extraire(chemin: Path, col: Collecteur, meta: dict[str, Any]) -> None:
    from openpyxl import load_workbook
    try:
        classeur = load_workbook(chemin, read_only=True, data_only=True)
    except MemoryError:
        raise
    except Exception as exc:  # noqa: BLE001 — toute erreur d'ouverture = fichier illisible
        raise ErreurExtracteur("corrupt", f"XLSX illisible : {type(exc).__name__}")
    try:
        meta["title"] = (classeur.properties.title or "").strip() or None
        feuilles = [f for f in classeur.worksheets if f.sheet_state == "visible"]
        for n, feuille in enumerate(feuilles):
            col.progresser(n, len(feuilles))
            lignes = _lignes(feuille, col)
            premiere = next(lignes, None)
            if premiere is None:
                continue
            col.ajouter("heading", feuille.title, level=1, sheet=feuille.title)
            col.tableau(chain([premiere], lignes), sheet=feuille.title)
        col.progresser(len(feuilles), len(feuilles))
    finally:
        classeur.close()
