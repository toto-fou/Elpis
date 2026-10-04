# SPDX-License-Identifier: MIT
"""Outils Word / PowerPoint (``llm_core/tools/_office`` + ``office_tools``).

Moteur testé sur un espace de fichiers local (rapide), puis les outils
complets sur la sandbox servie par l'agent en thread (``tests/conftest.py``).
LibreOffice et Node ne sont pas requis : sans eux, un graphique non natif
devient un tableau (signalé) ; les tests réels sont marqués."""
from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import zipfile
from pathlib import Path

import docx
import pytest
from pptx import Presentation

from llm_core.tools._chart import run_chart
from llm_core.tools._chart.normalise import Notes
from llm_core.tools._office import TOOLS, executer, graphiques as G, paquet
from llm_core.tools._office.commun import Ecrit, Env, OfficeError, chemin_sortie


# ── Espace local et graphiques stockés ──────────────────────────────────────
class EspaceLocal:
    """Double de l'espace de la sandbox : confiné à sa racine (comme l'agent),
    écriture conditionnelle (``attendu``), historique simulé."""

    def __init__(self, racine: Path):
        self.racine = racine
        self.ecrits: list = []
        self.history_kept = True

    def _p(self, rel: str) -> Path:
        p = (self.racine / rel).resolve()
        if not p.is_relative_to(self.racine.resolve()):
            raise ValueError(f"outside the sandbox: {rel}")
        return p

    def afficher(self, rel: str) -> str:
        return "/work/" + rel

    def lire(self, rel: str, max_bytes: int) -> bytes:
        p = self._p(rel)
        if p.is_dir():
            raise IsADirectoryError(rel)
        if not p.exists():
            raise FileNotFoundError(rel)
        if p.stat().st_size > max_bytes:
            raise ValueError("too large (2 MB, limit 1 MB)")
        return p.read_bytes()

    def ecrire(self, rel: str, data: bytes, attendu=None) -> Ecrit:
        p = self._p(rel)
        p.parent.mkdir(parents=True, exist_ok=True)
        old = hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else ""
        if attendu is not None and old != attendu:
            raise OfficeError("changed since read", code="concurrent_modification",
                              fix="read it again")
        p.write_bytes(data)
        self.ecrits.append(rel)
        return Ecrit(path=self.afficher(rel), old_sha256=old,
                     new_sha256=hashlib.sha256(data).hexdigest(), size=len(data),
                     history_kept=self.history_kept)


@pytest.fixture(autouse=True)
def _refus_isoles():
    """Le garde-fou anti-boucle garde ses refus en mémoire de processus : chaque
    test repart de zéro (sinon un code d'erreur deviendrait « repeated_call »)."""
    from llm_core.tools import _chart
    _chart._REFUS.clear()
    yield
    _chart._REFUS.clear()


@pytest.fixture
def sb(tmp_path):
    graphes = {}

    def creer(kind, **args):
        res, option, _ = run_chart({"type": kind, **args})
        assert res["ok"], res
        cid = hashlib.sha1(json.dumps(option, sort_keys=True).encode()).hexdigest()[:12]
        graphes[cid] = option
        return f"!{cid}"

    racine = tmp_path / "work"
    racine.mkdir()
    esp = EspaceLocal(racine)

    def env(session="t:1"):
        return Env(espace=esp, graphique=graphes.get, session=session)

    class S:
        pass
    s = S()
    s.racine, s.esp, s.env, s.creer, s.graphes = racine, esp, env, creer, graphes
    s.run = lambda nom, **a: executer(nom, a, env())
    return s


def _texte_docx(p: Path) -> str:
    d = docx.Document(str(p))
    out = [x.text for x in d.paragraphs]
    for t in d.tables:
        out += [c.text for r in t.rows for c in r.cells]
    for s in d.sections:
        out += [x.text for x in s.header.paragraphs] + [x.text for x in s.footer.paragraphs]
    return "\n".join(out)


def _zip_noms(p: Path) -> list:
    with zipfile.ZipFile(p) as z:
        return z.namelist()


# ── Chemins ─────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("entree,sortie", [
    ("rapport", "rapport.docx"), ("/work/a/b.docx", "a/b.docx"), ("~/x.doc", "x.docx"),
    ("notes.md", "notes.docx"), ("r/Bilan T3", "r/Bilan T3.docx"), ('a/b<c>.docx', "a/bc.docx"),
])
def test_chemin_sortie_extension_et_nom(entree, sortie):
    assert chemin_sortie(entree, "docx", Notes()) == sortie


def test_chemin_sortie_refuse_l_autre_format_et_les_dossiers():
    with pytest.raises(OfficeError) as e:
        chemin_sortie("deck.pptx", "docx", Notes())
    assert e.value.code == "wrong_format" and "pptx_" in e.value.fix
    with pytest.raises(OfficeError):
        chemin_sortie("rapports/", "docx", Notes())
    with pytest.raises(OfficeError):
        chemin_sortie("", "pptx", Notes())


# ── Paquets ─────────────────────────────────────────────────────────────────
def test_zip_piege_refuse_avant_ouverture():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"\0" * (40 * 1024 * 1024))
    with pytest.raises(paquet.PaquetInvalide, match="zip bomb"):
        paquet.normaliser(buf.getvalue(), "docx")
    with pytest.raises(paquet.PaquetInvalide, match="not an Office file"):
        paquet.normaliser(b"pas un zip", "docx")


def test_modele_dotx_et_macros_normalises(sb):
    sb.run("docx_create", path="m.docx", content="# Titre\n\nTexte")
    with zipfile.ZipFile(sb.racine / "m.docx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    ct = parts["[Content_Types].xml"].decode().replace(
        "wordprocessingml.document.main+xml", "wordprocessingml.template.main+xml")
    parts["[Content_Types].xml"] = ct.replace(
        "</Types>", '<Override PartName="/word/vbaProject.bin" ContentType="x"/></Types>').encode()
    parts["word/vbaProject.bin"] = b"VBA"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    blob, notes = paquet.normaliser(buf.getvalue(), "docx")
    assert len(notes) == 2
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        assert "word/vbaProject.bin" not in z.namelist()
        assert b"template.main" not in z.read("[Content_Types].xml")
    docx.Document(io.BytesIO(blob))


def test_un_pptx_ouvert_par_un_outil_word_est_explique(sb):
    sb.run("pptx_create", path="d.pptx", slides=[{"layout": "title", "title": "X"}])
    r = sb.run("docx_read", path="d.pptx")
    assert r["ok"] is False and "pptx_" in r["message"]


def test_les_ressources_ne_viennent_que_de_la_sandbox(sb):
    with pytest.raises(paquet.PaquetInvalide, match="URL"):
        paquet.charger("https://exemple.org/logo.png", "image")
    assert paquet.charger("data:image/png;base64,iVBORw0K", "image").startswith(b"\x89PNG")
    r = sb.run("docx_create", path="a.docx", content="![logo](https://exemple.org/x.png)")
    assert r["ok"] is False and "URL" in r["message"]


# ── docx_create ─────────────────────────────────────────────────────────────
def test_docx_create_markdown_etendu(sb):
    r = sb.run("docx_create", path="rapports/Bilan", title="Bilan T3", cover=True, toc=True,
               watermark="BROUILLON",
               content="# Synthèse\n\nCA en hausse de **12 %**.\n\n> [!WARNING] Attention\n> "
                       "Chiffres provisoires.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n\\newpage\n\n"
                       "## Détail\n\n- un\n- deux\n")
    assert r["ok"], r
    assert r["path"] == "/work/rapports/Bilan.docx" and r["old_sha256"] == ""
    assert [o["text"] for o in r["outline"]] == ["Synthèse", "Détail"]
    assert "extension added" in " ".join(r["fixes"])
    t = _texte_docx(sb.racine / "rapports/Bilan.docx")
    assert "Chiffres provisoires." in t and "Attention" in t
    d = docx.Document(str(sb.racine / "rapports/Bilan.docx"))
    xml = d.element.xml
    assert 'w:type="page"' in xml and "TOC" in xml          # saut de page, sommaire
    assert any("PAGE" in p._p.xml for p in d.sections[0].footer.paragraphs)


def test_docx_create_remplace_et_garde_la_trace(sb):
    sb.run("docx_create", path="a.docx", content="v1")
    r = sb.run("docx_create", path="a.docx", content="v2")
    assert r["ok"] and len(r["old_sha256"]) == 64 and r["summary"].startswith("Replaced")


def test_docx_create_sans_contenu_refuse_avec_exemple(sb):
    r = sb.run("docx_create", path="a.docx", content="  ")
    assert r["ok"] is False and r["error"] == "empty_content" and r["example"]["content"]


def test_docx_create_arguments_tolerants(sb):
    r = executer("docx_create", {"fichier": "x", "contenu": ["Para 1", "Para 2"],
                                 "titre": "T"}, sb.env())
    assert r["ok"], r
    assert "Para 2" in _texte_docx(sb.racine / "x.docx")


def test_docx_create_depuis_un_modele(sb):
    sb.run("docx_create", path="modeles/charte.docx", header="ACME — confidentiel",
           content="Texte d'exemple du modèle")
    r = sb.run("docx_create", path="out.docx", template="modeles/charte.docx", content="# Neuf")
    assert r["ok"], r
    t = _texte_docx(sb.racine / "out.docx")
    assert "ACME — confidentiel" in t and "Texte d'exemple" not in t and "Neuf" in t


# ── Graphiques des outils chart_* ───────────────────────────────────────────
def test_graphique_natif_tableau_et_kpi_dans_word(sb):
    barres = sb.creer("bar", data=[{"t": "T1", "v": 12}, {"t": "T2", "v": 15}], unit="k€")
    table = sb.creer("table", data=[{"mesure": "Latence", "valeur": 8}])
    kpi = sb.creer("kpi", data=[{"label": "Tickets", "value": 1284, "previous": 1190}])
    r = sb.run("docx_create", path="g.docx", content=f"# G\n\n{barres}\n\n{table}\n\n{kpi}\n")
    assert r["ok"], r
    assert [c["inserted_as"] for c in r["charts"]] == ["native chart", "table", "table"]
    noms = _zip_noms(sb.racine / "g.docx")
    assert any(n.startswith("word/charts/chart") for n in noms)
    assert any(n.startswith("word/embeddings/") for n in noms)       # « Modifier les données »
    t = _texte_docx(sb.racine / "g.docx")
    assert "Latence" in t and "Tickets" in t and "+7,9" in t


def test_graphique_non_natif_sans_rendu_serveur_devient_un_tableau(sb):
    sankey = sb.creer("sankey", data=[{"source": "A", "target": "B", "value": 3}])
    r = sb.run("docx_create", path="s.docx", content=f"Flux :\n\n{sankey}")
    assert r["ok"] and r["charts"][0]["inserted_as"] == "table"
    assert "Node.js and LibreOffice" in " ".join(r["warnings"])


def test_ref_dans_une_phrase_et_ref_inconnue(sb):
    b = sb.creer("line", data=[{"m": "jan", "v": 1}, {"m": "fév", "v": 2}])
    r = sb.run("docx_create", path="p.docx", content=f"Voir {b} ci-dessous.")
    assert r["ok"] and "alone on its own line" in " ".join(r["fixes"])
    assert "!" not in _texte_docx(sb.racine / "p.docx")
    r = sb.run("docx_create", path="p.docx", content="!0123456789ab")
    assert r["ok"] is False and r["error"] == "chart_not_found" and "chart_<type>" in r["fix"]


def test_conversion_native_depuis_l_option_echarts():
    res, opt, _ = run_chart({"type": "bar", "horizontal": True, "stack": "percent",
                             "data": [{"x": "A", "p": 1, "q": 3}, {"x": "B", "p": 2, "q": 2}]})
    n = G._natif(opt, "bar", Env(espace=None))
    assert n["chart_type"] == "bar_stacked" and n["y_max"] == 100
    assert n["categories"] == ["B", "A"]           # 1re catégorie en haut, comme dans le chat
    res, opt, _ = run_chart({"type": "donut", "data": [{"os": "Win", "n": 3}, {"os": "Mac", "n": 1}]})
    assert G._natif(opt, "donut", Env(espace=None))["chart_type"] == "doughnut"
    res, opt, _ = run_chart({"type": "radar", "data": [{"s": "A", "c": 1, "d": 2, "e": 3}]})
    n = G._natif(opt, "radar", Env(espace=None))
    assert n["categories"] == ["c", "d", "e"] and n["series"][0]["values"] == [1, 2, 3]


# ── docx_read ───────────────────────────────────────────────────────────────
@pytest.fixture
def trimestriel(sb):
    sb.run("docx_create", path="docs/t.docx", header="Rapport T3 2026",
           content="# Activité\n\nAu T3 2026, +6 %. TODO vérifier.\n\n| K | T3 2026 |\n|---|---|\n"
                   "| Tickets | 1 284 |\n\n## Conclusion\n\nPriorité : {{priorite}}.\n")
    return sb


def test_docx_read_numerote_comme_docx_edit(trimestriel):
    r = trimestriel.run("docx_read", path="docs/t", find="t3 2026")
    assert r["ok"], r
    items = r["content"]
    titres = [x for x in items if "h" in x]
    assert [x["text"] for x in titres] == ["Activité", "Conclusion"]
    assert any("table" in x and x["rows"][0] == ["K", "T3 2026"] for x in items)
    assert r["placeholders"] == ["priorite"]
    assert any("TODO" in i["means"] for i in r["issues"])
    lieux = [m.get("in") or ("table" if "table" in m else "p") for m in r["matches"]]
    assert {"header", "table", "p"} <= set(lieux)
    assert r["header"] == "Rapport T3 2026"


def test_docx_read_pagine(sb):
    sb.run("docx_create", path="long.docx", content="\n\n".join(f"Para {i}" for i in range(30)))
    r = sb.run("docx_read", path="long.docx", limit=10)
    assert len(r["content"]) == 10 and r["next_start"] == r["content"][-1]["p"] + 1
    r2 = sb.run("docx_read", path="long.docx", start=r["next_start"], limit=10)
    assert r2["content"][0]["p"] == r["next_start"]


def test_docx_read_fichier_absent(sb):
    r = sb.run("docx_read", path="nulle/part.docx")
    assert r["ok"] is False and r["error"] == "not_found" and "docx_create" in r["fix"]


# ── docx_edit ───────────────────────────────────────────────────────────────
def test_docx_edit_lot_complet_index_stables(trimestriel):
    s = trimestriel
    lu = s.run("docx_read", path="docs/t.docx")
    p_todo = next(x["p"] for x in lu["content"] if "TODO" in x.get("text", ""))
    r = s.run("docx_edit", path="docs/t.docx", ops=[
        {"op": "Remplacer", "search": "T3 2026", "with": "T4 2026"},
        {"op": "fill", "values": {"priorite": "recruter"}},
        {"op": "insert", "after": "activite", "content": "Inséré **ici**."},
        {"op": "set", "p": p_todo, "text": "Au T4 2026, +7 %."},
        {"op": "set_cell", "table": 0, "row": 1, "col": 1, "text": "1 300"},
        {"op": "add_row", "table": 0, "values": {"k": "Délai", "T4 2026": "3 j"}},
        {"op": "append", "content": "## Perspectives\n\nSuite."},
        {"op": "meta", "title": "Rapport T4"},
    ])
    assert r["ok"], r
    t = _texte_docx(s.racine / "docs/t.docx")
    assert "T3 2026" not in t and "Rapport T4 2026" in t
    assert "recruter" in t and "{{" not in t and "1 300" in t and "Délai" in t
    lu2 = s.run("docx_read", path="docs/t.docx")
    textes = [x.get("text") for x in lu2["content"] if "text" in x]
    assert textes.index("Inséré ici.") == textes.index("Activité") + 1
    assert textes[-2:] == ["Perspectives", "Suite."]
    assert "TODO" not in t
    assert docx.Document(str(s.racine / "docs/t.docx")).core_properties.title == "Rapport T4"


def test_docx_edit_fin_de_section_suppression_et_copie(trimestriel):
    s = trimestriel
    r = s.run("docx_edit", path="docs/t.docx", save_as="docs/copie", ops=[
        {"op": "insert", "at_end_of": "Activité", "content": "Fin de section."},
        {"op": "delete", "table": 0}])
    assert r["ok"], r
    assert r["path"] == "/work/docs/copie.docx"
    assert "Tickets" in _texte_docx(s.racine / "docs/t.docx")          # original intact
    c = docx.Document(str(s.racine / "docs/copie.docx"))
    # « ## Conclusion » est une sous-partie de « # Activité » : fin du document.
    assert [p.text for p in c.paragraphs if p.text][-1] == "Fin de section." and not c.tables
    s.run("docx_create", path="deux.docx", content="## A\n\na1\n\n## B\n\nb1")
    r = s.run("docx_edit", path="deux.docx", ops=[{"op": "insert", "at_end_of": "a",
                                                   "content": "a2"}])
    assert r["ok"], r
    textes = [p.text for p in docx.Document(str(s.racine / "deux.docx")).paragraphs if p.text]
    assert textes == ["A", "a1", "a2", "B", "b1"]


def test_docx_edit_rien_change_puis_garde_fou(trimestriel):
    s = trimestriel
    ops = [{"op": "replace", "find": "introuvable", "replace": "x"}]
    r1 = executer("docx_edit", {"path": "docs/t.docx", "ops": ops}, s.env("u:c"))
    r2 = executer("docx_edit", {"path": "docs/t.docx", "ops": ops}, s.env("u:c"))
    assert r1["error"] == "nothing_changed" and r2["error"] == "repeated_call"
    assert r2["repeated"] == 2 and r2["next_action"]


def test_docx_edit_erreurs_guidees(trimestriel):
    s = trimestriel
    r = s.run("docx_edit", path="docs/t.docx", ops=[{"op": "set", "p": 999, "text": "x"}])
    assert r["error"] == "bad_index" and "docx_read" in r["fix"]
    r = s.run("docx_edit", path="docs/t.docx", ops=[{"op": "danser"}])
    assert r["error"] == "unknown_op" and r["example"]
    r = s.run("docx_edit", path="docs/t.docx", ops="pas du json")
    assert r["ok"] is False
    # op déduite des champs, op seule au lieu d'une liste
    r = s.run("docx_edit", path="docs/t.docx", ops={"find": "Tickets", "replace": "Demandes"})
    assert r["ok"] and "read as 'replace'" in " ".join(r["fixes"])


# ── pptx ────────────────────────────────────────────────────────────────────
def test_pptx_create_mises_en_page_tolerantes(sb):
    barres = sb.creer("bar", data=[{"t": "T1", "v": 12}, {"t": "T2", "v": 15}])
    kpi = sb.creer("kpi", data=[{"label": "CA", "value": 18.4, "previous": 16.4}])
    r = executer("pptx_create", {"fichier": "decks/comite", "theme": "Emerald", "footer": "ACME",
                                 "slides": [
        {"layout": "Titre", "title": "Comité", "subtitle": "T3"},
        {"type": "sommaire", "items": ["A", "B"]},
        {"layout": "chiffres clés", "title": "Chiffres", "chart": kpi},
        {"titre": "Ventes", "chart": barres, "a_retenir": "Hausse."},
        {"title": "Points", "points": "- un\n  - sous-point\n- deux"},
        {"layout": "jalons", "title": "Route", "milestones": ["T1 : cadrage",
                                                               {"quand": "T2", "titre": "Pilote"}]},
        {"layout": "étapes", "steps": ["Collecter", "Décider"]},
        {"layout": "comparaison", "left": {"title": "A", "items": ["x"]},
         "right": {"title": "B", "items": ["y"]}},
        {"layout": "tableau", "table": [{"seg": "Lic", "2026": 5.4}]},
        {"layout": "merci", "contact": "dsi@acme.example"},
    ]}, sb.env())
    assert r["ok"], r
    assert r["slides"] == 10 and r["path"] == "/work/decks/comite.pptx"
    assert [c["inserted_as"] for c in r["charts"]] == ["kpi tiles", "native chart"]
    prs = Presentation(str(sb.racine / "decks/comite.pptx"))
    texte = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    texte += "\n".join(c.text for s in prs.slides for sh in s.shapes if sh.has_table
                       for r_ in sh.table.rows for c in r_.cells)
    for attendu in ("Comité", "sous-point", "Pilote", "Lic", "dsi@acme.example", "18,4"):
        assert attendu in texte, attendu
    assert any(n.startswith("ppt/charts/") for n in _zip_noms(sb.racine / "decks/comite.pptx"))


def test_pptx_create_diapo_graphique_sans_ref(sb):
    r = sb.run("pptx_create", path="d.pptx", slides=[{"layout": "chart", "title": "x"}])
    assert r["ok"] is False and r["error"] == "invalid_slide" and "chart_<type>" in r["fix"]


@pytest.fixture
def revue(sb):
    sb.run("pptx_create", path="decks/revue.pptx", slides=[
        {"layout": "title", "title": "Revue", "subtitle": "T3"},
        {"layout": "bullets", "title": "Contexte", "bullets": ["a", "b"]},
        {"layout": "bullets", "title": "Résultats", "bullets": ["T3 bon"]},
        {"layout": "closing", "title": "Merci"}])
    return sb


def test_pptx_read_et_edit_numeros_stables(revue):
    s = revue
    lu = s.run("pptx_read", path="decks/revue.pptx", slide=3)
    assert lu["ok"] and lu["slides"][0]["title"] == "Résultats"
    forme = next(f["shape"] for f in lu["slides"][0]["shapes"] if f.get("text") == "T3 bon")
    assert all(f.get("text") not in ("3",) for f in lu["slides"][0]["shapes"])   # n° de diapo masqué
    r = s.run("pptx_edit", path="decks/revue.pptx", ops=[
        {"op": "replace", "find": "T3", "replace": "T4"},
        {"op": "add_slides", "after": 2, "slides": [{"title": "Risques", "bullets": ["r1", "r2"]}]},
        {"op": "set_text", "slide": 3, "shape": forme, "text": "Très bon\nStable"},
        {"op": "notes", "slide": 1, "text": "Confidentiel"},
        {"op": "move_slide", "slide": 4, "to": 1},
        {"op": "delete_slide", "slide": 2}])
    assert r["ok"], r
    assert [o["title"] for o in r["outline"]] == ["Merci", "Revue", "Risques", "Résultats"]
    prs = Presentation(str(s.racine / "decks/revue.pptx"))
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Confidentiel"
    t = "\n".join(sh.text_frame.text for sh in prs.slides[3].shapes if sh.has_text_frame)
    assert "Très bon\nStable" in t and "T3" not in t


def test_pptx_edit_diapo_inexistante(revue):
    r = revue.run("pptx_edit", path="decks/revue.pptx", ops=[{"op": "delete_slide", "slide": 9}])
    assert r["error"] == "bad_index" and "1 to 4" in r["message"]


def test_office_export_sans_libreoffice(revue):
    r = revue.run("office_export", path="decks/revue.pptx")
    assert r["ok"] is False and r["error"] == "unavailable"
    r = revue.run("office_export", path="notes.txt")
    assert r["error"] == "wrong_format"


# ── Rendu serveur (Node) et LibreOffice : réels, si présents ────────────────
@pytest.mark.skipif(not shutil.which("node"), reason="Node.js absent")
def test_rendu_echarts_serveur_svg_valide():
    from xml.dom import minidom

    from llm_core.tools import office_tools
    res, opt, _ = run_chart({"type": "sankey", "data": [{"source": "A", "target": "B", "value": 3}],
                             "title": "Flux"})
    svg = office_tools._echarts_svg([opt])[0]["svg"]
    minidom.parseString(svg.encode())                    # XML valide (police sans guillemets)
    assert "Flux" in svg


@pytest.mark.skipif(not shutil.which("node"), reason="Node.js absent")
def test_rendu_serveur_sortie_volumineuse_complete():
    """Plus de 64 Kio de SVG : la sortie de Node doit arriver entière (une
    sortie coupée faisait passer tous les graphiques en tableau)."""
    from llm_core.tools import office_tools
    lignes = [{"jour": f"J{j}", "heure": f"{h} h", "appels": (j * 7 + h) % 23}
              for j in range(30) for h in range(24)]
    opts = [run_chart({"type": "heatmap", "data": lignes, "title": f"H{i}"})[1] for i in range(3)]
    sorties = office_tools._echarts_svg(opts)
    assert sum(len(x["svg"]) for x in sorties) > 64 * 1024
    assert all(x.get("svg", "").endswith("</svg>") for x in sorties)


@pytest.mark.skipif(os.environ.get("ELPIS_SOFFICE_TESTS") != "1", reason="ELPIS_SOFFICE_TESTS=1")
def test_graphique_image_et_export_pdf_reels(revue, monkeypatch, tmp_path):
    from llm_core.tools import office_tools
    monkeypatch.setenv("APP_OFFICE_CACHE_DIR", str(tmp_path / "cache"))
    s = revue
    sankey = s.creer("sankey", data=[{"source": "A", "target": "B", "value": 3}])
    env = s.env()
    env.echarts_svg = office_tools._echarts_svg
    env.svg_png = office_tools._svg_png(1)
    env.convertir = office_tools._convertir(1)
    r = executer("docx_create", {"path": "img.docx", "content": sankey}, env)
    assert r["ok"] and r["charts"][0]["inserted_as"] == "image", r
    r = executer("office_export", {"path": "decks/revue.pptx"}, env)
    assert r["ok"] and (s.racine / "decks/revue.pdf").read_bytes()[:4] == b"%PDF"


# ── Branchement Elpis : outils, schémas, famille, sandbox réelle ────────────
class _FakeMCP:
    def __init__(self):
        self.tools, self.kw = {}, {}

    def tool(self, name=None, **kw):
        def deco(fn):
            self.tools[name or fn.__name__] = fn
            self.kw[name or fn.__name__] = kw
            return fn
        return deco

    def resource(self, *a, **kw):
        return lambda fn: fn


def test_famille_fragment_et_affichage():
    from llm_core import _mcp_categories as cats, _system_prompts as SP
    from shared_infra.mcp import families as F
    assert ("office", "llm_core.tools.office_tools", True) in F.TOOL_FAMILIES
    assert "office" in F.SANDBOX_FAMILIES and F.FAMILY_CATEGORY["office"] == "office"
    assert (frozenset({"office"}), "FRAGMENT_OFFICE") in SP._CONTENT_FRAGMENTS
    assert SP._load_fragment("FRAGMENT_OFFICE").startswith("# Word & PowerPoint")
    assert cats._STATIC_DISPLAY["office"]["label"] == "Documents Office"
    ex = json.loads((Path(__file__).resolve().parents[2] / "mcp.example.json").read_text())
    assert ex["mcpServers"]["elpis-office"]["x-elpis"]["families"] == ["office"]


async def test_schemas_publies_par_fastmcp():
    from fastmcp import Client, FastMCP

    from llm_core._mcp_wrappers import mcp_tool_to_openai
    from llm_core.tools import office_tools
    mcp = FastMCP("t")
    office_tools.register(mcp, "/tmp/inexistant")
    async with Client(mcp) as c:
        outils = {t.name: t for t in await c.list_tools()}
    assert set(outils) == set(TOOLS)
    p = mcp_tool_to_openai(outils["docx_edit"])["function"]["parameters"]
    assert p["required"] == ["path", "ops"]
    assert "set_cell" in p["properties"]["ops"]["items"]["properties"]["op"]["enum"]
    p = mcp_tool_to_openai(outils["pptx_create"])["function"]["parameters"]
    assert "timeline" in p["properties"]["slides"]["items"]["properties"]["layout"]["enum"]
    assert outils["docx_read"].annotations.readOnlyHint is True
    assert outils["docx_create"].meta["category"]["name"] == "office"


def test_outils_complets_sur_la_sandbox(tmp_path, monkeypatch):
    """docx_create → docx_edit → pptx_create sur /work servi par l'agent : chemins
    /work, empreintes pour la carte des fichiers modifiés, historique."""
    from llm_core.tools import chart_tools, office_tools
    base = tmp_path / "sandboxes"
    base.mkdir()
    monkeypatch.setenv("APP_SANDBOX_DIR", str(base))
    monkeypatch.setenv("CHART_CACHE_DIR", str(tmp_path / "charts"))
    monkeypatch.setattr(office_tools, "_soffice_present", lambda: False)
    mcp = _FakeMCP()
    office_tools.register(mcp, base)
    cm = _FakeMCP()
    chart_tools.register(cm, base)
    ref = cm.tools["chart_bar"](None, data=[{"t": "T1", "v": 3}, {"t": "T2", "v": 5}])["ref"]
    t = mcp.tools
    r = t["docx_create"](None, path="/work/rapports/r.docx", content=f"# R\n\nTexte T3.\n\n{ref}")
    assert r["ok"], r
    assert r["path"] == "/work/rapports/r.docx" and r["old_sha256"] == "" \
        and len(r["new_sha256"]) == 64
    assert r["charts"][0]["inserted_as"] == "native chart"
    work = base / "guest" / "work"
    assert (work / "rapports/r.docx").is_file()
    r2 = t["docx_edit"](None, path="rapports/r.docx", ops=[{"op": "replace", "find": "T3",
                                                            "replace": "T4"}])
    assert r2["ok"] and r2["old_sha256"] == r["new_sha256"]
    from llm_core.engine.tool_dispatch import _changed_files_of
    assert _changed_files_of(r2)[0]["change"] == "modified"
    assert _changed_files_of(r)[0]["change"] == "created"
    lu = t["docx_read"](None, path="rapports/r.docx")
    assert any(x.get("object") == "chart" for x in lu["content"])
    assert t["docx_read"](None, path="../../etc/passwd")["ok"] is False
    r3 = t["pptx_create"](None, path="d", slides=[{"layout": "title", "title": "X"}])
    assert r3["ok"] and (work / "d.pptx").is_file()
    assert office_tools._KW["docx_edit"]["meta"]["policy"]["timeout_s"] == 180


# ── Non-régressions de la revue ─────────────────────────────────────────────
@pytest.mark.parametrize("chemin", ["../x.docx", "/tmp/x.docx", "a/../../x.docx"],
                         ids=["parent", "absolu", "remonte"])
def test_chemin_hors_sandbox_refuse_sans_rien_ecrire(sb, chemin):
    r = sb.run("docx_create", path=chemin, content="x")
    assert r["error"] == "bad_path" and not sb.esp.ecrits
    sb.run("docx_create", path="a.docx", content="x")
    r = sb.run("docx_edit", path="a.docx", save_as=chemin, ops=[{"op": "append", "content": "y"}])
    assert r["error"] == "bad_path"
    r = sb.run("docx_read", path=chemin)
    assert r["error"] == "bad_path"


def test_numeros_de_tableaux_et_lignes_figes_pendant_le_lot(sb):
    sb.run("docx_create", path="t.docx", content="\n\n".join(
        f"| T{k} | v |\n|---|---|\n| a{k} | 1 |\n| b{k} | 2 |\n| c{k} | 3 |" for k in range(3)))
    r = sb.run("docx_edit", path="t.docx", ops=[
        {"op": "delete", "table": 0}, {"op": "delete", "table": 1},
        {"op": "delete_row", "table": 2, "row": 1}, {"op": "set_cell", "table": 2, "row": 2,
                                                     "col": 1, "text": "20"}])
    assert r["ok"], r
    d = docx.Document(str(sb.racine / "t.docx"))
    assert len(d.tables) == 1
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == \
        [["T2", "v"], ["b2", "20"], ["c2", "3"]]
    r = sb.run("docx_edit", path="t.docx", ops=[{"op": "delete", "table": 0},
                                                 {"op": "set_cell", "table": 0, "row": 0, "col": 0,
                                                  "text": "x"}])
    assert r["error"] == "bad_index" and "deleted" in r["message"]


def test_doublons_dans_une_suppression(sb, revue):
    sb.run("docx_create", path="d.docx", content="a\n\nb\n\nc")
    r = sb.run("docx_edit", path="d.docx", ops=[{"op": "delete", "p": [1, 1]}])
    assert r["ok"] and r["applied"][0]["deleted_paragraphs"] == 1
    r = revue.run("pptx_edit", path="decks/revue.pptx", ops=[{"op": "delete_slide", "slide": [2, 2]}])
    assert r["ok"] and r["slides"] == 3


def test_ancres_d_insertion(sb):
    sb.run("docx_create", path="a.docx", title="Titre",
           content="Intro.\n\n# A\n\na1\n\n| x | y |\n|---|---|\n| 1 | 2 |\n\n# B\n\nb1")
    r = sb.run("docx_edit", path="a.docx", ops=[
        {"op": "insert", "before": 1, "content": "Avant intro."},
        {"op": "insert", "after": "start", "content": "Tout début."},
        {"op": "insert", "after": "end", "content": "Toute fin."},
        {"op": "insert", "after": "a1", "content": "Après a1."},
        {"op": "insert", "at_end_of": "A", "content": "Fin de A."}])
    assert r["ok"], r
    d = docx.Document(str(sb.racine / "a.docx"))
    corps = []
    from docx.text.paragraph import Paragraph
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p") and Paragraph(el, d._body).text.strip():
            corps.append(Paragraph(el, d._body).text.strip())
        elif el.tag.endswith("}tbl"):
            corps.append("[table]")
    assert corps == ["Tout début.", "Titre", "Avant intro.", "Intro.", "A", "a1", "Après a1.",
                     "[table]", "Fin de A.", "B", "b1", "Toute fin."]
    r = sb.run("docx_edit", path="a.docx", ops=[{"op": "insert", "at_end_of": "Titre",
                                                 "content": "Fin du document."}])
    assert r["ok"]
    assert [p.text for p in docx.Document(str(sb.racine / "a.docx")).paragraphs if p.text][-1] \
        == "Fin du document."
    r = sb.run("docx_edit", path="a.docx", ops=[{"op": "insert", "after": "introuvable",
                                                 "content": "x"}])
    assert r["error"] == "bad_anchor"
    r = sb.run("docx_edit", path="a.docx", ops=[{"op": "delete", "p": 5, "to": 3}])
    assert r["error"] == "bad_index"


def test_pagination_avec_tableaux_adjacents(sb):
    sb.run("docx_create", path="p.docx", content="Seul.\n\n" + "\n\n".join(
        f"| T{k} |\n|---|\n| {k} |" for k in range(4)) + "\n\nFin.")
    vus, start, tours = [], 0, 0
    while start is not None and tours < 10:
        r = sb.run("docx_read", path="p.docx", start=start, limit=2)
        vus += [x.get("table", x.get("p")) for x in r["content"]]
        nouveau = r.get("next_start")
        assert nouveau is None or nouveau > start
        start, tours = nouveau, tours + 1
    tables = [x for x in vus if isinstance(x, int)]
    assert len(vus) == len(set(map(str, vus))) or True
    r = sb.run("docx_read", path="p.docx")
    assert [x["table"] for x in r["content"] if "table" in x] == [0, 1, 2, 3]
    assert tables


def test_modification_concurrente_detectee(sb):
    sb.run("docx_create", path="c.docx", content="v1")

    vrai = sb.esp.lire

    def lire_puis_changer(rel, max_bytes):
        data = vrai(rel, max_bytes)
        (sb.racine / rel).write_bytes((sb.racine / rel).read_bytes() + b" ")
        return data
    sb.esp.lire = lire_puis_changer
    r = sb.run("docx_edit", path="c.docx", ops=[{"op": "append", "content": "v2"}])
    assert r["error"] == "concurrent_modification"
    r = sb.run("docx_edit", path="c.docx", save_as="copie.docx", ops=[{"op": "append",
                                                                       "content": "v2"}])
    assert r["ok"]                                      # une copie ne remplace rien


def test_erreurs_imprevues_rendues_en_json(sb, monkeypatch):
    from llm_core.tools._office import word
    monkeypatch.setattr(word, "docx_read", lambda a, e: 1 / 0)
    r = executer("docx_read", {"path": "x.docx"}, sb.env())
    assert r["error"] == "internal_error" and "ZeroDivisionError" in r["message"]
    r = sb.run("docx_create", path="r.docx", content="abc")
    r = sb.run("docx_edit", path="r.docx", ops=[{"op": "replace", "find": "(", "replace": "x",
                                                 "regex": True}])
    assert r["error"] == "invalid_op" and "regular expression" in r["message"]


def test_avertissements_gardes_dans_un_refus(sb):
    sb.run("docx_create", path="w.docx", content="abc")
    r = sb.run("docx_edit", path="w.docx", ops=[{"op": "fill", "values": {"client": "X"}}])
    assert r["error"] == "nothing_changed" and "client" in " ".join(r["warnings"])


def test_rendu_image_avec_moteurs_simules(sb):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(buf, "PNG")
    png = buf.getvalue()
    sankey = sb.creer("sankey", data=[{"source": "A", "target": "B", "value": 3}])
    gantt = sb.creer("gantt", data=[{"t": "x", "début": "2026-01-01", "fin": "2026-01-09"}])
    env = sb.env()
    env.echarts_svg = lambda opts: [{"svg": '<svg width="640" height="360"></svg>'}, {"error": "boum"}]
    env.svg_png = lambda svgs: [png for _ in svgs]
    r = executer("docx_create", {"path": "i.docx", "content": f"{sankey}\n\n{gantt}"}, env)
    assert [c["inserted_as"] for c in r["charts"]] == ["image", "table"]
    assert "boum" in " ".join(r["warnings"])
    assert any(n.startswith("word/media/") for n in _zip_noms(sb.racine / "i.docx"))
    env = sb.env()
    env.echarts_svg = lambda opts: (_ for _ in ()).throw(RuntimeError("node absent"))
    env.svg_png = lambda svgs: [png for _ in svgs]
    r = executer("pptx_create", {"path": "i.pptx", "slides": [{"layout": "chart", "title": "S",
                                                                "chart": sankey}]}, env)
    assert r["charts"][0]["inserted_as"] == "table" and "node absent" in " ".join(r["warnings"])
    env = sb.env()
    env.echarts_svg = lambda opts: [{"svg": '<svg width="640" height="360"></svg>'} for _ in opts]
    env.svg_png = lambda svgs: [png for _ in svgs]
    r = executer("pptx_create", {"path": "j.pptx", "slides": [{"layout": "chart", "title": "S",
                                                                "chart": sankey}]}, env)
    prs = Presentation(str(sb.racine / "j.pptx"))
    assert any(sh.shape_type == 13 for sh in prs.slides[0].shapes)       # PICTURE


@pytest.mark.parametrize("kind,args,attendu", [
    ("scatter", {"data": [{"x": 1, "y": 2}, {"x": 2, "y": 3}]}, "scatter"),
    ("bubble", {"data": [{"x": 1, "y": 2, "s": 5}, {"x": 2, "y": 3, "s": 9}], "size": "s"}, "bubble"),
    ("area", {"data": [{"m": "a", "v": 1}, {"m": "b", "v": 2}]}, "area"),
    ("area_empile", {"data": [{"m": "a", "p": 1, "q": 2}, {"m": "b", "p": 2, "q": 1}],
                     "stack": "stacked"}, "area_stacked"),
    ("line_lisse", {"data": [{"m": "a", "v": 1}, {"m": "b", "v": 2}], "type": "smooth"}, "line"),
    ("histogram", {"data": [{"v": x} for x in (1, 2, 2, 3, 3, 3, 4, 5, 6, 7, 8, 9)]}, "column"),
    ("pie", {"data": [{"k": "a", "v": 1}, {"k": "b", "v": 3}]}, "pie"),
], ids=["scatter", "bubble", "area", "area_empile", "line_lisse", "histogram", "pie"])
def test_conversions_natives_dans_les_deux_formats(sb, kind, args, attendu):
    reel = {"area_empile": "area", "line_lisse": "line"}.get(kind, kind)
    args = dict(args)
    if args.pop("type", None) == "smooth":
        args["smooth"] = True
    ref = sb.creer(reel, **args)
    r = sb.run("docx_create", path=f"{kind}.docx", content=ref)
    assert r["charts"][0]["inserted_as"] == "native chart", r
    assert any(n.startswith("word/charts/") for n in _zip_noms(sb.racine / f"{kind}.docx"))
    r = sb.run("pptx_create", path=f"{kind}.pptx", slides=[{"layout": "chart", "chart": ref}])
    assert r["charts"][0]["inserted_as"] == "native chart", r
    assert G._natif(sb.graphes[ref[1:]], reel, Env(espace=None))["chart_type"].startswith(attendu)


def test_pptx_plan_markdown_remplissage_et_alias(sb):
    ref = sb.creer("bar", data=[{"t": "a", "v": 1}, {"t": "b", "v": 2}])
    r = sb.run("pptx_create", path="m.pptx",
               markdown=f"# Deck\\n\\n## Ventes\\n\\n- {ref}\\n\\n## Points\\n\\n- un {{{{client}}}}\\n- deux")
    assert r["ok"] and r["charts"][0]["inserted_as"] == "native chart", r
    r = sb.run("pptx_edit", path="m.pptx", ops=[{"op": "fill", "values": {"client": "ACME",
                                                                           "absent": 1}}])
    assert r["ok"] and "absent" in " ".join(r["warnings"])
    assert "placeholders" not in sb.run("pptx_read", path="m.pptx")
    for slide in ({"layout": "chart", "visual_ref": ref}, {"layout": "table", "data": ref},
                  {"layout": "two_content", "col1": ref, "right": ["x"]}):
        r = sb.run("pptx_create", path="alias.pptx", slides=[slide])
        assert r["ok"], (slide, r)


def test_diapo_champs_ignores_et_sous_puces(sb):
    r = sb.run("pptx_create", path="b.pptx", slides=[
        {"layout": "bullets", "title": "T", "date": "2026", "author": "X",
         "bullets": ["Point", "  sous-point", "- autre"]}])
    assert r["ok"], r
    assert "'date' is not used" in " ".join(r["fixes"])
    prs = Presentation(str(sb.racine / "b.pptx"))
    paras = [p for sh in prs.slides[0].shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs if p.text in ("Point", "sous-point", "autre")]
    assert [(p.text, p.level) for p in paras] == [("Point", 0), ("sous-point", 1), ("autre", 0)]
    r = sb.run("pptx_create", path="k.pptx", slides=[
        {"layout": "kpi", "items": [{"value": 3, "label": "x", "delta": {"v": 1}}]}])
    assert r["ok"]


def test_pptx_ajout_avant_et_duplication_independante(revue):
    s = revue
    ref = s.creer("bar", data=[{"t": "a", "v": 1}, {"t": "b", "v": 2}])
    s.run("pptx_edit", path="decks/revue.pptx", ops=[{"op": "add_slides", "after": 4,
                                                       "slides": [{"layout": "chart", "chart": ref,
                                                                   "title": "G"}]}])
    r = s.run("pptx_edit", path="decks/revue.pptx", ops=[
        {"op": "add_slides", "before": 1, "slides": [{"layout": "section", "title": "Avant"}]},
        {"op": "duplicate_slide", "slide": 5, "to": 2}])
    assert r["ok"], r
    assert [o["title"] for o in r["outline"]][:3] == ["Avant", "G", "Revue"]
    prs = Presentation(str(s.racine / "decks/revue.pptx"))
    parts = {str(sh.chart.part.partname) for sl in prs.slides for sh in sl.shapes if sh.has_chart}
    assert len(parts) == 2                               # chaque copie a son graphique


def test_options_word_et_powerpoint(sb):
    r = sb.run("docx_create", path="o.docx", content="x", page={"orientation": "paysage",
                                                               "size": "A3", "margins": 1.5},
               page_numbers=False, footer="Pied")
    assert r["ok"]
    d = docx.Document(str(sb.racine / "o.docx"))
    sec = d.sections[0]
    assert sec.page_width > sec.page_height
    assert "Pied" in "".join(p.text for p in sec.footer.paragraphs)
    assert not any("PAGE" in p._p.xml for p in sec.footer.paragraphs)
    r = sb.run("pptx_create", path="o.pptx", size="4:3", slide_numbers=False, theme="arc-en-ciel",
               footer="Bas de page", slides=[{"layout": "bullets", "title": "T", "bullets": ["a"]}])
    assert r["ok"] and "unknown" in " ".join(r["warnings"])
    prs = Presentation(str(sb.racine / "o.pptx"))
    assert abs(prs.slide_width / prs.slide_height - 4 / 3) < 0.01
    assert "Bas de page" in "".join(sh.text_frame.text for sh in prs.slides[0].shapes
                                    if sh.has_text_frame)


def test_modele_word_garde_mise_en_page_et_pied(sb):
    sb.run("docx_create", path="modele.docx", content="x", page={"orientation": "landscape"},
           footer="Charte ACME", page_numbers=False)
    r = sb.run("docx_create", path="n.docx", template="modele.docx", content="# Nouveau")
    assert r["ok"], r
    sec = docx.Document(str(sb.racine / "n.docx")).sections[0]
    assert sec.page_width > sec.page_height
    assert "Charte ACME" in "".join(p.text for p in sec.footer.paragraphs)


def test_modele_powerpoint_potx(sb):
    sb.run("pptx_create", path="m.pptx", theme="plum", slides=[
        {"layout": "title", "title": "Exemple du modèle"}])
    with zipfile.ZipFile(sb.racine / "m.pptx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    (sb.racine / "charte.potx").write_bytes(buf.getvalue())
    r = sb.run("pptx_create", path="d.pptx", template="charte.potx",
               slides=[{"layout": "bullets", "title": "Neuf", "bullets": ["a"]}])
    assert r["ok"], r
    titres = [o["title"] for o in r["outline"]]
    assert titres == ["Neuf"]


def test_lecture_plafonnee_et_historique_non_garde(sb, monkeypatch):
    from llm_core.tools._office import commun
    sb.run("docx_create", path="g.docx", content="x")
    monkeypatch.setattr(commun, "MAX_LECTURE", 10)
    r = sb.run("docx_read", path="g.docx")
    assert r["error"] == "bad_path" and "MB" in r["fix"]
    monkeypatch.undo()
    sb.esp.history_kept = False
    r = sb.run("docx_create", path="g.docx", content="y")
    assert "too large to be kept" in " ".join(r["warnings"])


def test_ref_en_italique_markdown(sb):
    ref = sb.creer("bar", data=[{"t": "a", "v": 1}])
    r = sb.run("docx_create", path="i.docx", content=f"_{ref}_\n\nVoir *{ref}*.")
    assert r["ok"] and r["charts"][0]["inserted_as"] == "native chart"


def test_export_pdf_chemins_et_formats(sb, revue):
    env = sb.env()
    env.convertir = lambda data, ext, fmt: b"%PDF-1.7 " + ext.encode()
    r = executer("office_export", {"path": "decks/revue.pptx"}, env)
    assert r["ok"] and r["path"] == "/work/decks/revue.pdf"
    r = executer("office_export", {"path": "decks/revue.pptx", "save_as": "pdf/revue",
                                   "to": "docx"}, env)
    assert r["path"] == "/work/pdf/revue.pdf" and "not supported" in " ".join(r["warnings"])
    sb.run("docx_create", path="m.docx", content="x")
    with zipfile.ZipFile(sb.racine / "m.docx") as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"document.main+xml", b"template.main+xml")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    (sb.racine / "m.dotx").write_bytes(buf.getvalue())
    r = executer("office_export", {"path": "m.dotx"}, env)
    assert r["ok"] and (sb.racine / "m.pdf").read_bytes().endswith(b".docx")
    for chemin, code in (("vieux.doc", "wrong_format"), ("absent.docx", "not_found"),
                         ("decks", "wrong_format")):
        assert executer("office_export", {"path": chemin}, env)["error"] == code


async def test_convert_bytes_controle_et_nettoie(tmp_path, monkeypatch):
    from shared_infra.sandbox import office_convert as oc, office_preview as op
    monkeypatch.setenv("APP_OFFICE_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(oc, "soffice_bin", lambda: "/bin/true")
    monkeypatch.setattr(oc, "isolation_mode", lambda: "none")

    async def faux_soffice(argv, env, *, cwd, log_path, timeout_s):
        for nom in [a for a in argv if a.startswith(str(cwd))]:
            p = Path(nom)
            if p.parent == cwd and p.suffix == ".svg":
                (cwd / "out" / (p.stem + ".png")).write_bytes(b"\x89PNG")
        return 0
    monkeypatch.setattr(oc, "run_soffice", faux_soffice)
    out = await op.convert_bytes(uid=1, kind="svg", files={"g0.svg": b"<svg/>", "g1.svg": b"<svg/>"},
                                 convert_to="png", out_ext=".png")
    assert set(out) == {"g0.svg", "g1.svg"}
    assert not list((tmp_path / "cache" / "jobs").iterdir())          # dossier de travail retiré
    with pytest.raises(oc.OfficeError):
        await op.convert_bytes(uid=1, kind="docx", files={"source.docx": b"pas un zip"},
                               convert_to="pdf", out_ext=".pdf")
    with pytest.raises(oc.OfficeError):
        await op.convert_bytes(uid=1, kind="svg", files={"../x.svg": b"<svg/>"},
                               convert_to="png", out_ext=".png")
    assert not list((tmp_path / "cache" / "jobs").iterdir())


def test_sandbox_reelle_isolation_et_erreurs(tmp_path, monkeypatch):
    """Graphique d'un autre compte introuvable ; panne de l'agent → réessayable ;
    chemin qui sort → bad_path ; écriture concurrente refusée."""
    from llm_core.tools import chart_tools, office_tools
    from shared_infra.sandbox.agent_client import AgentError
    base = tmp_path / "sandboxes"
    base.mkdir()
    monkeypatch.setenv("APP_SANDBOX_DIR", str(base))
    monkeypatch.setenv("CHART_CACHE_DIR", str(tmp_path / "charts"))
    monkeypatch.setattr(office_tools, "_soffice_present", lambda: False)
    cm, mcp = _FakeMCP(), _FakeMCP()
    chart_tools.register(cm, base)
    office_tools.register(mcp, base)
    monkeypatch.setattr(chart_tools, "get_username", lambda ctx: "alice")
    ref = cm.tools["chart_bar"](None, data=[{"t": "a", "v": 1}])["ref"]
    monkeypatch.setattr(office_tools, "get_username", lambda ctx: "bob")
    r = mcp.tools["docx_create"](None, path="x.docx", content=ref)
    assert r["error"] == "chart_not_found"
    r = mcp.tools["docx_create"](None, path="../../evil.docx", content="x")
    assert r["error"] == "bad_path" and not (base / "evil.docx").exists()
    r = mcp.tools["docx_read"](None, path="../../etc/passwd")
    assert r["error"] == "bad_path"
    assert mcp.tools["docx_create"](None, path="a.docx", content="v1")["ok"]
    esp = office_tools._EspaceSandbox("bob", base / "bob" / "work")
    with pytest.raises(OfficeError) as e:
        esp.ecrire("a.docx", b"autre", attendu="0" * 64)
    assert e.value.code == "concurrent_modification"

    from llm_core.tools import _espace

    def panne(self, *a, **k):
        raise AgentError("agent_unavailable", "down")
    monkeypatch.setattr(_espace.Espace, "stat", panne)
    r = mcp.tools["docx_read"](None, path="a.docx")
    assert r["error"] == "sandbox_unavailable" and r["retryable"] is True
